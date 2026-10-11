#!/usr/bin/env python3
"""Create an editable DOCX projection from a production HTML artifact."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree as lxml_etree
from lxml import html as lxml_html
from PIL import Image


BLOCK_TAGS = {"blockquote", "div", "figure", "h1", "h2", "h3", "h4", "h5", "h6", "nav", "ol", "p", "section", "table", "ul"}
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_NAMESPACES = {"w": WORD_NS}


@dataclass
class BuildState:
    force_next_heading_page: bool = False
    paragraphs: int = 0
    tables: int = 0
    images: int = 0
    toc_entries: int = 0
    callouts: int = 0
    next_bookmark_id: int = 1
    missing_images: list[str] | None = None

    def __post_init__(self) -> None:
        if self.missing_images is None:
            self.missing_images = []


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def set_run_font(run: Any, name: str, size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)


def set_style_font(style: Any, name: str, size: float, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)


def remove_style_paragraph_borders(style: Any) -> None:
    paragraph_properties = style.element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(17)
    section.right_margin = Mm(17)

    normal = doc.styles["Normal"]
    set_style_font(normal, "Times New Roman", 10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True

    set_style_font(doc.styles["Title"], "Times New Roman", 22, True)
    remove_style_paragraph_borders(doc.styles["Title"])
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Title"].paragraph_format.space_after = Pt(12)

    heading_sizes = {1: 18, 2: 15, 3: 12.5, 4: 11.5, 5: 10.5, 6: 10.5}
    for level, size in heading_sizes.items():
        style = doc.styles[f"Heading {level}"]
        set_style_font(style, "Times New Roman", size, level <= 3)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(10 if level <= 2 else 7)
        style.paragraph_format.space_after = Pt(4)

    for style_name in ("List Bullet", "List Number"):
        set_style_font(doc.styles[style_name], "Times New Roman", 10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))
    set_run_font(run, "Times New Roman", 9)


def normalized_fragment(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", value.replace("\xa0", " "))


def bookmark_name(anchor: str) -> str:
    """Return a stable Word-safe bookmark name for an HTML fragment."""
    return f"nri_{hashlib.sha1(anchor.encode('utf-8')).hexdigest()[:32]}"


def add_bookmark(paragraph: Any, anchor: str, state: BuildState) -> None:
    bookmark_id = str(state.next_bookmark_id)
    state.next_bookmark_id += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), bookmark_name(anchor))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def wrap_runs_in_internal_hyperlink(paragraph: Any, runs: list[Any], anchor: str) -> None:
    if not runs:
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name(anchor))
    hyperlink.set(qn("w:history"), "1")
    first_index = paragraph._p.index(runs[0])
    paragraph._p.insert(first_index, hyperlink)
    for run in runs:
        hyperlink.append(run)


def add_text_run(paragraph: Any, value: str | None, *, bold: bool, italic: bool, code: bool, link: bool) -> None:
    text = normalized_fragment(value)
    if not text:
        return
    if not paragraph.text:
        text = text.lstrip()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if link:
        run.underline = True
        run.font.color.rgb = RGBColor(5, 99, 193)
    set_run_font(run, "Consolas" if code else "Times New Roman", 9.5 if code else None)


def add_inline(paragraph: Any, element: Any, *, bold: bool = False, italic: bool = False, code: bool = False, link: bool = False) -> None:
    add_text_run(paragraph, element.text, bold=bold, italic=italic, code=code, link=link)
    for child in element:
        tag = str(child.tag).lower()
        if tag == "br":
            paragraph.add_run().add_break()
        elif tag == "a" and str(child.get("href") or "").startswith("#"):
            before = list(paragraph._p)
            add_inline(
                paragraph,
                child,
                bold=bold,
                italic=italic,
                code=code,
                link=True,
            )
            inserted = [node for node in list(paragraph._p) if node not in before and node.tag == qn("w:r")]
            wrap_runs_in_internal_hyperlink(paragraph, inserted, str(child.get("href"))[1:])
        elif tag not in BLOCK_TAGS and tag != "img":
            add_inline(
                paragraph,
                child,
                bold=bold or tag in {"b", "strong"},
                italic=italic or tag in {"em", "i"},
                code=code or tag == "code",
                link=link or tag == "a",
            )
        add_text_run(paragraph, child.tail, bold=bold, italic=italic, code=code, link=link)


def add_paragraph(doc: Document, element: Any, state: BuildState, *, style: str | None = None, left_indent_mm: float = 0) -> Any:
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, element)
    paragraph.paragraph_format.left_indent = Mm(left_indent_mm) if left_indent_mm else None
    paragraph.paragraph_format.widow_control = True
    state.paragraphs += 1
    return paragraph


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table: Any, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), color)
        borders.append(item)


def suppress_table_borders(table: Any) -> None:
    """Make a table usable as a deterministic two-column layout container."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for child in list(borders):
        borders.remove(child)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)


def set_cell_margins(cell: Any, *, left_twips: int = 80, right_twips: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("left", left_twips), ("right", right_twips), ("top", 0), ("bottom", 0)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, element: Any, state: BuildState) -> None:
    rows = element.xpath("./thead/tr|./tbody/tr|./tr")
    if not rows:
        return
    width = max(len(row.xpath("./th|./td")) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        cells = row.xpath("./th|./td")
        for col_index, source_cell in enumerate(cells):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, source_cell)
            for run in paragraph.runs:
                set_run_font(run, "Times New Roman", 9.5)
                if row_index == 0 or str(source_cell.tag).lower() == "th":
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, "5B6573")
            elif row_index % 2:
                set_cell_shading(cell, "F3F5F7")
    state.tables += 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def image_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def add_picture_compatible(run: Any, image_path: Path, *, width: Any, height: Any) -> None:
    try:
        run.add_picture(str(image_path), width=width, height=height)
        return
    except Exception:
        pass
    with Image.open(image_path) as source:
        converted = source.convert("RGB") if source.mode not in {"RGB", "RGBA"} else source.copy()
        stream = BytesIO()
        converted.save(stream, format="PNG")
        stream.seek(0)
        run.add_picture(stream, width=width, height=height)


def figure_width_percent(classes: set[str]) -> int:
    for value in classes:
        match = re.fullmatch(r"artview-width-(\d+)", value)
        if match:
            return max(10, min(100, int(match.group(1))))
    if "art-body-right" in classes:
        return 25
    if "art-handout-page" in classes:
        return 80
    return 100


def add_figure(doc: Document, element: Any, html_path: Path, state: BuildState) -> None:
    image_nodes = element.xpath(".//img[@src]")
    if not image_nodes:
        return
    image_node = image_nodes[0]
    src = image_node.get("src")
    image_path = (html_path.parent / src).resolve()
    if not image_path.is_file():
        state.missing_images.append(src)
        return

    classes = set((element.get("class") or "").split())
    width_percent = figure_width_percent(classes)
    max_width_mm = 176.0 * width_percent / 100.0
    max_height_mm = 248.0 if ({"art-pregen-fullpage", "art-cover", "art-full-page"} & classes) else 220.0
    px_width, px_height = image_dimensions(image_path)
    width_mm = max_width_mm
    height_mm = width_mm * px_height / max(px_width, 1)
    if height_mm > max_height_mm:
        height_mm = max_height_mm
        width_mm = height_mm * px_width / max(px_height, 1)

    paragraph = doc.add_paragraph()
    if "artview-align-right" in classes or "art-body-right" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif "artview-align-left" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run()
    try:
        add_picture_compatible(run, image_path, width=Mm(width_mm), height=Mm(height_mm))
    except Exception as exc:
        raise ValueError(f"DOCX cannot embed image {src}: {type(exc).__name__}: {exc}") from exc
    state.images += 1
    state.paragraphs += 1


def add_cover(doc: Document, section: Any, html_path: Path, state: BuildState) -> None:
    style = section.get("style") or ""
    match = re.search(r"background-image:\s*url\(([^)]+)\)", style)
    if not match:
        return
    src = match.group(1).strip(" '\"")
    image_path = (html_path.parent / src).resolve()
    if not image_path.is_file():
        state.missing_images.append(src)
        return
    px_width, px_height = image_dimensions(image_path)
    max_width_mm, max_height_mm = 176.0, 265.0
    width_mm = max_width_mm
    height_mm = width_mm * px_height / max(px_width, 1)
    if height_mm > max_height_mm:
        height_mm = max_height_mm
        width_mm = height_mm * px_width / max(px_height, 1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    try:
        add_picture_compatible(paragraph.add_run(), image_path, width=Mm(width_mm), height=Mm(height_mm))
    except Exception as exc:
        raise ValueError(f"DOCX cannot embed cover {src}: {type(exc).__name__}: {exc}") from exc
    state.images += 1
    state.paragraphs += 1
    doc.add_page_break()


def add_list(doc: Document, element: Any, state: BuildState, *, ordered: bool, level: int = 0) -> None:
    style = "List Number" if ordered else "List Bullet"
    for item in element.xpath("./li"):
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Mm(6 + level * 6)
        paragraph.paragraph_format.first_line_indent = Mm(-3)
        add_inline(paragraph, item)
        state.paragraphs += 1
        for nested in item.xpath("./ul|./ol"):
            add_list(doc, nested, state, ordered=str(nested.tag).lower() == "ol", level=level + 1)


def toc_level(item: Any, fallback: int) -> int:
    for class_name in (item.get("class") or "").split():
        match = re.fullmatch(r"toc-level-(\d+)", class_name)
        if match:
            return max(1, min(6, int(match.group(1))))
    return fallback


def collect_toc_entries(element: Any, *, fallback_level: int = 1) -> list[tuple[int, str, str]]:
    entries: list[tuple[int, str, str]] = []
    for item in element.xpath("./li"):
        level = toc_level(item, fallback_level)
        anchors = item.xpath("./a[1]")
        source = anchors[0] if anchors else item
        text = normalized_fragment(source.text_content()).strip()
        if text:
            href = str(source.get("href") or "") if anchors else ""
            entries.append((level, text, href[1:] if href.startswith("#") else ""))
        for nested in item.xpath("./ol|./ul"):
            entries.extend(collect_toc_entries(nested, fallback_level=min(6, level + 1)))
    return entries


def add_toc_entry(cell: Any, *, ordinal: int, level: int, text: str, anchor: str, first: bool) -> None:
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.style = "Normal"
    level_offset_mm = max(0, level - 1) * 5.0
    text_start_mm = level_offset_mm + 9.0
    paragraph.paragraph_format.left_indent = Mm(text_start_mm)
    paragraph.paragraph_format.first_line_indent = Mm(-9.0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(text_start_mm), WD_TAB_ALIGNMENT.LEFT)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True

    primary = level == 1
    number_run = paragraph.add_run(f"{ordinal}.")
    number_run.bold = primary
    set_run_font(number_run, "Times New Roman", 9.5)
    number_run.add_tab()
    title_run = paragraph.add_run(text.upper() if primary else text)
    title_run.bold = primary
    set_run_font(title_run, "Times New Roman", 9.5 if level <= 2 else 9.0)
    if anchor:
        wrap_runs_in_internal_hyperlink(paragraph, [number_run._r, title_run._r], anchor)


def add_toc(doc: Document, element: Any, state: BuildState) -> None:
    entries = collect_toc_entries(element)
    if not entries:
        return

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    suppress_table_borders(table)
    split_at = (len(entries) + 1) // 2
    columns = (entries[:split_at], entries[split_at:])
    for column_index, (cell, column) in enumerate(zip(table.rows[0].cells, columns)):
        cell.width = Mm(86)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell)
        for position, (level, text, anchor) in enumerate(column):
            add_toc_entry(
                cell,
                ordinal=(1 if column_index == 0 else split_at + 1) + position,
                level=level,
                text=text,
                anchor=anchor,
                first=position == 0,
            )
    state.paragraphs += len(entries)
    state.toc_entries += len(entries)


def add_heading(doc: Document, element: Any, state: BuildState) -> None:
    level = max(1, min(6, int(str(element.tag)[1:])))
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if state.force_next_heading_page:
        paragraph.paragraph_format.page_break_before = True
        state.force_next_heading_page = False
    add_inline(paragraph, element)
    anchor = str(element.get("id") or "")
    if anchor:
        add_bookmark(paragraph, anchor, state)
    state.paragraphs += 1


def style_callout_paragraph(paragraph: Any, *, keep_with_next: bool) -> None:
    paragraph.paragraph_format.keep_with_next = keep_with_next
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F0DFBD")
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "8F3D19")


def add_blockquote(doc: Document, element: Any, state: BuildState) -> None:
    start = len(doc.paragraphs)
    children = [child for child in element if str(child.tag).lower() in BLOCK_TAGS]
    for index, child in enumerate(children):
        if str(child.tag).lower() == "p":
            paragraph = add_paragraph(doc, child, state, left_indent_mm=7)
            paragraph.paragraph_format.right_indent = Mm(4)
            if index == 0 and len(paragraph.text) < 120:
                for run in paragraph.runs:
                    run.bold = True
        else:
            add_block(doc, child, html_path=None, state=state)
    if "review-callout" in set((element.get("class") or "").split()):
        paragraphs = doc.paragraphs[start:]
        for index, paragraph in enumerate(paragraphs):
            style_callout_paragraph(paragraph, keep_with_next=index < len(paragraphs) - 1)
        state.callouts += 1


def add_nav(doc: Document, element: Any, state: BuildState) -> None:
    if element.get("id") == "toc":
        headings = element.xpath(".//h1[1]")
        if headings:
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, headings[0])
            state.paragraphs += 1
        lists = element.xpath(".//ol[1]|.//ul[1]")
        if lists:
            add_toc(doc, lists[0], state)
        doc.add_page_break()
        state.force_next_heading_page = False
        return

    for child in element:
        tag = str(child.tag).lower()
        if tag == "h1":
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, child)
            state.paragraphs += 1
        elif tag in {"ol", "ul"}:
            add_list(doc, child, state, ordered=tag == "ol")
        elif tag in BLOCK_TAGS:
            add_block(doc, child, html_path=None, state=state)
    doc.add_page_break()
    state.force_next_heading_page = False


def add_block(doc: Document, element: Any, html_path: Path | None, state: BuildState) -> None:
    tag = str(element.tag).lower()
    classes = set((element.get("class") or "").split())
    if tag == "section" and "production-landing" in classes:
        if html_path is not None:
            add_cover(doc, element, html_path, state)
        return
    if tag == "nav":
        add_nav(doc, element, state)
        return
    if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        add_heading(doc, element, state)
        return
    if tag == "p":
        add_paragraph(doc, element, state)
        return
    if tag == "figure":
        if html_path is not None:
            add_figure(doc, element, html_path, state)
        return
    if tag == "table":
        add_table(doc, element, state)
        return
    if tag in {"ul", "ol"}:
        add_list(doc, element, state, ordered=tag == "ol")
        return
    if tag == "blockquote":
        add_blockquote(doc, element, state)
        return
    if tag in {"section", "div"}:
        if "production-section-start" in classes:
            state.force_next_heading_page = True
        for child in element:
            if str(child.tag).lower() in BLOCK_TAGS:
                add_block(doc, child, html_path, state)


def _xml_text(element: Any) -> str:
    return normalized_fragment("".join(element.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))).strip()


def _paragraph_style_name(paragraph: Any) -> str:
    values = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=WORD_NAMESPACES)
    return str(values[0]) if values else ""


def _html_target_block(root: Any, anchor: str) -> Any:
    matches = root.xpath("//*[@id=$anchor]", anchor=anchor)
    if len(matches) != 1:
        raise ValueError(f"internal target must resolve exactly once: {anchor!r}")
    target = matches[0]
    block_tags = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li"}
    if str(target.tag).lower() in block_tags and normalized_fragment(target.text_content()).strip():
        return target
    parent = target.getparent()
    if parent is not None and str(parent.tag).lower() in block_tags and normalized_fragment(parent.text_content()).strip():
        return parent
    following = target.xpath(
        "following::*[self::h1 or self::h2 or self::h3 or self::h4 or self::h5 or self::h6 or self::p or self::li]"
    )
    for candidate in following:
        if normalized_fragment(candidate.text_content()).strip():
            return candidate
    raise ValueError(f"internal target has no rendered block: {anchor!r}")


def _html_internal_links(root: Any) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    toc_nodes = root.xpath("//nav[@id='toc']//a[starts-with(@href, '#')]")
    toc_identity = {id(node) for node in toc_nodes}
    body_nodes = [
        node
        for node in root.xpath("//main//a[starts-with(@href, '#')]")
        if id(node) not in toc_identity
    ]

    def rows(nodes: list[Any]) -> list[tuple[str, str]]:
        return [
            (str(node.get("href"))[1:], normalized_fragment(node.text_content()).strip())
            for node in nodes
        ]

    return rows(toc_nodes), rows(body_nodes)


def _find_target_paragraphs(root_html: Any, root_xml: Any, anchors: set[str]) -> dict[str, Any]:
    tree = root_html.getroottree()
    grouped: dict[str, dict[str, Any]] = {}
    for anchor in anchors:
        target = _html_target_block(root_html, anchor)
        target_path = tree.getpath(target)
        row = grouped.setdefault(
            target_path,
            {
                "target": target,
                "anchors": [],
                "order": int(target.xpath("count(preceding::*)")),
                "text": normalized_fragment(target.text_content()).strip(),
            },
        )
        row["anchors"].append(anchor)

    paragraphs = root_xml.xpath(".//w:body//w:p", namespaces=WORD_NAMESPACES)
    paragraph_texts = [_xml_text(paragraph) for paragraph in paragraphs]
    cursor = 0
    result: dict[str, Any] = {}
    for row in sorted(grouped.values(), key=lambda item: int(item["order"])):
        target = row["target"]
        target_text = str(row["text"])
        found = None
        for index in range(cursor, len(paragraphs)):
            if paragraph_texts[index].casefold() != target_text.casefold():
                continue
            found = index
            break
        if found is None:
            raise ValueError(
                f"DOCX target paragraph not found after position {cursor}: {target_text!r} ({sorted(row['anchors'])})"
            )
        for anchor in row["anchors"]:
            result[str(anchor)] = paragraphs[found]
        cursor = found + 1
    return result


def _is_visual_link_run(run: Any) -> bool:
    colors = run.xpath("./w:rPr/w:color/@w:val", namespaces=WORD_NAMESPACES)
    underlines = run.xpath("./w:rPr/w:u", namespaces=WORD_NAMESPACES)
    return bool(underlines and colors and str(colors[0]).upper() == "0563C1")


def _visual_link_groups(root_xml: Any) -> list[tuple[Any, list[Any], str]]:
    groups: list[tuple[Any, list[Any], str]] = []
    for paragraph in root_xml.xpath(".//w:body//w:p", namespaces=WORD_NAMESPACES):
        current: list[Any] = []
        for child in list(paragraph):
            if child.tag == qn("w:r") and _is_visual_link_run(child):
                current.append(child)
            elif child.tag == qn("w:r") and current:
                groups.append((
                    paragraph,
                    current,
                    normalized_fragment("".join(
                        "".join(run.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))
                        for run in current
                    )).strip(),
                ))
                current = []
        if current:
            groups.append((
                paragraph,
                current,
                normalized_fragment("".join(
                    "".join(run.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))
                    for run in current
                )).strip(),
            ))
    return groups


def _wrap_xml_runs(paragraph: Any, runs: list[Any], anchor: str) -> None:
    hyperlink = lxml_etree.Element(qn("w:hyperlink"))
    hyperlink.set(qn("w:anchor"), bookmark_name(anchor))
    hyperlink.set(qn("w:history"), "1")
    first_index = paragraph.index(runs[0])
    paragraph.insert(first_index, hyperlink)
    for run in runs:
        hyperlink.append(run)


def _add_xml_bookmarks(root_xml: Any, targets: dict[str, Any]) -> None:
    existing_ids = [
        int(value)
        for value in root_xml.xpath(".//w:bookmarkStart/@w:id", namespaces=WORD_NAMESPACES)
        if str(value).isdigit()
    ]
    next_id = max(existing_ids, default=0) + 1
    by_paragraph: dict[int, tuple[Any, list[str]]] = {}
    for anchor, paragraph in targets.items():
        row = by_paragraph.setdefault(id(paragraph), (paragraph, []))
        row[1].append(anchor)
    for paragraph, anchors in by_paragraph.values():
        insert_at = 1 if paragraph.find(qn("w:pPr")) is not None else 0
        ends: list[Any] = []
        for anchor in sorted(anchors):
            start = lxml_etree.Element(qn("w:bookmarkStart"))
            start.set(qn("w:id"), str(next_id))
            start.set(qn("w:name"), bookmark_name(anchor))
            end = lxml_etree.Element(qn("w:bookmarkEnd"))
            end.set(qn("w:id"), str(next_id))
            paragraph.insert(insert_at, start)
            insert_at += 1
            ends.append(end)
            next_id += 1
        for end in reversed(ends):
            paragraph.append(end)


def _style_xml_callout_paragraph(paragraph: Any, *, keep_with_next: bool) -> None:
    p_pr = paragraph.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = lxml_etree.Element(qn("w:pPr"))
        paragraph.insert(0, p_pr)
    keep = p_pr.find(qn("w:keepNext"))
    if keep_with_next and keep is None:
        keep = lxml_etree.Element(qn("w:keepNext"))
        p_pr.append(keep)
    elif not keep_with_next and keep is not None:
        p_pr.remove(keep)
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = lxml_etree.Element(qn("w:shd"))
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F0DFBD")
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = lxml_etree.Element(qn("w:pBdr"))
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = lxml_etree.Element(qn("w:left"))
        borders.append(left)
    for key, value in {"val": "single", "sz": "24", "space": "6", "color": "8F3D19"}.items():
        left.set(qn(f"w:{key}"), value)


def _apply_xml_callouts(root_html: Any, root_xml: Any) -> tuple[int, list[str]]:
    body_paragraphs = root_xml.xpath("./w:body/w:p", namespaces=WORD_NAMESPACES)
    body_texts = [_xml_text(paragraph) for paragraph in body_paragraphs]
    cursor = 0
    applied = 0
    missing: list[str] = []
    for blockquote in root_html.xpath("//blockquote[contains(concat(' ', normalize-space(@class), ' '), ' review-callout ')]"):
        chunks = [
            normalized_fragment(child.text_content()).strip()
            for child in blockquote.xpath("./p|./ol/li|./ul/li")
            if normalized_fragment(child.text_content()).strip()
        ]
        expected = normalized_fragment(" ".join(chunks)).strip()
        matched: list[int] | None = None
        for start in range(cursor, len(body_paragraphs)):
            if not body_texts[start]:
                continue
            combined = ""
            indices: list[int] = []
            for end in range(start, len(body_paragraphs)):
                if not body_texts[end]:
                    continue
                combined = normalized_fragment(f"{combined} {body_texts[end]}").strip()
                indices.append(end)
                if combined == expected:
                    matched = indices
                    break
                if not expected.startswith(combined):
                    break
            if matched is not None:
                break
        if matched is None:
            missing.append(expected[:160])
            continue
        for index, paragraph_index in enumerate(matched):
            _style_xml_callout_paragraph(
                body_paragraphs[paragraph_index], keep_with_next=index < len(matched) - 1
            )
        cursor = matched[-1] + 1
        applied += 1
    return applied, missing


def patch_existing_docx(html_path: Path, docx_path: Path, output_path: Path | None = None) -> dict[str, Any]:
    """Patch navigation and callout presentation into an existing edited DOCX.

    Only ``word/document.xml`` is replaced, so user formatting, media, headers,
    footers, sections and all other package members remain byte-for-byte intact.
    """
    html_path = html_path.resolve()
    docx_path = docx_path.resolve()
    output_path = (output_path or docx_path).resolve()
    root_html = lxml_html.parse(str(html_path)).getroot()
    toc_links, body_links = _html_internal_links(root_html)
    expected_links = len(toc_links) + len(body_links)

    with zipfile.ZipFile(docx_path) as source:
        if source.testzip() is not None:
            raise ValueError("source DOCX archive is corrupt")
        root_xml = lxml_etree.fromstring(source.read("word/document.xml"))
        existing_links = root_xml.xpath(".//w:hyperlink", namespaces=WORD_NAMESPACES)
        expected_anchor_sequence = [bookmark_name(anchor) for anchor, _ in toc_links + body_links]
        if existing_links:
            if len(existing_links) != expected_links:
                raise ValueError(
                    f"existing DOCX has a partial hyperlink set: {len(existing_links)} of {expected_links}"
                )
            actual_anchor_sequence = [str(link.get(qn("w:anchor")) or "") for link in existing_links]
            if actual_anchor_sequence != expected_anchor_sequence:
                raise ValueError("existing DOCX hyperlink targets differ from the HTML navigation sequence")
        else:
            toc_table = root_xml.xpath(".//w:body/w:tbl[1]", namespaces=WORD_NAMESPACES)
            if len(toc_table) != 1:
                raise ValueError("DOCX must contain one leading two-column TOC table")
            toc_paragraphs = [
                paragraph
                for paragraph in toc_table[0].xpath("./w:tr/w:tc/w:p", namespaces=WORD_NAMESPACES)
                if _xml_text(paragraph)
            ]
            if len(toc_paragraphs) != len(toc_links):
                raise ValueError(f"DOCX TOC has {len(toc_paragraphs)} entries, expected {len(toc_links)}")
            for paragraph, (anchor, _text) in zip(toc_paragraphs, toc_links):
                runs = paragraph.xpath("./w:r", namespaces=WORD_NAMESPACES)
                _wrap_xml_runs(paragraph, runs, anchor)

            visual_groups = _visual_link_groups(root_xml)
            if len(visual_groups) != len(body_links):
                raise ValueError(
                    f"DOCX has {len(visual_groups)} visual body-link groups, expected {len(body_links)}"
                )
            for (paragraph, runs, actual_text), (anchor, expected_text) in zip(visual_groups, body_links):
                if actual_text.casefold() != expected_text.casefold():
                    raise ValueError(
                        f"body-link order mismatch: DOCX {actual_text!r}, HTML {expected_text!r}"
                    )
                _wrap_xml_runs(paragraph, runs, anchor)

        target_anchors = {anchor for anchor, _ in toc_links + body_links}
        existing_bookmarks = set(
            root_xml.xpath(".//w:bookmarkStart/@w:name", namespaces=WORD_NAMESPACES)
        )
        missing_target_anchors = {
            anchor for anchor in target_anchors if bookmark_name(anchor) not in existing_bookmarks
        }
        if missing_target_anchors:
            targets = _find_target_paragraphs(root_html, root_xml, missing_target_anchors)
            _add_xml_bookmarks(root_xml, targets)

        callouts_applied, callouts_missing = _apply_xml_callouts(root_html, root_xml)
        document_xml = lxml_etree.tostring(
            root_xml, xml_declaration=True, encoding="UTF-8", standalone="yes"
        )
        temp_path = output_path.with_name(f".{output_path.name}.navigation.tmp")
        with zipfile.ZipFile(temp_path, "w") as destination:
            for info in source.infolist():
                payload = document_xml if info.filename == "word/document.xml" else source.read(info.filename)
                destination.writestr(info, payload)
    temp_path.replace(output_path)
    return {
        "status": "pass",
        "output": output_path.as_posix(),
        "sha256": sha256(output_path),
        "bytes": output_path.stat().st_size,
        "toc_hyperlinks": len(toc_links),
        "body_hyperlinks": len(body_links),
        "bookmarks": len({anchor for anchor, _ in toc_links + body_links}),
        "callouts_applied": callouts_applied,
        "callouts_missing": callouts_missing,
    }


def build_docx(html_path: Path, output_path: Path) -> dict[str, Any]:
    html_path = html_path.resolve()
    output_path = output_path.resolve()
    if not html_path.is_file():
        raise FileNotFoundError(html_path)

    root = lxml_html.parse(str(html_path)).getroot()
    main_nodes = root.xpath("//main")
    if len(main_nodes) != 1:
        raise ValueError("production HTML must contain exactly one main element")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Вирус"
    doc.core_properties.subject = "Сценарий Call of Cthulhu"
    doc.core_properties.author = "NRI"
    state = BuildState()
    for element in main_nodes[0]:
        if str(element.tag).lower() in BLOCK_TAGS:
            add_block(doc, element, html_path, state)
    if state.missing_images:
        raise FileNotFoundError(f"DOCX source images are missing: {sorted(state.missing_images)}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f".{output_path.name}.tmp")
    doc.save(temp_path)
    semantic_result = patch_existing_docx(html_path, temp_path)
    if semantic_result["callouts_missing"]:
        raise ValueError(
            f"generated DOCX lost review callouts: {semantic_result['callouts_missing']}"
        )
    with zipfile.ZipFile(temp_path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names or "[Content_Types].xml" not in names:
            raise ValueError("generated DOCX is missing required OOXML members")
    temp_path.replace(output_path)
    return {
        "status": "pass",
        "output": output_path.as_posix(),
        "sha256": sha256(output_path),
        "bytes": output_path.stat().st_size,
        "paragraphs": state.paragraphs,
        "tables": state.tables,
        "images": state.images,
        "toc_entries": state.toc_entries,
        "callouts": state.callouts,
        "internal_hyperlinks": semantic_result["toc_hyperlinks"] + semantic_result["body_hyperlinks"],
        "bookmarks": semantic_result["bookmarks"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("html", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--patch-existing", action="store_true")
    args = parser.parse_args()
    try:
        result = (
            patch_existing_docx(args.html, args.output)
            if args.patch_existing
            else build_docx(args.html, args.output)
        )
        print(json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2))
        return 0
    except Exception as exc:
        print(json.dumps({"status": "fail", "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False, indent=2))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
