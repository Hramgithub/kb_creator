#!/usr/bin/env python3
"""Create an editable DOCX projection from a production HTML artifact."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from lxml import etree as lxml_etree
from lxml import html as lxml_html
from PIL import Image


BLOCK_TAGS = {"blockquote", "div", "figure", "h1", "h2", "h3", "h4", "h5", "h6", "nav", "ol", "p", "section", "table", "ul"}
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_NAMESPACES = {"w": WORD_NS}
ROLE_STYLE_NAMES = {
    "table_body": "NRI Table Body",
    "table_header": "NRI Table Header",
    "caption": "NRI Caption",
    "statblock": "NRI Statblock",
    "toc": "NRI TOC",
    "quote": "NRI Quote",
}


@dataclass
class BuildState:
    force_next_heading_page: bool = False
    paragraphs: int = 0
    tables: int = 0
    images: int = 0
    toc_entries: int = 0
    callouts: int = 0
    callout_presentation: str = "legacy_shaded"
    next_bookmark_id: int = 1
    missing_images: list[str] | None = None
    figure_policy: dict[str, Any] | None = None
    statblocks_keep_together: bool = False
    page_usable_width_mm: float | None = None
    page_usable_height_mm: float | None = None
    typography: dict[str, Any] | None = None
    typography_configured: bool = False
    bookmarked_anchors: set[str] | None = None
    statblock_depth: int = 0
    role_styles_initialized: bool = False
    ordered_list_abstract_num_id: int | None = None
    reader_statblock_groups: dict[str, list[Any]] | None = None

    def __post_init__(self) -> None:
        if self.missing_images is None:
            self.missing_images = []
        if self.bookmarked_anchors is None:
            self.bookmarked_anchors = set()
        if self.reader_statblock_groups is None:
            self.reader_statblock_groups = {}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def set_run_font(run: Any, name: str, size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{key}"), None)


def set_style_font(style: Any, name: str, size: float, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)
    # Template theme references override explicit fonts in Word/LibreOffice.
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{key}"), None)
    cs_size = r_pr.find(qn("w:szCs"))
    if cs_size is None:
        cs_size = OxmlElement("w:szCs")
        r_pr.append(cs_size)
    cs_size.set(qn("w:val"), str(round(size * 2)))


def remove_style_paragraph_borders(style: Any) -> None:
    paragraph_properties = style.element.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is not None:
        paragraph_properties.remove(borders)


def _closed(value: object, keys: set[str], label: str) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != keys:
        raise ValueError(f"{label} has unknown or missing fields")
    return value


def _figure_policy(value: object, label: str) -> dict[str, Any]:
    required = {"max_width_mm", "max_height_mm", "keep_with_next"}
    optional = {"crop_fraction"}
    if not isinstance(value, dict) or not required.issubset(value) or set(value) - required - optional:
        raise ValueError(f"{label} has unknown or missing fields")
    policy = value
    dimensions = (policy["max_width_mm"], policy["max_height_mm"])
    if (
        any(isinstance(item, bool) or not isinstance(item, (int, float)) or not math.isfinite(item) or item <= 0 for item in dimensions)
        or not isinstance(policy["keep_with_next"], bool)
    ):
        raise ValueError(f"{label} contains invalid typed values")
    if "crop_fraction" in policy:
        crop = _closed(policy["crop_fraction"], {"left", "top", "right", "bottom"}, f"{label} crop_fraction")
        values = [crop[edge] for edge in ("left", "top", "right", "bottom")]
        if any(isinstance(item, bool) or not isinstance(item, (int, float)) or not math.isfinite(item) or not 0 <= item < 1 for item in values):
            raise ValueError(f"{label} crop_fraction contains invalid values")
        if crop["left"] + crop["right"] >= 1 or crop["top"] + crop["bottom"] >= 1:
            raise ValueError(f"{label} crop_fraction leaves no visible image area")
    return policy


def _figure_override_source(value: object) -> str:
    if not isinstance(value, str) or not value or value != value.strip() or "\\" in value:
        raise ValueError("DOCX figure override source must be a contained relative image path")
    path = PurePosixPath(value)
    if (
        path.is_absolute()
        or not path.parts
        or any(part in {"", ".", ".."} for part in path.parts)
        or ":" in value
        or "//" in value
    ):
        raise ValueError("DOCX figure override source must be a contained relative image path")
    return value


def _validate_figure_settings(value: object) -> dict[str, Any]:
    figures = _closed(value, {"default", "overrides"}, "DOCX figures")
    _figure_policy(figures["default"], "DOCX figures default")
    if not isinstance(figures["overrides"], dict):
        raise ValueError("DOCX figure overrides must be an object")
    for source, policy in figures["overrides"].items():
        _figure_override_source(source)
        _figure_policy(policy, f"DOCX figure override {source!r}")
    return figures


def _default_typography(fonts: dict[str, str], sizes: dict[str, float], spacing: dict[str, float]) -> dict[str, dict[str, Any]]:
    """Return the legacy-compatible role profile used when roles are omitted."""
    body = {
        "font": fonts["body"], "size_pt": sizes["body_pt"], "space_before_pt": 0.0,
        "space_after_pt": spacing["body_after_pt"], "line_spacing": spacing["body_line"],
    }
    return {
        "table_body": {**body, "size_pt": 9.5, "space_after_pt": 0.0, "line_spacing": 1.0},
        "table_header": {**body, "size_pt": 9.5, "space_after_pt": 0.0, "line_spacing": 1.0},
        "caption": {**body, "size_pt": 9.5, "space_after_pt": 3.0, "line_spacing": 1.0},
        "statblock": body.copy(),
        "toc": {**body, "size_pt": 9.5, "space_after_pt": 1.5, "line_spacing": 1.0},
        "quote": body.copy(),
    }


def _validate_typography(value: object) -> dict[str, dict[str, Any]]:
    roles = _closed(value, set(ROLE_STYLE_NAMES), "DOCX typography")
    fields = {"font", "size_pt", "space_before_pt", "space_after_pt", "line_spacing"}
    for role, settings in roles.items():
        _closed(settings, fields, f"DOCX typography {role}")
        if not isinstance(settings["font"], str) or not settings["font"].strip():
            raise ValueError(f"DOCX typography {role} font must be a non-empty string")
        numeric = [settings[key] for key in fields - {"font"}]
        if any(isinstance(item, bool) or not isinstance(item, (int, float)) or not math.isfinite(item) or item < 0 for item in numeric):
            raise ValueError(f"DOCX typography {role} contains invalid numeric values")
        if settings["size_pt"] <= 0 or settings["line_spacing"] <= 0:
            raise ValueError(f"DOCX typography {role} size and line spacing must be positive")
    return roles


def validate_document_config(value: object) -> dict[str, Any]:
    required = {"schema", "page", "margins", "fonts", "sizes", "spacing", "headings", "footer", "metadata"}
    optional = {"figures", "statblocks", "typography", "callouts"}
    if not isinstance(value, dict) or not required.issubset(value) or set(value) - required - optional:
        raise ValueError("DOCX config has unknown or missing fields")
    config = value
    if config["schema"] != "kb_creator.production_docx_config.v1":
        raise ValueError("unsupported DOCX config schema")
    page = _closed(config["page"], {"width_mm", "height_mm", "orientation"}, "DOCX page")
    margins = _closed(config["margins"], {"top_mm", "bottom_mm", "left_mm", "right_mm"}, "DOCX margins")
    fonts = _closed(config["fonts"], {"body", "title", "headings", "footer"}, "DOCX fonts")
    sizes = _closed(config["sizes"], {"body_pt", "title_pt", "footer_pt"}, "DOCX sizes")
    spacing = _closed(config["spacing"], {"body_after_pt", "body_line"}, "DOCX spacing")
    footer = _closed(config["footer"], {"page_numbers"}, "DOCX footer")
    metadata = _closed(config["metadata"], {"title", "subject", "author", "keywords", "comments"}, "DOCX metadata")
    headings = config["headings"]
    if not isinstance(headings, dict) or set(headings) != {str(level) for level in range(1, 7)}:
        raise ValueError("DOCX headings must define exact levels 1-6")
    for level in range(1, 7):
        _closed(headings[str(level)], {"size_pt", "bold", "space_before_pt", "space_after_pt"}, f"DOCX heading {level}")
    numeric = [page["width_mm"], page["height_mm"], *margins.values(), *sizes.values(), spacing["body_after_pt"], spacing["body_line"]]
    numeric.extend(item[key] for item in headings.values() for key in ("size_pt", "space_before_pt", "space_after_pt"))
    if any(isinstance(item, bool) or not isinstance(item, (int, float)) or item <= 0 for item in numeric):
        raise ValueError("DOCX numeric layout values must be positive")
    if page["orientation"] not in {"portrait", "landscape"} or not all(isinstance(value, str) and value.strip() for value in fonts.values()) or not isinstance(footer["page_numbers"], bool) or any(not isinstance(value, str) for value in metadata.values()) or any(not isinstance(item["bold"], bool) for item in headings.values()):
        raise ValueError("DOCX config contains invalid typed values")
    if "figures" in config:
        _validate_figure_settings(config["figures"])
    if "statblocks" in config:
        statblocks = _closed(config["statblocks"], {"keep_together"}, "DOCX statblocks")
        if not isinstance(statblocks["keep_together"], bool):
            raise ValueError("DOCX statblocks contains invalid typed values")
    if "typography" in config:
        _validate_typography(config["typography"])
    if "callouts" in config:
        callouts = _closed(config["callouts"], {"presentation"}, "DOCX callouts")
        if callouts["presentation"] not in {"plain", "legacy_shaded"}:
            raise ValueError("DOCX callout presentation is unsupported")
    return config


def build_state_for_profile(profile: dict[str, Any] | None) -> BuildState:
    if profile is None:
        return BuildState()
    figures = profile.get("figures")
    if figures is None:
        return BuildState(
            statblocks_keep_together=bool(profile.get("statblocks", {}).get("keep_together", False)),
            typography=profile.get("typography"), typography_configured="typography" in profile,
            callout_presentation=profile.get("callouts", {}).get("presentation", "legacy_shaded"),
        )
    page = profile["page"]
    margins = profile["margins"]
    return BuildState(
        figure_policy=figures,
        statblocks_keep_together=bool(profile.get("statblocks", {}).get("keep_together", False)),
        page_usable_width_mm=page["width_mm"] - margins["left_mm"] - margins["right_mm"],
        page_usable_height_mm=page["height_mm"] - margins["top_mm"] - margins["bottom_mm"],
        typography=profile.get("typography"),
        typography_configured="typography" in profile,
        callout_presentation=profile.get("callouts", {}).get("presentation", "legacy_shaded"),
    )


def _configure_role_style(doc: Document, style_name: str, settings: dict[str, Any], *, bold: bool | None = None) -> None:
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(style, settings["font"], settings["size_pt"], bold)
    style.paragraph_format.space_before = Pt(settings["space_before_pt"])
    style.paragraph_format.space_after = Pt(settings["space_after_pt"])
    style.paragraph_format.line_spacing = settings["line_spacing"]
    style.paragraph_format.widow_control = True


def configure_typography_styles(doc: Document, fonts: dict[str, str], sizes: dict[str, float], spacing: dict[str, float], profile: dict[str, Any] | None) -> None:
    typography = profile["typography"] if profile and "typography" in profile else _default_typography(fonts, sizes, spacing)
    for role, style_name in ROLE_STYLE_NAMES.items():
        _configure_role_style(doc, style_name, typography[role], bold=True if role == "table_header" else None)


def ensure_role_styles(doc: Document, state: BuildState) -> None:
    """Support focused in-memory block tests without changing normal build setup."""
    if state.role_styles_initialized:
        return
    defaults = _default_typography(
        {"body": "Times New Roman"},
        {"body_pt": 10.5},
        {"body_after_pt": 4.0, "body_line": 1.05},
    )
    typography = state.typography or defaults
    for role, style_name in ROLE_STYLE_NAMES.items():
        try:
            doc.styles[style_name]
        except KeyError:
            _configure_role_style(doc, style_name, typography[role], bold=True if role == "table_header" else None)
    state.role_styles_initialized = True


def configure_document(doc: Document, profile: dict[str, Any] | None = None) -> None:
    if profile is not None:
        profile = validate_document_config(profile)
    section = doc.sections[0]
    page = profile["page"] if profile else {"width_mm": 210, "height_mm": 297, "orientation": "portrait"}
    margins = profile["margins"] if profile else {"top_mm": 15, "bottom_mm": 15, "left_mm": 17, "right_mm": 17}
    section.orientation = WD_ORIENT.LANDSCAPE if page["orientation"] == "landscape" else WD_ORIENT.PORTRAIT
    section.page_width = Mm(page["width_mm"])
    section.page_height = Mm(page["height_mm"])
    section.top_margin = Mm(margins["top_mm"])
    section.bottom_margin = Mm(margins["bottom_mm"])
    section.left_margin = Mm(margins["left_mm"])
    section.right_margin = Mm(margins["right_mm"])

    fonts = profile["fonts"] if profile else {"body": "Times New Roman", "title": "Times New Roman", "headings": "Times New Roman", "footer": "Times New Roman"}
    sizes = profile["sizes"] if profile else {"body_pt": 10.5, "title_pt": 22, "footer_pt": 9}
    spacing = profile["spacing"] if profile else {"body_after_pt": 4, "body_line": 1.05}

    normal = doc.styles["Normal"]
    set_style_font(normal, fonts["body"], sizes["body_pt"])
    normal.paragraph_format.space_after = Pt(spacing["body_after_pt"])
    normal.paragraph_format.line_spacing = spacing["body_line"]
    normal.paragraph_format.widow_control = True

    set_style_font(doc.styles["Title"], fonts["title"], sizes["title_pt"], True)
    remove_style_paragraph_borders(doc.styles["Title"])
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Title"].paragraph_format.space_after = Pt(12)

    heading_profiles = profile["headings"] if profile else {str(level): {"size_pt": size, "bold": level <= 3, "space_before_pt": 10 if level <= 2 else 7, "space_after_pt": 4} for level, size in {1: 18, 2: 15, 3: 12.5, 4: 11.5, 5: 10.5, 6: 10.5}.items()}
    for level in range(1, 7):
        heading = heading_profiles[str(level)]
        style = doc.styles[f"Heading {level}"]
        set_style_font(style, fonts["headings"], heading["size_pt"], heading["bold"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(heading["space_before_pt"])
        style.paragraph_format.space_after = Pt(heading["space_after_pt"])

    for style_name in ("List Bullet", "List Number"):
        set_style_font(doc.styles[style_name], fonts["body"], sizes["body_pt"])

    configure_typography_styles(doc, fonts, sizes, spacing, profile)

    if profile is None or profile["footer"]["page_numbers"]:
        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instruction, end))
        set_run_font(run, fonts["footer"], sizes["footer_pt"])


def normalized_fragment(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", value.replace("\xa0", " "))


def validate_unique_html_ids(root: Any) -> None:
    """Reject ambiguous fragment targets before generating or patching a DOCX."""
    seen: set[str] = set()
    for element in root.xpath("//*[@id]"):
        anchor = str(element.get("id") or "")
        if not anchor or anchor in seen:
            raise ValueError(f"HTML ids must be non-empty and unique: {anchor!r}")
        seen.add(anchor)


def bookmark_name(anchor: str) -> str:
    """Return a stable Word-safe bookmark name for an HTML fragment."""
    return f"nri_{hashlib.sha1(anchor.encode('utf-8')).hexdigest()[:32]}"


def add_bookmark(paragraph: Any, anchor: str, state: BuildState) -> None:
    if anchor in state.bookmarked_anchors:
        raise ValueError(f"HTML id was rendered more than once: {anchor!r}")
    state.bookmarked_anchors.add(anchor)
    bookmark_id = str(state.next_bookmark_id)
    state.next_bookmark_id += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), bookmark_name(anchor))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_element_bookmark(paragraph: Any, element: Any, state: BuildState) -> None:
    anchor = str(element.get("id") or "")
    if anchor:
        add_bookmark(paragraph, anchor, state)


def add_rendered_bookmarks(paragraph: Any, element: Any, state: BuildState) -> None:
    """Anchor a leaf element and inline descendants at its rendered paragraph."""
    add_element_bookmark(paragraph, element, state)
    for child in element.xpath(".//*[@id]"):
        add_element_bookmark(paragraph, child, state)


def add_container_bookmark(doc: Document, element: Any, state: BuildState, *, paragraph_start: int, table_start: int) -> None:
    if len(doc.paragraphs) > paragraph_start:
        add_element_bookmark(doc.paragraphs[paragraph_start], element, state)
    elif len(doc.tables) > table_start:
        add_element_bookmark(doc.tables[table_start].cell(0, 0).paragraphs[0], element, state)


def wrap_runs_in_internal_hyperlink(paragraph: Any, runs: list[Any], anchor: str) -> None:
    if not runs:
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name(anchor))
    hyperlink.set(qn("w:history"), "1")
    first_index = paragraph._p.index(runs[0])
    paragraph._p.insert(first_index, hyperlink)
    for run in runs:
        hyperlink.append(run)


def add_text_run(paragraph: Any, value: str | None, *, bold: bool, italic: bool, code: bool, link: bool) -> None:
    text = normalized_fragment(value)
    if not text:
        return
    if not paragraph.text:
        text = text.lstrip()
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if link:
        run.underline = True
        run.font.color.rgb = RGBColor(5, 99, 193)
    if code:
        set_run_font(run, "Consolas", 9.5)


def add_inline(paragraph: Any, element: Any, *, bold: bool = False, italic: bool = False, code: bool = False, link: bool = False) -> None:
    add_text_run(paragraph, element.text, bold=bold, italic=italic, code=code, link=link)
    for child in element:
        tag = str(child.tag).lower()
        if tag == "br":
            paragraph.add_run().add_break()
        elif tag == "a" and str(child.get("href") or "").startswith("#"):
            before = list(paragraph._p)
            add_inline(
                paragraph,
                child,
                bold=bold,
                italic=italic,
                code=code,
                link=True,
            )
            inserted = [node for node in list(paragraph._p) if node not in before and node.tag == qn("w:r")]
            wrap_runs_in_internal_hyperlink(paragraph, inserted, str(child.get("href"))[1:])
        elif tag not in BLOCK_TAGS and tag != "img":
            add_inline(
                paragraph,
                child,
                bold=bold or tag in {"b", "strong"},
                italic=italic or tag in {"em", "i"},
                code=code or tag == "code",
                link=link or tag == "a",
            )
        add_text_run(paragraph, child.tail, bold=bold, italic=italic, code=code, link=link)


def add_paragraph(doc: Document, element: Any, state: BuildState, *, style: str | None = None, left_indent_mm: float = 0) -> Any:
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, element)
    paragraph.paragraph_format.left_indent = Mm(left_indent_mm) if left_indent_mm else None
    paragraph.paragraph_format.widow_control = True
    if element.get("data-semantic-kind") == "label" and element.get("data-keep-with-next") == "true":
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
    add_rendered_bookmarks(paragraph, element, state)
    state.paragraphs += 1
    return paragraph


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table: Any, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), color)
        borders.append(item)


def suppress_table_borders(table: Any) -> None:
    """Make a table usable as a deterministic two-column layout container."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for child in list(borders):
        borders.remove(child)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)


def set_cell_margins(cell: Any, *, left_twips: int = 80, right_twips: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("left", left_twips), ("right", right_twips), ("top", 0), ("bottom", 0)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def leading_table_header_rows(rows: list[Any], explicit_headers: set[int] | None) -> int:
    """Count only the contiguous leading header rows eligible for pagination binding."""
    count = 0
    for row_index, row in enumerate(rows):
        if explicit_headers is not None:
            is_header = row.get("data-table-header") == "true"
        else:
            cells = row.xpath("./th|./td")
            is_header = row_index == 0 or bool(cells) and all(
                str(cell.tag).lower() == "th" for cell in cells
            )
        if not is_header:
            break
        count += 1
    return count


def add_table(doc: Document, element: Any, state: BuildState) -> None:
    rows = element.xpath("./thead/tr|./tbody/tr|./tr")
    if not rows:
        return
    explicit_policy = element.get("data-table-header-policy")
    if explicit_policy not in {None, "explicit"}:
        raise ValueError("unsupported table header policy")
    explicit_headers: set[int] | None = None
    if explicit_policy == "explicit":
        raw_headers = element.get("data-table-source-header-rows")
        if raw_headers is None:
            raise ValueError("explicit table header policy has no source header rows")
        values = [] if raw_headers == "" else raw_headers.split(",")
        if any(not value.isdigit() or int(value) <= 0 for value in values):
            raise ValueError("explicit table header rows are invalid")
        explicit_headers = {int(value) for value in values}
        if len(explicit_headers) != len(values):
            raise ValueError("explicit table header rows are duplicated")
        source_rows = []
        for row in rows:
            source_row = row.get("data-table-source-row")
            marker = row.get("data-table-header")
            if (source_row is None or not source_row.isdigit() or int(source_row) <= 0
                    or marker not in {"true", "false"}):
                raise ValueError("explicit table row marker is invalid")
            source_rows.append(int(source_row))
            if (int(source_row) in explicit_headers) != (marker == "true"):
                raise ValueError("explicit table header marker disagrees with policy")
        if source_rows != sorted(set(source_rows)) or not explicit_headers.issubset(source_rows):
            raise ValueError("explicit table source rows are invalid")
        if len(explicit_headers) == len(rows):
            raise ValueError("explicit table has no data row")
    leading_headers = leading_table_header_rows(rows, explicit_headers)
    bind_leading_headers = 0 < leading_headers < len(rows)
    width = max(len(row.xpath("./th|./td")) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        cells = row.xpath("./th|./td")
        for col_index, source_cell in enumerate(cells):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            header = (row.get("data-table-header") == "true"
                      if explicit_headers is not None
                      else row_index == 0 or str(source_cell.tag).lower() == "th")
            paragraph.style = ROLE_STYLE_NAMES["table_header" if header else "table_body"]
            add_inline(paragraph, source_cell)
            add_rendered_bookmarks(paragraph, source_cell, state)
            for run in paragraph.runs:
                if header:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if header:
                set_cell_shading(cell, "5B6573")
            elif row_index % 2:
                set_cell_shading(cell, "F3F5F7")
            if bind_leading_headers and row_index < leading_headers:
                paragraph.paragraph_format.keep_with_next = True
    state.tables += 1
    spacer = doc.add_paragraph(style=ROLE_STYLE_NAMES["table_body"])
    spacer.paragraph_format.space_after = Pt(2)
    add_element_bookmark(table.cell(0, 0).paragraphs[0], element, state)


def image_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def add_picture_compatible(run: Any, image_path: Path, *, width: Any, height: Any, preserve_bytes: bool = False) -> Any:
    try:
        return run.add_picture(str(image_path), width=width, height=height)
    except Exception:
        if preserve_bytes:
            raise
    with Image.open(image_path) as source:
        converted = source.convert("RGB") if source.mode not in {"RGB", "RGBA"} else source.copy()
        stream = BytesIO()
        converted.save(stream, format="PNG")
        stream.seek(0)
        return run.add_picture(stream, width=width, height=height)


def apply_picture_crop(inline_shape: Any, crop_fraction: dict[str, Any] | None) -> None:
    """Apply DrawingML source cropping without changing the embedded image bytes."""
    if crop_fraction is None:
        return
    blip_fill = next(
        (node for node in inline_shape._inline.iter() if node.tag == qn("pic:blipFill")),
        None,
    )
    if blip_fill is None:
        raise ValueError("DOCX embedded picture has no DrawingML blipFill")
    for child in list(blip_fill):
        if child.tag == qn("a:srcRect"):
            blip_fill.remove(child)
    source_rect = OxmlElement("a:srcRect")
    for edge, key in (("l", "left"), ("t", "top"), ("r", "right"), ("b", "bottom")):
        source_rect.set(edge, str(round(float(crop_fraction[key]) * 100000)))
    insert_at = next(
        (index + 1 for index, child in enumerate(blip_fill) if child.tag == qn("a:blip")),
        0,
    )
    blip_fill.insert(insert_at, source_rect)


def figure_width_percent(classes: set[str]) -> int:
    for value in classes:
        match = re.fullmatch(r"artview-width-(\d+)", value)
        if match:
            return max(10, min(100, int(match.group(1))))
    if "art-body-right" in classes:
        return 25
    if "art-handout-page" in classes:
        return 80
    return 100


def _caption_reserve_height_mm(captions: list[Any], state: BuildState) -> float:
    """Estimate the space captions need below a figure within one usable page.

    Word's keep-with-next controls prevent an orphaned caption; this cap also
    prevents a full-height image from making that chain impossible to place.
    The conservative minimum includes the image paragraph's spacing and a
    one-line caption, while longer captions reserve additional wrapped lines.
    """
    if not captions or state.page_usable_width_mm is None:
        return 0.0
    defaults = _default_typography(
        {"body": "Times New Roman"},
        {"body_pt": 10.5},
        {"body_after_pt": 4.0, "body_line": 1.05},
    )
    typography = state.typography or defaults
    settings = typography["caption"]
    point_to_mm = 25.4 / 72.0
    average_character_mm = float(settings["size_pt"]) * point_to_mm * 0.5
    line_capacity = max(1, int(state.page_usable_width_mm / average_character_mm))
    caption_lines = sum(
        max(1, math.ceil(len(normalized_fragment(caption.text_content()).strip()) / line_capacity))
        for caption in captions
    )
    line_height_mm = float(settings["size_pt"]) * float(settings["line_spacing"]) * point_to_mm
    caption_spacing_mm = sum(
        (float(settings["space_before_pt"]) + float(settings["space_after_pt"])) * point_to_mm
        for _caption in captions
    )
    image_spacing_mm = (3.0 + 5.0) * point_to_mm
    # Allow for line-wrap variance and Word's paragraph layout rounding.
    return max(15.0, image_spacing_mm + caption_lines * line_height_mm + caption_spacing_mm + 2.0)


def _figure_dimensions(
    image_path: Path, classes: set[str], source: str, state: BuildState,
    captions: list[Any],
) -> tuple[float, float, bool]:
    if state.figure_policy is None:
        width_percent = figure_width_percent(classes)
        max_width_mm = 176.0 * width_percent / 100.0
        max_height_mm = 248.0 if ({"art-pregen-fullpage", "art-cover", "art-full-page"} & classes) else 220.0
        keep_with_next = False
    else:
        policy = state.figure_policy["overrides"].get(source, state.figure_policy["default"])
        max_width_mm = float(policy["max_width_mm"])
        max_height_mm = float(policy["max_height_mm"])
        keep_with_next = bool(policy["keep_with_next"])
        if state.page_usable_width_mm is not None:
            max_width_mm = min(max_width_mm, state.page_usable_width_mm)
        if state.page_usable_height_mm is not None:
            max_height_mm = min(max_height_mm, state.page_usable_height_mm)
            if captions:
                max_height_mm = max(1.0, max_height_mm - _caption_reserve_height_mm(captions, state))
    px_width, px_height = image_dimensions(image_path)
    crop_fraction = policy.get("crop_fraction") if state.figure_policy is not None else None
    if crop_fraction is not None:
        px_width *= 1 - float(crop_fraction["left"]) - float(crop_fraction["right"])
        px_height *= 1 - float(crop_fraction["top"]) - float(crop_fraction["bottom"])
    width_mm = max_width_mm
    height_mm = width_mm * px_height / max(px_width, 1)
    if height_mm > max_height_mm:
        height_mm = max_height_mm
        width_mm = height_mm * px_width / max(px_height, 1)
    return width_mm, height_mm, keep_with_next


def add_figure(doc: Document, element: Any, html_path: Path, state: BuildState) -> None:
    image_nodes = element.xpath(".//img[@src]")
    if not image_nodes:
        return
    image_node = image_nodes[0]
    src = image_node.get("src")
    image_path = (html_path.parent / src).resolve()
    if not image_path.is_file():
        state.missing_images.append(src)
        return

    classes = set((element.get("class") or "").split())
    captions = element.xpath("./figcaption")
    width_mm, height_mm, keep_with_next = _figure_dimensions(image_path, classes, src, state, captions)

    paragraph = doc.add_paragraph()
    if "artview-align-right" in classes or "art-body-right" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif "artview-align-left" in classes:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if state.figure_policy is not None or captions:
        paragraph.paragraph_format.keep_with_next = True if captions else keep_with_next
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run()
    policy = state.figure_policy["overrides"].get(src, state.figure_policy["default"]) if state.figure_policy else None
    crop_fraction = policy.get("crop_fraction") if policy else None
    fidelity_id = image_node.get("data-fidelity-placement-id")
    if fidelity_id is not None:
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", fidelity_id):
            raise ValueError("invalid DOCX fidelity placement ID")
        try:
            width_emu = int(image_node.get("data-fidelity-width-emu", ""))
            height_emu = int(image_node.get("data-fidelity-height-emu", ""))
            crop_fraction = {key: float(image_node.get(f"data-fidelity-crop-{key}", ""))
                             for key in ("left", "top", "right", "bottom")}
        except (TypeError, ValueError) as exc:
            raise ValueError("invalid DOCX fidelity geometry") from exc
        if (width_emu <= 0 or height_emu <= 0 or
                any(not math.isfinite(value) or not 0 <= value < 1 for value in crop_fraction.values()) or
                crop_fraction["left"] + crop_fraction["right"] >= 1 or
                crop_fraction["top"] + crop_fraction["bottom"] >= 1):
            raise ValueError("invalid DOCX fidelity geometry")
    try:
        inline_shape = add_picture_compatible(
            run, image_path, width=Emu(width_emu) if fidelity_id else Mm(width_mm),
            height=Emu(height_emu) if fidelity_id else Mm(height_mm),
            preserve_bytes=crop_fraction is not None,
        )
        apply_picture_crop(inline_shape, crop_fraction)
        if fidelity_id:
            inline_shape._inline.docPr.set("descr", fidelity_id)
    except Exception as exc:
        raise ValueError(f"DOCX cannot embed image {src}: {type(exc).__name__}: {exc}") from exc
    state.images += 1
    state.paragraphs += 1
    add_element_bookmark(paragraph, element, state)
    for image in image_nodes:
        add_rendered_bookmarks(paragraph, image, state)
    for index, caption in enumerate(captions):
        caption_paragraph = add_paragraph(doc, caption, state, style=ROLE_STYLE_NAMES["caption"])
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.keep_with_next = index < len(captions) - 1


def add_cover(doc: Document, section: Any, html_path: Path, state: BuildState) -> None:
    style = section.get("style") or ""
    match = re.search(r"background-image:\s*url\(([^)]+)\)", style)
    if not match:
        return
    src = match.group(1).strip(" '\"")
    image_path = (html_path.parent / src).resolve()
    if not image_path.is_file():
        state.missing_images.append(src)
        return
    px_width, px_height = image_dimensions(image_path)
    max_width_mm, max_height_mm = 176.0, 265.0
    width_mm = max_width_mm
    height_mm = width_mm * px_height / max(px_width, 1)
    if height_mm > max_height_mm:
        height_mm = max_height_mm
        width_mm = height_mm * px_width / max(px_height, 1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    try:
        add_picture_compatible(paragraph.add_run(), image_path, width=Mm(width_mm), height=Mm(height_mm))
    except Exception as exc:
        raise ValueError(f"DOCX cannot embed cover {src}: {type(exc).__name__}: {exc}") from exc
    state.images += 1
    state.paragraphs += 1
    add_element_bookmark(paragraph, section, state)
    doc.add_page_break()


def _ordered_list_numbering(doc: Document, state: BuildState, start: int) -> int:
    numbering = doc.part.numbering_part.element
    if state.ordered_list_abstract_num_id is None:
        abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(abstract_ids, default=-1) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi_level = OxmlElement("w:multiLevelType")
        multi_level.set(qn("w:val"), "singleLevel")
        abstract.append(multi_level)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("suff", "space"), ("lvlJc", "left")):
            child = OxmlElement(f"w:{tag}")
            child.set(qn("w:val"), value)
            level.append(child)
        abstract.append(level)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            first_num.addprevious(abstract)
        state.ordered_list_abstract_num_id = abstract_id

    concrete = numbering.add_num(state.ordered_list_abstract_num_id)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    concrete.append(override)
    return int(concrete.get(qn("w:numId")))


def _html_list_number(element: Any, attribute: str, default: int) -> int:
    raw = element.get(attribute)
    if raw is None:
        return default
    if not re.fullmatch(r"[+-]?\d+", str(raw).strip()):
        raise ValueError(f"HTML ordered-list {attribute} must be an integer")
    return int(str(raw).strip())


def add_list(doc: Document, element: Any, state: BuildState, *, ordered: bool, level: int = 0) -> None:
    style = "List Number" if ordered else "List Bullet"
    start = len(doc.paragraphs)
    num_id = _ordered_list_numbering(doc, state, _html_list_number(element, "start", 1)) if ordered else None
    for item in element.xpath("./li"):
        if ordered and item.get("value") is not None:
            num_id = _ordered_list_numbering(doc, state, _html_list_number(item, "value", 1))
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Mm(6 + level * 6)
        paragraph.paragraph_format.first_line_indent = Mm(-3)
        if num_id is not None:
            num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num_id
        add_inline(paragraph, item)
        add_rendered_bookmarks(paragraph, item, state)
        state.paragraphs += 1
        for nested in item.xpath("./ul|./ol"):
            add_list(doc, nested, state, ordered=str(nested.tag).lower() == "ol", level=level + 1)
    if len(doc.paragraphs) > start:
        add_element_bookmark(doc.paragraphs[start], element, state)


def toc_level(item: Any, fallback: int) -> int:
    for class_name in (item.get("class") or "").split():
        match = re.fullmatch(r"toc-level-(\d+)", class_name)
        if match:
            return max(1, min(6, int(match.group(1))))
    return fallback


def collect_toc_entries(element: Any, *, fallback_level: int = 1) -> list[tuple[int, str, str]]:
    entries: list[tuple[int, str, str]] = []
    for item in element.xpath("./li"):
        level = toc_level(item, fallback_level)
        anchors = item.xpath("./a[1]")
        source = anchors[0] if anchors else item
        text = normalized_fragment(source.text_content()).strip()
        if text:
            href = str(source.get("href") or "") if anchors else ""
            entries.append((level, text, href[1:] if href.startswith("#") else ""))
        for nested in item.xpath("./ol|./ul"):
            entries.extend(collect_toc_entries(nested, fallback_level=min(6, level + 1)))
    return entries


def collect_toc_groups(element: Any) -> list[list[tuple[int, str, str]]]:
    groups: list[list[tuple[int, str, str]]] = []
    for item in element.xpath("./li"):
        wrapper = lxml_html.Element("ul")
        wrapper.append(item)
        groups.append(collect_toc_entries(wrapper))
    return [group for group in groups if group]


def add_toc_entry(
    cell: Any, *, ordinal: int, level: int, text: str, anchor: str, first: bool,
    numbered: bool = True, keep_with_next: bool = False,
    source_flat: bool = False,
) -> None:
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.style = ROLE_STYLE_NAMES["toc"]
    level_offset_mm = max(0, level - 1) * 5.0
    text_start_mm = level_offset_mm + (9.0 if numbered else 0.0)
    paragraph.paragraph_format.left_indent = Mm(text_start_mm)
    paragraph.paragraph_format.first_line_indent = Mm(-9.0 if numbered else 0.0)
    if numbered:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(text_start_mm), WD_TAB_ALIGNMENT.LEFT)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next

    link_runs: list[Any] = []
    if numbered:
        number_run = paragraph.add_run(f"{ordinal}.")
        number_run.add_tab()
        link_runs.append(number_run._r)
    primary = level == 1 and not source_flat
    title_run = paragraph.add_run(text.upper() if primary and numbered else text)
    if primary:
        # The first TOC level is a semantic heading; all non-heading runs inherit.
        title_run.bold = True
    link_runs.append(title_run._r)
    if anchor:
        wrap_runs_in_internal_hyperlink(paragraph, link_runs, anchor)


def add_toc(
    doc: Document, element: Any, state: BuildState, *, numbered: bool = True,
    hierarchical: bool = False,
    source_flat: bool = False,
) -> None:
    entries = collect_toc_entries(element)
    if not entries:
        return

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    suppress_table_borders(table)
    if hierarchical:
        groups = collect_toc_groups(element)
        if len(groups) <= 1:
            columns = (entries, [])
        else:
            total = len(entries)
            split_group = min(range(1, len(groups)), key=lambda value: abs(sum(len(group) for group in groups[:value]) - total / 2))
            columns = (sum(groups[:split_group], []), sum(groups[split_group:], []))
    else:
        split_at = (len(entries) + 1) // 2
        columns = (entries[:split_at], entries[split_at:])
    preceding_entries = 0
    for column_index, (cell, column) in enumerate(zip(table.rows[0].cells, columns)):
        cell.width = Mm(86)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell)
        for position, (level, text, anchor) in enumerate(column):
            next_level = column[position + 1][0] if position + 1 < len(column) else None
            add_toc_entry(
                cell,
                ordinal=preceding_entries + position + 1,
                level=level,
                text=text,
                anchor=anchor,
                first=position == 0,
                numbered=numbered,
                keep_with_next=bool(hierarchical and next_level is not None and (level == 1 or next_level > 1)),
                source_flat=source_flat,
            )
        preceding_entries += len(column)
    state.paragraphs += len(entries)
    state.toc_entries += len(entries)


def _rich_toc_entries(element: Any) -> list[tuple[Any, Any, Any]]:
    entries: list[tuple[Any, Any, Any]] = []
    for item in element.xpath("./li"):
        children = [child for child in item if isinstance(child.tag, str)]
        if (len(children) != 3 or str(children[0].tag).lower() != "a"
                or str(children[1].tag).lower() != "p" or str(children[2].tag).lower() != "p"):
            raise ValueError("rich source TOC entry must be title, byline, synopsis")
        title, byline, synopsis = children
        if (byline.get("class") or "").split() != ["byline"] or (synopsis.get("class") or "").split() != ["synopsis"]:
            raise ValueError("rich source TOC entry roles are malformed")
        href = str(title.get("href") or "")
        if not href.startswith("#") or len(href) == 1:
            raise ValueError("rich source TOC title requires an internal anchor")
        if byline.xpath(".//a") or synopsis.xpath(".//a"):
            raise ValueError("rich source TOC metadata must not contain hyperlinks")
        if not all(normalized_fragment(value.text_content()).strip() for value in (title, byline, synopsis)):
            raise ValueError("rich source TOC entry text is empty")
        entries.append((title, byline, synopsis))
    if not entries:
        raise ValueError("rich source TOC has no entries")
    return entries


def _add_rich_toc_paragraph(cell: Any, element: Any, *, first: bool, title: bool = False) -> None:
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.style = ROLE_STYLE_NAMES["toc"]
    paragraph.paragraph_format.space_before = Pt(0); paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0; paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = title or "byline" in (element.get("class") or "").split()
    add_inline(paragraph, element)
    if title:
        for run in paragraph.runs:
            run.bold = True
        wrap_runs_in_internal_hyperlink(paragraph, [run._r for run in paragraph.runs], str(element.get("href"))[1:])


def add_rich_source_toc(doc: Document, element: Any, state: BuildState) -> None:
    entries = _rich_toc_entries(element)
    split_at = (len(entries) + 1) // 2
    table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    suppress_table_borders(table)
    for cell, column in zip(table.rows[0].cells, (entries[:split_at], entries[split_at:])):
        cell.width = Mm(86); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP; set_cell_margins(cell)
        for position, (title, byline, synopsis) in enumerate(column):
            _add_rich_toc_paragraph(cell, title, first=position == 0, title=True)
            _add_rich_toc_paragraph(cell, byline, first=False)
            _add_rich_toc_paragraph(cell, synopsis, first=False)
    state.paragraphs += len(entries) * 3; state.toc_entries += len(entries)


def add_heading(doc: Document, element: Any, state: BuildState) -> None:
    level = max(1, min(6, int(str(element.tag)[1:])))
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    if state.force_next_heading_page:
        paragraph.paragraph_format.page_break_before = True
        state.force_next_heading_page = False
    add_inline(paragraph, element)
    add_element_bookmark(paragraph, element, state)
    if state.statblock_depth:
        # Preserve heading outline/navigation while using compact card typography.
        role = doc.styles[ROLE_STYLE_NAMES["statblock"]]
        for run in paragraph.runs:
            set_run_font(run, role.font.name, role.font.size.pt)
        paragraph.paragraph_format.space_before = role.paragraph_format.space_before
        paragraph.paragraph_format.space_after = role.paragraph_format.space_after
        paragraph.paragraph_format.line_spacing = role.paragraph_format.line_spacing
    state.paragraphs += 1


def style_callout_paragraph(paragraph: Any, *, keep_with_next: bool, presentation: str = "legacy_shaded") -> None:
    paragraph.paragraph_format.keep_with_next = keep_with_next
    p_pr = paragraph._p.get_or_add_pPr()
    if presentation == "plain":
        for child in list(p_pr):
            if child.tag in {qn("w:shd"), qn("w:pBdr")}:
                p_pr.remove(child)
        return
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F0DFBD")
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "8F3D19")


def add_blockquote(doc: Document, element: Any, html_path: Path | None, state: BuildState) -> None:
    start = len(doc.paragraphs)
    children = [child for child in element if str(child.tag).lower() in BLOCK_TAGS]
    for index, child in enumerate(children):
        if str(child.tag).lower() == "p":
            paragraph = add_paragraph(doc, child, state, style=ROLE_STYLE_NAMES["quote"], left_indent_mm=7)
            paragraph.paragraph_format.right_indent = Mm(4)
            if state.callout_presentation == "legacy_shaded" and index == 0 and len(paragraph.text) < 120:
                for run in paragraph.runs:
                    run.bold = True
        else:
            add_block(doc, child, html_path=html_path, state=state)
    if "review-callout" in set((element.get("class") or "").split()):
        paragraphs = doc.paragraphs[start:]
        for index, paragraph in enumerate(paragraphs):
            style_callout_paragraph(paragraph, keep_with_next=index < len(paragraphs) - 1,
                                    presentation=state.callout_presentation)
        state.callouts += 1


def add_nav(doc: Document, element: Any, state: BuildState) -> None:
    if element.get("id") == "toc":
        headings = element.xpath(".//h1[1]")
        if headings:
            paragraph = doc.add_paragraph(style=ROLE_STYLE_NAMES["toc"])
            # The TOC table may not fit in the remaining cover-page space.
            # Keep its title attached so Word/LibreOffice move both together
            # instead of orphaning the title below the cover artwork.
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, headings[0])
            state.paragraphs += 1
        lists = element.xpath(".//ol[1]|.//ul[1]")
        if lists:
            numbered = str(element.get("data-toc-numbering") or "numbered").strip().casefold() != "none"
            layout_mode = str(element.get("data-toc-layout") or "").strip().casefold()
            hierarchical = layout_mode == "hierarchical"
            if not state.typography_configured:
                raw_font_size = str(element.get("data-toc-font-size-pt") or "9.5").strip()
                try:
                    font_size_pt = float(raw_font_size)
                except ValueError as exc:
                    raise ValueError("DOCX TOC font metadata must be numeric points") from exc
                if not 1 <= font_size_pt <= 72:
                    raise ValueError("DOCX TOC font metadata must be between 1 and 72 points")
                toc_style = doc.styles[ROLE_STYLE_NAMES["toc"]]
                toc_style.font.size = Pt(font_size_pt)
            if layout_mode == "rich-source":
                if numbered:
                    raise ValueError("rich source TOC must disable numbering")
                add_rich_source_toc(doc, lists[0], state)
            else:
                add_toc(doc, lists[0], state, numbered=numbered, hierarchical=hierarchical,
                        source_flat=layout_mode == "flat")
        doc.add_page_break()
        state.force_next_heading_page = False
        return

    for child in element:
        tag = str(child.tag).lower()
        if tag == "h1":
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, child)
            state.paragraphs += 1
        elif tag in {"ol", "ul"}:
            add_list(doc, child, state, ordered=tag == "ol")
        elif tag in BLOCK_TAGS:
            add_block(doc, child, html_path=None, state=state)
    doc.add_page_break()
    state.force_next_heading_page = False


def keep_paragraph_group(paragraphs: list[Any]) -> None:
    """Keep a finite group while explicitly releasing its final paragraph."""
    for index, paragraph in enumerate(paragraphs):
        properties = paragraph.get_or_add_pPr()
        for tag, value in (("keepLines", True), ("keepNext", index < len(paragraphs) - 1)):
            flag = properties.find(qn(f"w:{tag}"))
            if flag is None:
                flag = OxmlElement(f"w:{tag}")
                properties.append(flag)
            flag.set(qn("w:val"), "1" if value else "0")


def finalize_reader_statblock_groups(doc: Document, state: BuildState) -> None:
    stream = [n for n in doc._element.body if n.tag in {qn("w:p"), qn("w:tbl")}]
    for paragraphs in state.reader_statblock_groups.values():
        positions = [stream.index(p) for p in paragraphs]
        if (len(positions) < 2
                or positions != list(range(positions[0], positions[0] + len(positions)))
                or any(p.xpath('.//w:br[@w:type="page"] | ./w:pPr/w:pageBreakBefore[not(@w:val="0")]')
                       for p in paragraphs)):
            raise ValueError("reviewed statblock group has unsupported Word gaps or page breaks")
        if state.statblocks_keep_together:
            keep_paragraph_group(paragraphs)


def add_block(doc: Document, element: Any, html_path: Path | None, state: BuildState) -> None:
    group_id = element.get("data-reader-statblock-group")
    if group_id is None:
        _add_block(doc, element, html_path, state)
        return
    from kb_creator.pipeline.reader_statblock_grouping import TAGS, member_anchor
    if (str(element.tag).lower() not in TAGS
            or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,95}", group_id)):
        raise ValueError("reviewed statblock group member is unsupported")
    before = len(doc._element.body)
    _add_block(doc, element, html_path, state)
    emitted = list(doc._element.body)[before - 1:-1]  # Final sectPr stays last.
    if len(emitted) != 1 or emitted[0].tag != qn("w:p"):
        raise ValueError("reviewed statblock member must emit one body paragraph")
    members = state.reader_statblock_groups.setdefault(group_id, [])
    add_bookmark(doc.paragraphs[-1], member_anchor(group_id, len(members)), state)
    members.append(emitted[0])


def _add_block(doc: Document, element: Any, html_path: Path | None, state: BuildState) -> None:
    ensure_role_styles(doc, state)
    tag = str(element.tag).lower()
    classes = set((element.get("class") or "").split())
    if tag == "section" and "production-landing" in classes:
        if html_path is not None:
            add_cover(doc, element, html_path, state)
        return
    if tag == "nav":
        add_nav(doc, element, state)
        return
    if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        add_heading(doc, element, state)
        return
    if tag == "p":
        add_paragraph(doc, element, state, style=ROLE_STYLE_NAMES["statblock"] if state.statblock_depth else None)
        return
    if tag == "figure":
        if html_path is not None:
            add_figure(doc, element, html_path, state)
        return
    if tag == "table":
        add_table(doc, element, state)
        return
    if tag in {"ul", "ol"}:
        add_list(doc, element, state, ordered=tag == "ol")
        return
    if tag == "blockquote":
        add_blockquote(doc, element, html_path, state)
        return
    if tag in {"section", "div"}:
        statblock = tag == "section" and "statblock-card" in classes
        block_start = len(doc.paragraphs)
        table_start = len(doc.tables)
        body_start = (
            len([child for child in doc._element.body if child.tag in {qn("w:p"), qn("w:tbl")}])
            if statblock and state.statblocks_keep_together else None
        )
        if "production-section-start" in classes:
            state.force_next_heading_page = True
        if statblock:
            state.statblock_depth += 1
        try:
            for child in element:
                if str(child.tag).lower() in BLOCK_TAGS:
                    add_block(doc, child, html_path, state)
        finally:
            if statblock:
                state.statblock_depth -= 1
        if body_start is not None:
            paragraphs = [
                paragraph
                for child in [child for child in doc._element.body if child.tag in {qn("w:p"), qn("w:tbl")}][body_start:]
                for paragraph in child.iter(qn("w:p"))
            ]
            keep_paragraph_group(paragraphs)
        add_container_bookmark(doc, element, state, paragraph_start=block_start, table_start=table_start)


def _xml_text(element: Any) -> str:
    return normalized_fragment("".join(element.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))).strip()


def _paragraph_style_name(paragraph: Any) -> str:
    values = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=WORD_NAMESPACES)
    return str(values[0]) if values else ""


def _html_target_block(root: Any, anchor: str) -> Any:
    matches = root.xpath("//*[@id=$anchor]", anchor=anchor)
    if len(matches) != 1:
        raise ValueError(f"internal target must resolve exactly once: {anchor!r}")
    target = matches[0]
    block_tags = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li"}
    if str(target.tag).lower() in block_tags and normalized_fragment(target.text_content()).strip():
        return target
    parent = target.getparent()
    if parent is not None and str(parent.tag).lower() in block_tags and normalized_fragment(parent.text_content()).strip():
        return parent
    following = target.xpath(
        "following::*[self::h1 or self::h2 or self::h3 or self::h4 or self::h5 or self::h6 or self::p or self::li]"
    )
    for candidate in following:
        if normalized_fragment(candidate.text_content()).strip():
            return candidate
    raise ValueError(f"internal target has no rendered block: {anchor!r}")


def _html_internal_links(root: Any) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    toc_nodes = root.xpath("//nav[@id='toc']//a[starts-with(@href, '#')]")
    toc_identity = {id(node) for node in toc_nodes}
    body_nodes = [
        node
        for node in root.xpath("//main//a[starts-with(@href, '#')]")
        if id(node) not in toc_identity
    ]

    def rows(nodes: list[Any]) -> list[tuple[str, str]]:
        return [
            (str(node.get("href"))[1:], normalized_fragment(node.text_content()).strip())
            for node in nodes
        ]

    return rows(toc_nodes), rows(body_nodes)


def _find_target_paragraphs(root_html: Any, root_xml: Any, anchors: set[str]) -> dict[str, Any]:
    tree = root_html.getroottree()
    grouped: dict[str, dict[str, Any]] = {}
    for anchor in anchors:
        target = _html_target_block(root_html, anchor)
        target_path = tree.getpath(target)
        row = grouped.setdefault(
            target_path,
            {
                "target": target,
                "anchors": [],
                "order": int(target.xpath("count(preceding::*)")),
                "text": normalized_fragment(target.text_content()).strip(),
            },
        )
        row["anchors"].append(anchor)

    paragraphs = root_xml.xpath(".//w:body//w:p", namespaces=WORD_NAMESPACES)
    paragraph_texts = [_xml_text(paragraph) for paragraph in paragraphs]
    cursor = 0
    result: dict[str, Any] = {}
    for row in sorted(grouped.values(), key=lambda item: int(item["order"])):
        target = row["target"]
        target_text = str(row["text"])
        found = None
        for index in range(cursor, len(paragraphs)):
            if paragraph_texts[index].casefold() != target_text.casefold():
                continue
            found = index
            break
        if found is None:
            raise ValueError(
                f"DOCX target paragraph not found after position {cursor}: {target_text!r} ({sorted(row['anchors'])})"
            )
        for anchor in row["anchors"]:
            result[str(anchor)] = paragraphs[found]
        cursor = found + 1
    return result


def _is_visual_link_run(run: Any) -> bool:
    colors = run.xpath("./w:rPr/w:color/@w:val", namespaces=WORD_NAMESPACES)
    underlines = run.xpath("./w:rPr/w:u", namespaces=WORD_NAMESPACES)
    return bool(underlines and colors and str(colors[0]).upper() == "0563C1")


def _visual_link_groups(root_xml: Any) -> list[tuple[Any, list[Any], str]]:
    groups: list[tuple[Any, list[Any], str]] = []
    for paragraph in root_xml.xpath(".//w:body//w:p", namespaces=WORD_NAMESPACES):
        current: list[Any] = []
        for child in list(paragraph):
            if child.tag == qn("w:r") and _is_visual_link_run(child):
                current.append(child)
            elif child.tag == qn("w:r") and current:
                groups.append((
                    paragraph,
                    current,
                    normalized_fragment("".join(
                        "".join(run.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))
                        for run in current
                    )).strip(),
                ))
                current = []
        if current:
            groups.append((
                paragraph,
                current,
                normalized_fragment("".join(
                    "".join(run.xpath(".//w:t/text()", namespaces=WORD_NAMESPACES))
                    for run in current
                )).strip(),
            ))
    return groups


def _wrap_xml_runs(paragraph: Any, runs: list[Any], anchor: str) -> None:
    hyperlink = lxml_etree.Element(qn("w:hyperlink"))
    hyperlink.set(qn("w:anchor"), bookmark_name(anchor))
    hyperlink.set(qn("w:history"), "1")
    first_index = paragraph.index(runs[0])
    paragraph.insert(first_index, hyperlink)
    for run in runs:
        hyperlink.append(run)


def _add_xml_bookmarks(root_xml: Any, targets: dict[str, Any]) -> None:
    existing_ids = [
        int(value)
        for value in root_xml.xpath(".//w:bookmarkStart/@w:id", namespaces=WORD_NAMESPACES)
        if str(value).isdigit()
    ]
    next_id = max(existing_ids, default=0) + 1
    by_paragraph: dict[int, tuple[Any, list[str]]] = {}
    for anchor, paragraph in targets.items():
        row = by_paragraph.setdefault(id(paragraph), (paragraph, []))
        row[1].append(anchor)
    for paragraph, anchors in by_paragraph.values():
        insert_at = 1 if paragraph.find(qn("w:pPr")) is not None else 0
        ends: list[Any] = []
        for anchor in sorted(anchors):
            start = lxml_etree.Element(qn("w:bookmarkStart"))
            start.set(qn("w:id"), str(next_id))
            start.set(qn("w:name"), bookmark_name(anchor))
            end = lxml_etree.Element(qn("w:bookmarkEnd"))
            end.set(qn("w:id"), str(next_id))
            paragraph.insert(insert_at, start)
            insert_at += 1
            ends.append(end)
            next_id += 1
        for end in reversed(ends):
            paragraph.append(end)


def _style_xml_callout_paragraph(paragraph: Any, *, keep_with_next: bool, presentation: str = "legacy_shaded") -> None:
    p_pr = paragraph.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = lxml_etree.Element(qn("w:pPr"))
        paragraph.insert(0, p_pr)
    keep = p_pr.find(qn("w:keepNext"))
    if keep_with_next and keep is None:
        keep = lxml_etree.Element(qn("w:keepNext"))
        p_pr.append(keep)
    elif not keep_with_next and keep is not None:
        p_pr.remove(keep)
    if presentation == "plain":
        for child in list(p_pr):
            if child.tag in {qn("w:shd"), qn("w:pBdr")}:
                p_pr.remove(child)
        return
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = lxml_etree.Element(qn("w:shd"))
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F0DFBD")
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = lxml_etree.Element(qn("w:pBdr"))
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = lxml_etree.Element(qn("w:left"))
        borders.append(left)
    for key, value in {"val": "single", "sz": "24", "space": "6", "color": "8F3D19"}.items():
        left.set(qn(f"w:{key}"), value)


def _callout_list_paragraph_chunks(element: Any) -> list[str]:
    chunks: list[str] = []
    for item in element.xpath("./li"):
        text = normalized_fragment(item.text_content()).strip()
        if text:
            chunks.append(text)
        for nested in item.xpath("./ol|./ul"):
            chunks.extend(_callout_list_paragraph_chunks(nested))
    return chunks


def _callout_paragraph_chunks(element: Any) -> list[str]:
    """Return the visible body paragraphs that ``add_blockquote`` creates.

    Tables have their own OOXML containers.  Figures create an empty image
    paragraph and may also create visible caption paragraphs, which must take
    part in the matcher.  Containers recurse through the same block types
    that ``add_block`` accepts.
    """
    chunks: list[str] = []
    for child in element:
        tag = str(child.tag).lower()
        if tag == "p" or tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            text = normalized_fragment(child.text_content()).strip()
            if text:
                chunks.append(text)
        elif tag == "figure":
            # Images have no body-paragraph text, but add_figure emits each
            # figcaption as a caption paragraph after the image paragraph.
            for caption in child.xpath("./figcaption"):
                text = normalized_fragment(caption.text_content()).strip()
                if text:
                    chunks.append(text)
        elif tag in {"ol", "ul"}:
            chunks.extend(_callout_list_paragraph_chunks(child))
        elif tag in {"blockquote", "section", "div"}:
            chunks.extend(_callout_paragraph_chunks(child))
    return chunks


def _apply_xml_callouts(root_html: Any, root_xml: Any, *, presentation: str = "legacy_shaded") -> tuple[int, list[str]]:
    body_paragraphs = root_xml.xpath("./w:body/w:p", namespaces=WORD_NAMESPACES)
    body_texts = [_xml_text(paragraph) for paragraph in body_paragraphs]
    cursor = 0
    applied = 0
    missing: list[str] = []
    for blockquote in root_html.xpath("//blockquote[contains(concat(' ', normalize-space(@class), ' '), ' review-callout ')]"):
        # Semantic transport can leave an empty quote shell around a moved block.
        # Skip only a truly empty leaf, never unsupported or unmatched content.
        if not len(blockquote) and not (blockquote.text or "").strip():
            continue
        chunks = _callout_paragraph_chunks(blockquote)
        expected = normalized_fragment(" ".join(chunks)).strip()
        matched: list[int] | None = None
        for start in range(cursor, len(body_paragraphs)):
            if not body_texts[start]:
                continue
            combined = ""
            indices: list[int] = []
            for end in range(start, len(body_paragraphs)):
                if not body_texts[end]:
                    continue
                combined = normalized_fragment(f"{combined} {body_texts[end]}").strip()
                indices.append(end)
                if combined == expected:
                    matched = indices
                    break
                if not expected.startswith(combined):
                    break
            if matched is not None:
                break
        if matched is None:
            missing.append(expected[:160])
            continue
        for index, paragraph_index in enumerate(matched):
            _style_xml_callout_paragraph(
                body_paragraphs[paragraph_index], keep_with_next=index < len(matched) - 1,
                presentation=presentation,
            )
        cursor = matched[-1] + 1
        applied += 1
    return applied, missing


def patch_existing_docx_projection(root_html: Any, docx_path: Path, output_path: Path | None = None,
                                   *, callout_presentation: str = "legacy_shaded") -> dict[str, Any]:
    """Patch navigation and callouts from a validated in-memory projection."""
    if callout_presentation not in {"plain", "legacy_shaded"}:
        raise ValueError("DOCX callout presentation is unsupported")
    docx_path = docx_path.resolve()
    output_path = (output_path or docx_path).resolve()
    validate_unique_html_ids(root_html)
    toc_links, body_links = _html_internal_links(root_html)
    expected_links = len(toc_links) + len(body_links)

    with zipfile.ZipFile(docx_path) as source:
        if source.testzip() is not None:
            raise ValueError("source DOCX archive is corrupt")
        root_xml = lxml_etree.fromstring(source.read("word/document.xml"))
        if root_html.xpath("//nav[@id='toc' and @data-toc-layout='flat']"):
            toc_tables = root_xml.xpath(".//w:body/w:tbl[1]", namespaces=WORD_NAMESPACES)
            if len(toc_tables) != 1:
                raise ValueError("flat source TOC requires one leading DOCX table")
            toc_texts = [
                re.sub(r"^\d+\.\s*", "", _xml_text(paragraph))
                for paragraph in toc_tables[0].xpath("./w:tr/w:tc/w:p", namespaces=WORD_NAMESPACES)
                if _xml_text(paragraph)
            ]
            if toc_texts != [text for _anchor, text in toc_links]:
                raise ValueError("DOCX flat TOC display text differs from accepted HTML")
        existing_links = root_xml.xpath(".//w:hyperlink", namespaces=WORD_NAMESPACES)
        expected_anchor_sequence = [bookmark_name(anchor) for anchor, _ in toc_links + body_links]
        if existing_links:
            if len(existing_links) != expected_links:
                raise ValueError(
                    f"existing DOCX has a partial hyperlink set: {len(existing_links)} of {expected_links}"
                )
            actual_anchor_sequence = [str(link.get(qn("w:anchor")) or "") for link in existing_links]
            if actual_anchor_sequence != expected_anchor_sequence:
                raise ValueError("existing DOCX hyperlink targets differ from the HTML navigation sequence")
        elif expected_links:
            toc_table = root_xml.xpath(".//w:body/w:tbl[1]", namespaces=WORD_NAMESPACES)
            if len(toc_table) != 1:
                raise ValueError("DOCX must contain one leading two-column TOC table")
            toc_paragraphs = [
                paragraph
                for paragraph in toc_table[0].xpath("./w:tr/w:tc/w:p", namespaces=WORD_NAMESPACES)
                if _xml_text(paragraph)
            ]
            if len(toc_paragraphs) != len(toc_links):
                raise ValueError(f"DOCX TOC has {len(toc_paragraphs)} entries, expected {len(toc_links)}")
            for paragraph, (anchor, _text) in zip(toc_paragraphs, toc_links):
                runs = paragraph.xpath("./w:r", namespaces=WORD_NAMESPACES)
                _wrap_xml_runs(paragraph, runs, anchor)

            visual_groups = _visual_link_groups(root_xml)
            if len(visual_groups) != len(body_links):
                raise ValueError(
                    f"DOCX has {len(visual_groups)} visual body-link groups, expected {len(body_links)}"
                )
            for (paragraph, runs, actual_text), (anchor, expected_text) in zip(visual_groups, body_links):
                if actual_text.casefold() != expected_text.casefold():
                    raise ValueError(
                        f"body-link order mismatch: DOCX {actual_text!r}, HTML {expected_text!r}"
                    )
                _wrap_xml_runs(paragraph, runs, anchor)

        target_anchors = {anchor for anchor, _ in toc_links + body_links}
        existing_bookmarks = set(
            root_xml.xpath(".//w:bookmarkStart/@w:name", namespaces=WORD_NAMESPACES)
        )
        missing_target_anchors = {
            anchor for anchor in target_anchors if bookmark_name(anchor) not in existing_bookmarks
        }
        if missing_target_anchors:
            targets = _find_target_paragraphs(root_html, root_xml, missing_target_anchors)
            _add_xml_bookmarks(root_xml, targets)

        callouts_applied, callouts_missing = _apply_xml_callouts(root_html, root_xml, presentation=callout_presentation)
        document_xml = lxml_etree.tostring(
            root_xml, xml_declaration=True, encoding="UTF-8", standalone="yes"
        )
        temp_path = output_path.with_name(f".{output_path.name}.navigation.tmp")
        with zipfile.ZipFile(temp_path, "w") as destination:
            for info in source.infolist():
                payload = document_xml if info.filename == "word/document.xml" else source.read(info.filename)
                destination.writestr(info, payload)
    temp_path.replace(output_path)
    return {
        "status": "pass",
        "output": output_path.as_posix(),
        "sha256": sha256(output_path),
        "bytes": output_path.stat().st_size,
        "toc_hyperlinks": len(toc_links),
        "body_hyperlinks": len(body_links),
        "bookmarks": len({anchor for anchor, _ in toc_links + body_links}),
        "callouts_applied": callouts_applied,
        "callouts_missing": callouts_missing,
    }


def patch_existing_docx(html_path: Path, docx_path: Path, output_path: Path | None = None,
                        *, callout_presentation: str = "legacy_shaded") -> dict[str, Any]:
    """Legacy file-backed patch adapter for archived publication tooling."""
    root = lxml_html.parse(str(html_path.resolve())).getroot()
    return patch_existing_docx_projection(root, docx_path, output_path,
                                          callout_presentation=callout_presentation)


def build_docx_projection(root: Any, asset_base: Path, output_path: Path,
                          *, config: Path | dict[str, Any] | None = None) -> dict[str, Any]:
    """Build Word from an in-memory projection and explicit asset directory."""
    output_path = output_path.resolve()
    asset_base = asset_base.resolve(strict=True)
    if not asset_base.is_dir():
        raise ValueError("projection asset base must be a directory")
    validate_unique_html_ids(root)
    main_nodes = root.xpath("//main")
    if len(main_nodes) != 1:
        raise ValueError("production HTML must contain exactly one main element")
    # The Word construction pass may move nodes while composing nested TOCs.
    # Keep the original projection for the semantic hyperlink/callout pass.
    semantic_root = copy.deepcopy(root)

    profile = None
    if config is not None:
        profile = json.loads(config.read_text(encoding="utf-8-sig")) if isinstance(config, Path) else config
        profile = validate_document_config(profile)
    doc = Document()
    configure_document(doc, profile)
    metadata = profile["metadata"] if profile else {"title": "Вирус", "subject": "Сценарий Call of Cthulhu", "author": "NRI", "keywords": "", "comments": ""}
    doc.core_properties.title = metadata["title"]
    doc.core_properties.subject = metadata["subject"]
    doc.core_properties.author = metadata["author"]
    doc.core_properties.keywords = metadata["keywords"]
    doc.core_properties.comments = metadata["comments"]
    state = build_state_for_profile(profile)
    projection_location = asset_base / ".in-memory-projection"
    for element in main_nodes[0]:
        if str(element.tag).lower() in BLOCK_TAGS:
            add_block(doc, element, projection_location, state)
    finalize_reader_statblock_groups(doc, state)
    if state.missing_images:
        raise FileNotFoundError(f"DOCX source images are missing: {sorted(state.missing_images)}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f".{output_path.name}.tmp")
    doc.save(temp_path)
    semantic_result = patch_existing_docx_projection(semantic_root, temp_path, callout_presentation=state.callout_presentation)
    if semantic_result["callouts_missing"]:
        raise ValueError(
            f"generated DOCX lost review callouts: {semantic_result['callouts_missing']}"
        )
    with zipfile.ZipFile(temp_path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names or "[Content_Types].xml" not in names:
            raise ValueError("generated DOCX is missing required OOXML members")
    temp_path.replace(output_path)
    return {
        "status": "pass",
        "output": output_path.as_posix(),
        "sha256": sha256(output_path),
        "bytes": output_path.stat().st_size,
        "paragraphs": state.paragraphs,
        "tables": state.tables,
        "images": state.images,
        "toc_entries": state.toc_entries,
        "callouts": state.callouts,
        "internal_hyperlinks": semantic_result["toc_hyperlinks"] + semantic_result["body_hyperlinks"],
        "bookmarks": semantic_result["bookmarks"],
        "config_sha256": hashlib.sha256(json.dumps(profile, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest().upper() if profile is not None else None,
    }


def build_docx(html_path: Path, output_path: Path, *, config: Path | dict[str, Any] | None = None) -> dict[str, Any]:
    """Legacy file-backed adapter; new normal builds call build_docx_projection."""
    html_path = html_path.resolve()
    if not html_path.is_file():
        raise FileNotFoundError(html_path)
    root = lxml_html.parse(str(html_path)).getroot()
    return build_docx_projection(root, html_path.parent, output_path, config=config)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("html", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--patch-existing", action="store_true")
    args = parser.parse_args()
    try:
        result = (
            patch_existing_docx(args.html, args.output)
            if args.patch_existing
            else build_docx(args.html, args.output, config=args.config)
        )
        print(json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2))
        return 0
    except Exception as exc:
        print(json.dumps({"status": "fail", "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False, indent=2))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
