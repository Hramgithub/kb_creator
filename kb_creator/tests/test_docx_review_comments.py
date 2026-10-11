from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from kb_creator.pipeline import docx_review_comments as comments_module
from kb_creator.pipeline.docx_review_comments import DocxReviewCommentsError, annotate_docx_review
from kb_creator.pipeline.identity import sha256_file
from kb_creator.scripts.production_docx import bookmark_name


def _bookmark(paragraph, anchor: str, identifier: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), bookmark_name(anchor))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _replace_member(path: Path, member: str, payload: bytes) -> None:
    temporary = path.with_suffix(".rewrite.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            target.writestr(info, payload if info.filename == member else source.read(info.filename))
    temporary.replace(path)


def _existing_comment(path: Path) -> None:
    comments = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <w:comment w:id="7" w:author="Reader" w:initials="R"><w:p><w:r><w:t>Old note</w:t></w:r></w:p></w:comment>
</w:comments>'''
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/comments.xml", comments)


def _hyperlink(paragraph, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


@pytest.fixture
def source_docx(tmp_path: Path) -> Path:
    path = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("До ")
    bold = paragraph.add_run("точный ")
    bold.bold = True
    italic = paragraph.add_run("текст")
    italic.italic = True
    paragraph.add_run(" после")
    _bookmark(paragraph, "body-anchor", 1)
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0).paragraphs[0]
    cell.add_run("Табличная ")
    target = cell.add_run("цель")
    target.bold = True
    _bookmark(cell, "table-anchor", 2)
    document.save(path)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/media/retained.bin", b"retained-media")
    return path


def _issues(path: Path, issues: list[dict[str, object]] | None = None) -> dict[str, object]:
    return {
        "schema": "kb_creator.docx_review_issues.v1",
        "docx_sha256": sha256_file(path),
        "issues": issues or [
            {"issue_id": "I-1", "reader_anchor": "body-anchor", "target_text": "точный текст", "reason": "Проверить чтение", "proposal": "точное чтение", "source_page": "12", "en_excerpt": "exact text"},
            {"issue_id": "I-2", "reader_anchor": "table-anchor", "target_text": "цель", "reason": "Проверить таблицу", "source_page": "13", "en_excerpt": "table target"},
        ],
    }


def _write_issues(path: Path, value: dict[str, object]) -> Path:
    registry = path.with_name("issues.json")
    registry.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    return registry


def test_annotates_cyrillic_across_formatted_runs_table_and_existing_comments(source_docx: Path) -> None:
    _existing_comment(source_docx)
    source_hash = sha256_file(source_docx)
    issues_path = _write_issues(source_docx, _issues(source_docx))
    output = source_docx.with_name("annotated.docx")

    report = annotate_docx_review(source_docx, issues_path, output)

    assert sha256_file(source_docx) == source_hash
    assert [row["comment_id"] for row in report["bindings"]] == [8, 9]
    assert report["output"]["sha256"] == sha256_file(output)
    with zipfile.ZipFile(source_docx) as before, zipfile.ZipFile(output) as after:
        assert after.read("word/media/retained.bin") == before.read("word/media/retained.bin")
        document = after.read("word/document.xml").decode("utf-8")
        comments = after.read("word/comments.xml").decode("utf-8")
        rels = after.read("word/_rels/document.xml.rels").decode("utf-8")
        types = after.read("[Content_Types].xml").decode("utf-8")
    assert document.count('w:commentRangeStart') == 2
    assert 'w:commentReference w:id="8"' in document
    assert '<w:b/>' in document and '<w:i/>' in document
    assert "Old note" in comments and "I-1" in comments and "I-2" in comments
    assert rels.count("relationships/comments") == 1 and 'Target="comments.xml"' in rels
    assert "wordprocessingml.comments+xml" in types


@pytest.mark.parametrize("fault", ["missing", "duplicate", "repeated"])
def test_refuses_missing_duplicate_or_ambiguous_bindings(source_docx: Path, fault: str) -> None:
    issues = _issues(source_docx)
    if fault == "missing":
        issues["issues"][0]["reader_anchor"] = "absent"
    elif fault == "duplicate":
        with zipfile.ZipFile(source_docx) as archive:
            document = archive.read("word/document.xml").replace(
                b"</w:body>",
                f'<w:p><w:bookmarkStart w:id="88" w:name="{bookmark_name("body-anchor")}"/><w:r><w:t>copy</w:t></w:r><w:bookmarkEnd w:id="88"/></w:p></w:body>'.encode(),
            )
        _replace_member(source_docx, "word/document.xml", document)
        issues["docx_sha256"] = sha256_file(source_docx)
    else:
        issues["issues"][0]["target_text"] = "текст"
        # Add a second exact occurrence to the same bookmark paragraph.
        with zipfile.ZipFile(source_docx) as archive:
            document = archive.read("word/document.xml").replace(" после".encode(), " текст после".encode())
        _replace_member(source_docx, "word/document.xml", document)
        issues["docx_sha256"] = sha256_file(source_docx)
    with pytest.raises(DocxReviewCommentsError):
        annotate_docx_review(source_docx, _write_issues(source_docx, issues), source_docx.with_name("out.docx"))


def test_refuses_stale_source_hash_and_output_overwrite(source_docx: Path) -> None:
    stale = _issues(source_docx)
    stale["docx_sha256"] = "0" * 64
    with pytest.raises(DocxReviewCommentsError, match="docx_sha256"):
        annotate_docx_review(source_docx, _write_issues(source_docx, stale), source_docx.with_name("out.docx"))
    output = source_docx.with_name("already.docx")
    output.write_bytes(b"keep")
    with pytest.raises(FileExistsError):
        annotate_docx_review(source_docx, _write_issues(source_docx, _issues(source_docx)), output)
    assert output.read_bytes() == b"keep"


def test_cli_writes_annotation_and_binding_report(source_docx: Path) -> None:
    issues = _write_issues(source_docx, _issues(source_docx))
    output, report = source_docx.with_name("cli.docx"), source_docx.with_name("binding.json")
    command = [sys.executable, "kb_creator/scripts/annotate_docx_review.py", "--source", str(source_docx), "--issues", str(issues), "--output", str(output), "--report", str(report)]
    completed = subprocess.run(command, cwd=Path(__file__).parents[2], text=True, capture_output=True, check=False)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    written = json.loads(report.read_text(encoding="utf-8"))
    assert written["source"]["sha256"] == sha256_file(source_docx)
    assert written["output"]["sha256"] == sha256_file(output)


@pytest.mark.parametrize(
    "construct",
    [b"<w:br/>", b"<w:tab/>", b'<w:fldChar w:fldCharType="begin"/>', b"<w:drawing/>", b"<w:control/>"],
)
def test_refuses_unmodeled_paragraph_constructs(source_docx: Path, construct: bytes) -> None:
    with zipfile.ZipFile(source_docx) as archive:
        document = archive.read("word/document.xml").replace(
            b"</w:r><w:r>", b"</w:r><w:r>" + construct + b"</w:r><w:r>", 1
        )
    _replace_member(source_docx, "word/document.xml", document)
    with pytest.raises(DocxReviewCommentsError, match="unsupported"):
        annotate_docx_review(source_docx, _write_issues(source_docx, _issues(source_docx)), source_docx.with_name("out.docx"))


@pytest.mark.parametrize("style_inherited", [False, True])
def test_refuses_direct_or_style_inherited_hidden_text(tmp_path: Path, style_inherited: bool) -> None:
    source = tmp_path / "styled-hidden.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("невидимая цель")
    if style_inherited:
        hidden_style = document.styles.add_style("Hidden review text", WD_STYLE_TYPE.CHARACTER)
        hidden_style.font.hidden = True
        run.style = hidden_style
    else:
        run.font.hidden = True
    _bookmark(paragraph, "hidden-anchor", 1)
    document.save(source)
    issues = _issues(source, [{"issue_id": "hidden", "reader_anchor": "hidden-anchor", "target_text": "невидимая цель", "reason": "hidden", "source_page": 1, "en_excerpt": "hidden"}])

    with pytest.raises(DocxReviewCommentsError, match="hidden text"):
        annotate_docx_review(source, _write_issues(source, issues), source.with_name("out.docx"))


def test_preserves_exact_hyperlink_target_without_widening(tmp_path: Path) -> None:
    source = tmp_path / "hyperlink.docx"
    document = Document()
    paragraph = document.add_paragraph("До ")
    _hyperlink(paragraph, "ссылка")
    paragraph.add_run(" далее")
    _bookmark(paragraph, "hyperlink-anchor", 1)
    document.save(source)
    issues = _issues(source, [{"issue_id": "link", "reader_anchor": "hyperlink-anchor", "target_text": "ссылка далее", "reason": "link", "source_page": 1, "en_excerpt": "link"}])
    output = source.with_name("annotated.docx")

    annotate_docx_review(source, _write_issues(source, issues), output)

    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:hyperlink><w:commentRangeStart" in xml
    assert "<w:t>ссылка</w:t>" in xml and "<w:t xml:space=\"preserve\"> далее</w:t>" in xml
    assert xml.index("commentRangeStart") < xml.index("commentRangeEnd") < xml.index("commentReference")


def test_plain_comment_reference_follows_range_end(source_docx: Path) -> None:
    issues = _issues(source_docx, [_issues(source_docx)["issues"][0]])
    output = source_docx.with_name("plain.docx")

    annotate_docx_review(source_docx, _write_issues(source_docx, issues), output)

    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.index("commentRangeStart") < xml.index("commentRangeEnd") < xml.index("commentReference")


def test_refuses_all_source_registry_output_report_aliases(source_docx: Path) -> None:
    registry = _write_issues(source_docx, _issues(source_docx))
    fresh_output = source_docx.with_name("fresh.docx")
    with pytest.raises(DocxReviewCommentsError, match="alias"):
        annotate_docx_review(source_docx, registry, source_docx)
    with pytest.raises(DocxReviewCommentsError, match="alias"):
        annotate_docx_review(source_docx, registry, registry)
    with pytest.raises(DocxReviewCommentsError, match="alias"):
        comments_module.write_annotation_report(source_docx, registry, fresh_output, fresh_output)
    with pytest.raises(DocxReviewCommentsError, match="alias"):
        comments_module.write_annotation_report(source_docx, registry, fresh_output, source_docx)
    with pytest.raises(DocxReviewCommentsError, match="alias"):
        comments_module.write_annotation_report(source_docx, registry, fresh_output, registry)


@pytest.mark.parametrize("target", ["other-comments.xml", "comments.xml"])
def test_rejects_invalid_or_duplicate_existing_comments_relationship(source_docx: Path, target: str) -> None:
    with zipfile.ZipFile(source_docx) as archive:
        relationships = archive.read("word/_rels/document.xml.rels")
    relationship = (
        b'<Relationship Id="rId999" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
        + f'Target="{target}"/>'.encode()
    )
    extra = relationship if target != "comments.xml" else relationship + relationship.replace(b'rId999', b'rId998')
    _replace_member(source_docx, "word/_rels/document.xml.rels", relationships.replace(b"</Relationships>", extra + b"</Relationships>"))
    with pytest.raises(DocxReviewCommentsError, match="comments relationship"):
        annotate_docx_review(source_docx, _write_issues(source_docx, _issues(source_docx)), source_docx.with_name("out.docx"))


def test_refuses_existing_temporary_sibling_and_concurrent_output(source_docx: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    output = source_docx.with_name("race.docx")
    temporary = output.with_name(f".{output.name}.tmp")
    temporary.write_bytes(b"prior-build")
    with pytest.raises(FileExistsError):
        annotate_docx_review(source_docx, _write_issues(source_docx, _issues(source_docx)), output)
    assert temporary.read_bytes() == b"prior-build"
    temporary.unlink()

    real_link = comments_module.os.link

    def concurrent_creator(source: Path, destination: Path) -> None:
        destination.write_bytes(b"concurrent-owner")
        real_link(source, destination)

    monkeypatch.setattr(comments_module.os, "link", concurrent_creator)
    with pytest.raises(FileExistsError):
        annotate_docx_review(source_docx, _write_issues(source_docx, _issues(source_docx)), output)
    assert output.read_bytes() == b"concurrent-owner"
    assert not temporary.exists()
