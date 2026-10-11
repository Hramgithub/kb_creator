from docx import Document

from kb_creator.pipeline.five_step_docx import _default_config
from kb_creator.scripts.production_docx import build_docx


def test_statblock_heading_uses_card_size_and_preserves_outline_and_inline(tmp_path):
    source = tmp_path / 'source.html'
    source.write_text('<html><head><meta charset="utf-8"></head><body><main>'
        '<nav id="toc"><h1>Содержание</h1><ol><li><a href="#chapter">Глава</a></li></ol></nav>'
        '<h2 id="chapter">Глава</h2>'
        '<section class="statblock-card" data-semantic-kind="statblock">'
        '<h2 id="damage">Бонус: <em>+1D4</em></h2><p>ПЗ: 10</p>'
        '</section><h2>После карточки</h2></main></body></html>', encoding='utf-8')
    output = tmp_path / 'result.docx'
    config = _default_config('Карточка')
    config['sizes']['body_pt'] = 11
    build_docx(source, output, config=config)
    doc = Document(output)
    p = next(p for p in doc.paragraphs if p.text == 'Бонус: +1D4')
    assert p.style.name == 'Heading 2'
    assert all(r.font.size.pt == 11 for r in p.runs if r.text)
    assert next(r for r in p.runs if r.text == '+1D4').italic
    assert p._p.xpath('.//w:bookmarkStart')
    after = next(p for p in doc.paragraphs if p.text == 'После карточки')
    assert after.style.name == 'Heading 2'
    assert all(r.font.size is None for r in after.runs if r.text)
    assert after.style.font.size.pt == config['headings']['2']['size_pt']
