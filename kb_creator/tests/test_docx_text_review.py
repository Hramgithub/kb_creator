from __future__ import annotations

import copy
import json
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from kb_creator.pipeline.docx_text_review import (
    DocxTextReviewError,
    prepare_docx_text_review,
    validate_docx_text_review,
)
from kb_creator.pipeline.docx_review_comments import annotate_docx_review
from kb_creator.pipeline.identity import sha256_file
from kb_creator.scripts.production_docx import BuildState, add_bookmark


def _hyperlink(paragraph, text: str) -> None:
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def _add_textbox(paragraph, text: str) -> None:
    # A minimal DrawingML text-box tree is sufficient for Word OOXML traversal.
    drawing = OxmlElement("w:drawing")
    txbx = OxmlElement("w:txbxContent")
    inner = OxmlElement("w:p")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    inner.append(run)
    txbx.append(inner)
    drawing.append(txbx)
    paragraph._p.append(drawing)


def _replace_zip_member(path: Path, member: str, data: bytes) -> None:
    temp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, data if item.filename == member else source.read(item.filename))
    temp.replace(path)


@pytest.fixture
def review_docx(tmp_path: Path) -> Path:
    path = tmp_path / "review.docx"
    document = Document()
    document.add_paragraph("Основной абзац")
    linked = document.add_paragraph("Ссылка: ")
    _hyperlink(linked, "видимый гипертекст")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Ячейка таблицы"
    textbox_host = document.add_paragraph()
    _add_textbox(textbox_host, "Текст в рамке")
    header = document.sections[0].header.paragraphs[0]
    header.text = "Верхний колонтитул"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.save(path)
    footnotes = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:footnote w:id="1"><w:p><w:r><w:t>Сноска</w:t></w:r></w:p></w:footnote>
    </w:footnotes>'''.encode("utf-8")
    endnotes = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:endnote w:id="1"><w:p><w:r><w:t>Концевая сноска</w:t></w:r></w:p></w:endnote>
    </w:endnotes>'''.encode("utf-8")
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/footnotes.xml", footnotes)
        archive.writestr("word/endnotes.xml", endnotes)
    return path


@pytest.fixture
def annotated_review_docx(tmp_path: Path) -> Path:
    """Create native review comments through the production annotation path."""
    source = tmp_path / "comment-source.docx"
    document = Document()
    paragraph = document.add_paragraph("Проверяемая формулировка")
    add_bookmark(paragraph, "comment-target", BuildState(bookmarked_anchors=set()))
    document.save(source)
    registry = tmp_path / "issues.json"
    registry.write_text(
        json.dumps(
            {
                "schema": "kb_creator.docx_review_issues.v1",
                "docx_sha256": sha256_file(source),
                "issues": [{
                    "issue_id": "issue-1",
                    "reader_anchor": "comment-target",
                    "target_text": "Проверяемая формулировка",
                    "reason": "Проверить формулировку",
                    "en_excerpt": "Review the wording",
                    "source_page": 1,
                }],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    annotated = tmp_path / "comment-annotated.docx"
    annotate_docx_review(source, registry, annotated)

    # Native comments can contain several paragraphs; retain each as review text.
    with zipfile.ZipFile(annotated) as archive:
        comments = ET.fromstring(archive.read("word/comments.xml"))
    comment = next(element for element in comments if element.tag.endswith("comment"))
    paragraph = ET.SubElement(comment, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
    run = ET.SubElement(paragraph, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    ET.SubElement(run, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t").text = "Продолжение комментария"
    _replace_zip_member(annotated, "word/comments.xml", ET.tostring(comments, encoding="utf-8", xml_declaration=True))
    return annotated


def _submission(task: dict) -> dict:
    return {
        "schema": "kb_creator.docx_text_review_submission.v1",
        "task_sha256": task["task_sha256"],
        "docx": copy.deepcopy(task["docx"]),
        "reviewer_role": "agent",
        "blocks": [
            {"block_id": block["block_id"], "text_sha256": block["text_sha256"], "read": True}
            for block in task["blocks"]
        ],
        "findings": [],
    }


def test_extracts_full_visible_text_coverage_across_containers(review_docx: Path) -> None:
    task = prepare_docx_text_review(review_docx)
    text = [block["text"] for block in task["blocks"]]
    assert {"Основной абзац", "Ссылка: видимый гипертекст", "Ячейка таблицы", "Текст в рамке", "Верхний колонтитул", "Сноска", "Концевая сноска"}.issubset(text)
    assert any(block["container"] == "table_cell" for block in task["blocks"])
    assert any(block["container"] == "textbox" for block in task["blocks"])
    assert text.count("Текст в рамке") == 1
    assert len(task["blocks"]) == len({block["block_id"] for block in task["blocks"]})
    record = validate_docx_text_review(task, _submission(task))
    assert record["schema"] == "kb_creator.docx_text_review_record.v1"


@pytest.mark.parametrize("defect", ["omitted", "reordered", "stale", "tampered"])
def test_rejects_nonexhaustive_or_stale_block_evidence(review_docx: Path, defect: str) -> None:
    task = prepare_docx_text_review(review_docx)
    submission = _submission(task)
    if defect == "omitted":
        submission["blocks"].pop()
    elif defect == "reordered":
        submission["blocks"][0], submission["blocks"][1] = submission["blocks"][1], submission["blocks"][0]
    elif defect == "stale":
        submission["task_sha256"] = "0" * 64
    else:
        submission["blocks"][0]["text_sha256"] = "0" * 64
    with pytest.raises(DocxTextReviewError):
        validate_docx_text_review(task, submission)


def test_rejects_unresolved_finding_and_modified_docx_identity(review_docx: Path) -> None:
    task = prepare_docx_text_review(review_docx)
    submission = _submission(task)
    submission["findings"] = [{"finding_id": "f1", "block_id": task["blocks"][0]["block_id"], "status": "open", "resolution": ""}]
    with pytest.raises(DocxTextReviewError, match="unresolved"):
        validate_docx_text_review(task, submission)
    modified = Document(review_docx)
    modified.add_paragraph("Новая строка")
    modified.save(review_docx)
    new_task = prepare_docx_text_review(review_docx)
    with pytest.raises(DocxTextReviewError, match="stale"):
        validate_docx_text_review(new_task, _submission(task))


def test_rejects_non_agent_reviewer_role(review_docx: Path) -> None:
    task = prepare_docx_text_review(review_docx)
    submission = _submission(task)
    submission["reviewer_role"] = "human"
    with pytest.raises(DocxTextReviewError, match="reviewer_role"):
        validate_docx_text_review(task, submission)


def test_covers_native_multi_paragraph_comments(annotated_review_docx: Path) -> None:
    task = prepare_docx_text_review(annotated_review_docx)
    comments = [block for block in task["blocks"] if block["part"] == "word/comments.xml"]
    assert [block["text"] for block in comments][-2:] == ["[issue-1] Проверить формулировку\nИсточник: стр. 1; EN: Review the wording", "Продолжение комментария"]
    assert {block["comment_id"] for block in comments} == {"0"}
    assert {block["container"] for block in comments} == {"comment"}
    assert validate_docx_text_review(task, _submission(task))["schema"] == "kb_creator.docx_text_review_record.v1"


def test_rejects_missing_or_stale_native_comment_coverage(annotated_review_docx: Path) -> None:
    task = prepare_docx_text_review(annotated_review_docx)
    missing = _submission(task)
    missing["blocks"] = [block for block in missing["blocks"] if block["block_id"] != next(item["block_id"] for item in task["blocks"] if item["part"] == "word/comments.xml")]
    with pytest.raises(DocxTextReviewError, match="full DOCX block coverage"):
        validate_docx_text_review(task, missing)

    with zipfile.ZipFile(annotated_review_docx) as archive:
        changed = archive.read("word/comments.xml").replace("Проверить формулировку".encode("utf-8"), "Изменённый комментарий".encode("utf-8"))
    _replace_zip_member(annotated_review_docx, "word/comments.xml", changed)
    changed_task = prepare_docx_text_review(annotated_review_docx)
    with pytest.raises(DocxTextReviewError, match="stale"):
        validate_docx_text_review(changed_task, _submission(task))


def test_fails_closed_for_unknown_word_text_part(review_docx: Path) -> None:
    with zipfile.ZipFile(review_docx, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/custom-visible.xml", '<w:custom xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>Неизвестный текст</w:t></w:r></w:p></w:custom>'.encode("utf-8"))
    with pytest.raises(DocxTextReviewError, match="custom-visible.xml"):
        prepare_docx_text_review(review_docx)


def test_detects_tampered_task_blocks(review_docx: Path) -> None:
    task = prepare_docx_text_review(review_docx)
    task["blocks"][0]["text_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="task_sha256 mismatch"):
        validate_docx_text_review(task, _submission(task))
