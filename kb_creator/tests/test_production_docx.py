from pathlib import Path
import re
import zipfile

import pytest
from lxml import etree as lxml_etree
from lxml import html as lxml_html

pytest.importorskip("docx")
pytest.importorskip("lxml")
pytest.importorskip("PIL")

from docx import Document
from docx.shared import Mm, RGBColor
from PIL import Image

from kb_creator.scripts.production_docx import (
    WORD_NS,
    BuildState,
    ROLE_STYLE_NAMES,
    _apply_xml_callouts,
    add_block,
    bookmark_name,
    build_state_for_profile,
    build_docx,
    build_docx_projection,
    patch_existing_docx,
    validate_document_config,
)
from kb_creator.scripts.build_viral_production import inspect_docx_semantics, inspect_docx_toc


def test_build_docx_preserves_core_blocks_and_images(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (120, 80), "white").save(assets / "sample.png")
    html_path = tmp_path / "sample.html"
    html_path.write_text(
        """<!doctype html><html><head><meta charset="utf-8"></head><body><main>
        <section class="production-landing" style="background-image: url(assets/sample.png);"></section>
        <nav id="toc"><h1>Оглавление</h1><ol>
        <li class="toc-level-1"><a href="#one">Первый раздел</a></li>
        <li class="toc-level-2"><a href="#one-a">Подраздел</a></li>
        <li class="toc-level-1"><a href="#two">Второй раздел</a></li>
        <li class="toc-level-2"><a href="#two-a">Ещё один подраздел</a></li>
        </ol></nav>
        <h2 id="one">Первый раздел</h2>
        <h3 id="one-a">Подраздел</h3>
        <p>Обычный <strong>жирный</strong> и <em>курсивный</em> текст.</p>
        <p>См. <a href="#two">Второй раздел</a>.</p>
        <figure class="artview-decision artview-width-25 artview-align-right"><img src="assets/sample.png"></figure>
        <table><thead><tr><th>Параметр</th><th>Значение</th></tr></thead><tbody><tr><td>СИЛ</td><td>50</td></tr></tbody></table>
        <h2 id="two">Второй раздел</h2>
        <h3 id="two-a">Ещё один подраздел</h3>
        <blockquote class="review-callout"><p>Заметки Хранителя</p><p>Выделенный текст.</p></blockquote>
        </main></body></html>""",
        encoding="utf-8",
    )
    output = tmp_path / "sample.docx"

    result = build_docx(html_path, output)

    assert result["status"] == "pass"
    assert result["config_sha256"] is None
    assert result["images"] == 2
    assert result["tables"] == 1
    assert result["toc_entries"] == 4
    document = Document(output)
    assert document.core_properties.title == "Вирус"
    assert document.core_properties.subject == "Сценарий Call of Cthulhu"
    assert document.core_properties.author == "NRI"
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert round(document.sections[0].page_width.mm) == 210
    assert round(document.sections[0].page_height.mm) == 297
    assert "Первый раздел" in [paragraph.text for paragraph in document.paragraphs]
    body = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Обычный"))
    assert any(run.text == "жирный" and run.bold for run in body.runs)
    assert any(run.text == "курсивный" and run.italic for run in body.runs)
    assert len(document.inline_shapes) == 2
    assert len(document.tables) == 2
    toc_table = document.tables[0]
    assert len(toc_table.columns) == 2
    first_column = [paragraph for paragraph in toc_table.cell(0, 0).paragraphs if paragraph.text]
    second_column = [paragraph for paragraph in toc_table.cell(0, 1).paragraphs if paragraph.text]
    assert [paragraph.text for paragraph in first_column] == [
        "1.\tПЕРВЫЙ РАЗДЕЛ",
        "2.\tПодраздел",
    ]
    assert [paragraph.text for paragraph in second_column] == [
        "3.\tВТОРОЙ РАЗДЕЛ",
        "4.\tЕщё один подраздел",
    ]
    assert first_column[1].paragraph_format.left_indent > first_column[0].paragraph_format.left_indent
    with output.open("rb") as stream:
        with zipfile.ZipFile(stream) as archive:
            toc = inspect_docx_toc(archive, expected_entries=4)
            semantics = inspect_docx_semantics(archive)
    assert toc["entries"] == 4
    assert toc["two_columns"] is True
    assert toc["hierarchy"] is True
    assert toc["hyperlinks"] == 4
    assert len(set(toc["hyperlink_targets"])) == 4
    assert semantics["hyperlinks"] == 5
    assert semantics["bookmarks"] == 4
    assert semantics["callout_paragraphs"] == 2

    projected_output = tmp_path / "projected.docx"
    projected = build_docx_projection(lxml_html.fromstring(html_path.read_text(encoding="utf-8")),
                                      tmp_path, projected_output)
    assert projected["status"] == "pass"
    assert not (tmp_path / ".in-memory-projection").exists()
    with zipfile.ZipFile(projected_output) as archive:
        projected_toc = inspect_docx_toc(archive, expected_entries=4)
        projected_semantics = inspect_docx_semantics(archive)
    assert projected_toc == toc
    assert projected_semantics == semantics
    projected_document = Document(projected_output)
    assert len(projected_document.inline_shapes) == len(document.inline_shapes)
    assert len(projected_document.tables) == len(document.tables)
    projected_body = next(paragraph for paragraph in projected_document.paragraphs if paragraph.text.startswith("Обычный"))
    assert [(run.text, run.bold, run.italic) for run in projected_body.runs] == [
        (run.text, run.bold, run.italic) for run in body.runs]


def test_build_docx_preserves_unnumbered_hierarchical_toc_metadata(tmp_path: Path) -> None:
    html_path = tmp_path / "hierarchical.html"
    html_path.write_text(
        """<html><head><meta charset="utf-8"></head><body><main>
        <nav id="toc" class="toc toc-hierarchical" data-toc-numbering="none" data-toc-layout="hierarchical" data-toc-font-size-pt="12.5"><h1>Оглавление</h1><ul class="toc-root">
        <li class="toc-group toc-level-1"><a href="#one">Первый раздел</a><ul class="toc-children">
        <li class="toc-level-2"><a href="#one-a">Первый подраздел</a></li>
        <li class="toc-level-2"><a href="#one-b">Второй подраздел</a></li></ul></li>
        <li class="toc-group toc-level-1"><a href="#two">Второй раздел</a><ul class="toc-children">
        <li class="toc-level-2"><a href="#two-a">Третий подраздел</a></li></ul></li>
        </ul></nav>
        <h2 id="one">Первый раздел</h2><h3 id="one-a">Первый подраздел</h3><h3 id="one-b">Второй подраздел</h3>
        <h2 id="two">Второй раздел</h2><h3 id="two-a">Третий подраздел</h3>
        </main></body></html>""",
        encoding="utf-8",
    )
    output = tmp_path / "hierarchical.docx"

    result = build_docx(html_path, output)

    assert result["toc_entries"] == 5
    document = Document(output)
    toc_title = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Оглавление")
    assert toc_title.paragraph_format.keep_together is True
    assert toc_title.paragraph_format.keep_with_next is True
    toc_table = document.tables[0]
    first_column = [paragraph for paragraph in toc_table.cell(0, 0).paragraphs if paragraph.text]
    second_column = [paragraph for paragraph in toc_table.cell(0, 1).paragraphs if paragraph.text]
    assert [paragraph.text for paragraph in first_column] == ["Первый раздел", "Первый подраздел", "Второй подраздел"]
    assert [paragraph.text for paragraph in second_column] == ["Второй раздел", "Третий подраздел"]
    assert all(not re.match(r"\d+\.\s", paragraph.text) for paragraph in first_column + second_column)
    assert first_column[0].paragraph_format.left_indent < first_column[1].paragraph_format.left_indent
    assert first_column[0].paragraph_format.keep_with_next is True
    assert first_column[1].paragraph_format.keep_with_next is True
    assert first_column[2].paragraph_format.keep_with_next is False
    assert '<w:b' in first_column[0]._p.xml
    assert '<w:b w:val="0"' not in first_column[1]._p.xml
    assert document.styles[ROLE_STYLE_NAMES["toc"]].font.size.pt == 12.5
    assert 'w:val="25"' not in first_column[1]._p.xml


def test_build_docx_preserves_explicit_flat_source_toc_case_and_peer_style(tmp_path: Path) -> None:
    html_path = tmp_path / "flat.html"
    html_path.write_text(
        """<html><head><meta charset="utf-8"></head><body><main>
        <nav id="toc" data-toc-layout="flat" data-toc-numbering="numbered">
        <h1>Содержание</h1><ol>
        <li class="toc-level-1"><a href="#intro">Введение</a></li>
        <li class="toc-level-1"><a href="#appendix">Приложение A — характеристики персонажей</a></li>
        </ol></nav>
        <h2 id="intro">Введение</h2><h2 id="appendix">Приложение A</h2>
        </main></body></html>""",
        encoding="utf-8",
    )
    output = tmp_path / "flat.docx"

    result = build_docx(html_path, output)

    assert result["toc_entries"] == 2
    paragraphs = [
        paragraph
        for cell in Document(output).tables[0].rows[0].cells
        for paragraph in cell.paragraphs
        if paragraph.text
    ]
    assert [paragraph.text for paragraph in paragraphs] == [
        "1.\tВведение",
        "2.\tПриложение A — характеристики персонажей",
    ]
    assert all(not any(run.bold for run in paragraph.runs) for paragraph in paragraphs)

    drifted = tmp_path / "flat-drifted.docx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(drifted, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace("Введение".encode(), "ВВЕДЕНИЕ".encode(), 1)
            target.writestr(info, payload)
    with pytest.raises(ValueError, match="flat TOC display text"):
        patch_existing_docx(html_path, drifted)


def test_build_docx_rejects_invalid_toc_font_metadata_before_output(tmp_path: Path) -> None:
    html_path = tmp_path / "invalid-toc.html"
    html_path.write_text(
        '<html><body><main><nav id="toc" data-toc-font-size-pt="huge"><ul>'
        '<li><a href="#one">One</a></li></ul></nav><h2 id="one">One</h2></main></body></html>',
        encoding="utf-8",
    )
    output = tmp_path / "invalid-toc.docx"

    with pytest.raises(ValueError, match="TOC font metadata"):
        build_docx(html_path, output)

    assert not output.exists()


def test_patch_existing_docx_preserves_user_document_and_adds_semantics(tmp_path: Path) -> None:
    html_path = tmp_path / "sample.html"
    html_path.write_text(
        """<!doctype html><html><head><meta charset="utf-8"></head><body><main>
        <nav id="toc"><h1>Оглавление</h1><ol>
        <li class="toc-level-1"><a href="#one">Первый раздел</a></li>
        <li class="toc-level-2"><a href="#one-a">Подраздел</a></li>
        <li class="toc-level-1"><a href="#two">Второй раздел</a></li>
        <li class="toc-level-2"><a href="#two-a">Ещё один подраздел</a></li>
        </ol></nav>
        <h2 id="one">Первый раздел</h2>
        <p>См. <a href="#two">Второй раздел</a>.</p>
        <h3 id="one-a">Подраздел</h3>
        <h2 id="two">Второй раздел</h2>
        <h3 id="two-a">Ещё один подраздел</h3>
        <blockquote class="review-callout"><p>Заметки Хранителя</p><p>Выделенный текст.</p></blockquote>
        </main></body></html>""",
        encoding="utf-8",
    )
    document = Document()
    document.add_paragraph("Оглавление", style="Title")
    table = document.add_table(rows=1, cols=2)
    for cell, rows in zip(
        table.rows[0].cells,
        (
            ("1.\tПЕРВЫЙ РАЗДЕЛ", "2.\tПодраздел"),
            ("3.\tВТОРОЙ РАЗДЕЛ", "4.\tЕщё один подраздел"),
        ),
    ):
        cell.paragraphs[0].text = rows[0]
        cell.add_paragraph(rows[1])
    document.add_paragraph("Первый раздел", style="Heading 2")
    link_paragraph = document.add_paragraph("См. ")
    link_run = link_paragraph.add_run("Второй раздел")
    link_run.underline = True
    link_run.font.color.rgb = RGBColor(5, 99, 193)
    link_paragraph.add_run(".")
    document.add_paragraph("Подраздел", style="Heading 3")
    document.add_paragraph("Второй раздел", style="Heading 2")
    document.add_paragraph("Ещё один подраздел", style="Heading 3")
    document.add_paragraph("Заметки Хранителя")
    document.add_paragraph("Выделенный текст.")
    docx_path = tmp_path / "edited.docx"
    document.save(docx_path)

    result = patch_existing_docx(html_path, docx_path)

    assert result["toc_hyperlinks"] == 4
    assert result["body_hyperlinks"] == 1
    assert result["bookmarks"] == 4
    assert result["callouts_applied"] == 1
    assert result["callouts_missing"] == []
    with zipfile.ZipFile(docx_path) as archive:
        semantics = inspect_docx_semantics(archive)
    assert semantics["hyperlinks"] == 5
    assert semantics["bookmarks"] == 4
    assert semantics["callout_paragraphs"] == 2


def test_build_docx_styles_callout_headings_and_nested_containers(tmp_path: Path) -> None:
    html_path = tmp_path / "mixed-callout.html"
    html_path.write_text(
        """<html><body><main>
        <nav id="toc"><h1>Содержание</h1><ol>
          <li class="toc-level-1"><a href="#section">Раздел</a></li>
        </ol></nav>
        <h2 id="section">Раздел</h2>
        <blockquote class="review-callout">
          <p>Вступительная строка.</p>
          <div><h3>Заголовок раздатки</h3><p>Строка после заголовка.</p></div>
          <ol><li>Первый пункт.</li><li>Второй пункт.</li></ol>
        </blockquote>
        </main></body></html>""",
        encoding="utf-8",
    )
    output = tmp_path / "mixed-callout.docx"

    result = build_docx(html_path, output)

    assert result["callouts"] == 1
    with zipfile.ZipFile(output) as archive:
        semantics = inspect_docx_semantics(archive)
    assert semantics["callout_paragraphs"] == 5


def test_apply_xml_callouts_fails_closed_when_heading_is_missing() -> None:
    root_html = lxml_html.fromstring(
        """<html><body><blockquote class="review-callout">
        <p>Вступительная строка.</p><h3>Заголовок раздатки</h3><p>Строка после заголовка.</p>
        </blockquote></body></html>"""
    )
    root_xml = lxml_etree.fromstring(
        f"""<w:document xmlns:w="{WORD_NS}"><w:body>
        <w:p><w:r><w:t>Вступительная строка.</w:t></w:r></w:p>
        <w:p><w:r><w:t>Строка после заголовка.</w:t></w:r></w:p>
        </w:body></w:document>"""
    )

    applied, missing = _apply_xml_callouts(root_html, root_xml)

    assert applied == 0
    assert missing == ["Вступительная строка. Заголовок раздатки Строка после заголовка."]


def test_add_blockquote_preserves_nested_figure_without_saving_docx(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (120, 80), "white").save(assets / "handout.png")
    html_path = tmp_path / "internal.html"
    blockquote = lxml_html.fromstring(
        """<blockquote class="review-callout"><p>Раздаточный материал</p>
        <figure><img src="assets/handout.png"></figure><p>Подпись.</p></blockquote>"""
    )
    document = Document()
    state = BuildState()

    add_block(document, blockquote, html_path, state)

    assert state.images == 1
    assert state.missing_images == []
    assert len(document.inline_shapes) == 1


def test_build_docx_matches_review_callout_with_figure_caption(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (120, 80), "white").save(assets / "handout.png")
    html_path = tmp_path / "callout-figure.html"
    html_path.write_text(
        """<html><body><main>
        <nav id="toc"><h1>Содержание</h1><ol>
          <li class="toc-level-1"><a href="#section">Раздел</a></li>
        </ol></nav>
        <h2 id="section">Раздел</h2>
        <blockquote class="review-callout">
          <p>Текст перед раздаткой.</p>
          <figure><img src="assets/handout.png"><figcaption>Подпись раздатки.</figcaption></figure>
          <p>Текст после раздатки.</p>
        </blockquote>
        </main></body></html>""",
        encoding="utf-8",
    )
    output = tmp_path / "callout-figure.docx"

    result = build_docx(html_path, output)

    assert result["callouts"] == 1
    with zipfile.ZipFile(output) as archive:
        semantics = inspect_docx_semantics(archive)
    # The callout styles its image, caption, and two text paragraphs; only the
    # three nonempty text paragraphs participate in the XML text matcher.
    assert semantics["callout_paragraphs"] == 4


def _layout_profile() -> dict[str, object]:
    return {
        "schema": "kb_creator.production_docx_config.v1",
        "page": {"width_mm": 100, "height_mm": 120, "orientation": "portrait"},
        "margins": {"top_mm": 20, "bottom_mm": 20, "left_mm": 20, "right_mm": 20},
        "fonts": {"body": "Times New Roman", "title": "Times New Roman", "headings": "Times New Roman", "footer": "Times New Roman"},
        "sizes": {"body_pt": 10.5, "title_pt": 22, "footer_pt": 9},
        "spacing": {"body_after_pt": 4, "body_line": 1.05},
        "headings": {
            str(level): {"size_pt": 10 + level, "bold": level < 4, "space_before_pt": 4, "space_after_pt": 4}
            for level in range(1, 7)
        },
        "footer": {"page_numbers": True},
        "metadata": {"title": "Test", "subject": "Test", "author": "NRI", "keywords": "", "comments": ""},
    }


def test_figure_policy_bounds_aspect_ratio_and_exact_override_without_saving_docx(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (200, 100), "white").save(assets / "wide.png")
    Image.new("RGB", (200, 400), "white").save(assets / "portrait.png")
    profile = _layout_profile()
    profile["figures"] = {
        "default": {"max_width_mm": 100, "max_height_mm": 100, "keep_with_next": False},
        "overrides": {"assets/portrait.png": {"max_width_mm": 50, "max_height_mm": 70, "keep_with_next": True}},
    }
    state = build_state_for_profile(validate_document_config(profile))
    document = Document()

    add_block(document, lxml_html.fromstring('<figure><img src="assets/wide.png"></figure>'), tmp_path / "internal.html", state)
    add_block(document, lxml_html.fromstring('<figure><img src="assets/portrait.png"></figure>'), tmp_path / "internal.html", state)

    wide, portrait = document.inline_shapes
    assert wide.width == Mm(60)
    assert wide.height == Mm(30)
    assert portrait.width == Mm(35)
    assert portrait.height == Mm(70)
    assert document.paragraphs[0].paragraph_format.keep_with_next is False
    assert document.paragraphs[1].paragraph_format.keep_with_next is True


@pytest.mark.parametrize(
    "figures",
    [
        {"default": {"max_width_mm": 80, "max_height_mm": 100}, "overrides": {}},
        {"default": {"max_width_mm": float("nan"), "max_height_mm": 100, "keep_with_next": False}, "overrides": {}},
        {"default": {"max_width_mm": 80, "max_height_mm": 100, "keep_with_next": False}, "overrides": {"../asset.png": {"max_width_mm": 80, "max_height_mm": 100, "keep_with_next": False}}},
        {"default": {"max_width_mm": 80, "max_height_mm": 100, "keep_with_next": False}, "overrides": {"https://example.test/asset.png": {"max_width_mm": 80, "max_height_mm": 100, "keep_with_next": False}}},
    ],
)
def test_figure_policy_rejects_incomplete_or_uncontained_overrides(figures: dict[str, object]) -> None:
    profile = _layout_profile()
    profile["figures"] = figures

    with pytest.raises(ValueError):
        validate_document_config(profile)


@pytest.mark.parametrize(
    "crop_fraction",
    [
        {"left": 1, "top": 0, "right": 0, "bottom": 0},
        {"left": 0.5, "top": 0, "right": 0.5, "bottom": 0},
        {"left": -0.1, "top": 0, "right": 0, "bottom": 0},
        {"left": 0, "top": 0, "right": float("nan"), "bottom": 0},
        {"left": 0, "top": 0, "right": 0},
    ],
)
def test_figure_policy_rejects_invalid_crop_fractions(crop_fraction: dict[str, object]) -> None:
    profile = _layout_profile()
    profile["figures"] = {
        "default": {
            "max_width_mm": 100, "max_height_mm": 200, "keep_with_next": False,
            "crop_fraction": crop_fraction,
        },
        "overrides": {},
    }

    with pytest.raises(ValueError, match="crop_fraction"):
        validate_document_config(profile)


def test_figure_crop_uses_native_src_rect_and_retains_original_asset(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    original = assets / "character-sheet.png"
    Image.new("RGB", (2048, 1198), "white").save(original)
    profile = _layout_profile()
    profile["figures"] = {
        "default": {
            "max_width_mm": 100, "max_height_mm": 200, "keep_with_next": False,
            "crop_fraction": {"left": 0, "top": 0, "right": 0.75, "bottom": 0},
        },
        "overrides": {},
    }
    state = build_state_for_profile(validate_document_config(profile))
    document = Document()

    add_block(document, lxml_html.fromstring('<figure><img src="assets/character-sheet.png"></figure>'), tmp_path / "internal.html", state)

    shape = document.inline_shapes[0]
    assert shape.width == Mm(80 * (2048 * 0.25) / 1198)
    assert shape.height == Mm(80)
    assert "a:srcRect" in shape._inline.xml
    assert 'r="75000"' in shape._inline.xml
    output = tmp_path / "cropped.docx"
    document.save(output)
    with zipfile.ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == 1
        assert archive.read(media[0]) == original.read_bytes()


def test_legacy_figure_build_state_preserves_existing_dimensions_without_saving_docx(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (100, 100), "white").save(assets / "legacy.png")
    document = Document()
    state = build_state_for_profile(validate_document_config(_layout_profile()))

    add_block(document, lxml_html.fromstring('<figure><img src="assets/legacy.png"></figure>'), tmp_path / "internal.html", state)

    assert state.figure_policy is None
    assert document.inline_shapes[0].width == Mm(176)
    assert document.inline_shapes[0].height == Mm(176)
    assert document.paragraphs[0].paragraph_format.keep_with_next is None


def test_captioned_figure_keeps_caption_chain_and_reserves_page_height(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (80, 160), "white").save(assets / "full-height.png")
    profile = _layout_profile()
    profile["figures"] = {
        "default": {"max_width_mm": 80, "max_height_mm": 80, "keep_with_next": False},
        "overrides": {},
    }
    html_path = tmp_path / "captioned-figure.html"
    html_path.write_text(
        "<html><head><meta charset='utf-8'></head><body><main><figure><img src='assets/full-height.png'/><figcaption>Первая подпись.</figcaption>"
        "<figcaption>Вторая подпись.</figcaption></figure><p>Следующий абзац.</p></main></body></html>",
        encoding="utf-8",
    )
    output = tmp_path / "captioned-figure.docx"

    build_docx(html_path, output, config=profile)
    document = Document(output)

    assert document.inline_shapes[0].height == Mm(65)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "", "Первая подпись.", "Вторая подпись.", "Следующий абзац.",
    ]
    assert [paragraph.paragraph_format.keep_with_next for paragraph in document.paragraphs] == [
        True, True, False, None,
    ]


def test_uncaptioned_figure_keeps_existing_policy_and_does_not_bind_following_prose(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (80, 160), "white").save(assets / "uncaptioned.png")
    profile = _layout_profile()
    profile["figures"] = {
        "default": {"max_width_mm": 80, "max_height_mm": 80, "keep_with_next": False},
        "overrides": {},
    }
    state = build_state_for_profile(validate_document_config(profile))
    document = Document()

    add_block(
        document,
        lxml_html.fromstring("<div><figure><img src='assets/uncaptioned.png'/></figure><p>Проза.</p></div>"),
        tmp_path / "internal.html",
        state,
    )

    assert document.inline_shapes[0].height == Mm(80)
    assert [paragraph.paragraph_format.keep_with_next for paragraph in document.paragraphs] == [False, None]


def test_statblock_policy_sets_group_keep_flags_without_saving_docx() -> None:
    profile = _layout_profile()
    profile["statblocks"] = {"keep_together": True}
    state = build_state_for_profile(validate_document_config(profile))
    document = Document()
    statblock = lxml_html.fromstring(
        """<section class="statblock-card"><h3>Статблок</h3><p>СИЛ 50</p><p>Навык 40%</p></section>"""
    )

    add_block(document, statblock, None, state)

    assert [paragraph.text for paragraph in document.paragraphs] == ["Статблок", "СИЛ 50", "Навык 40%"]
    assert [paragraph.style.name for paragraph in document.paragraphs[1:]] == [ROLE_STYLE_NAMES["statblock"]] * 2
    assert [paragraph.paragraph_format.keep_together for paragraph in document.paragraphs] == [True, True, True]
    assert [paragraph.paragraph_format.keep_with_next for paragraph in document.paragraphs] == [True, True, False]


def test_statblock_policy_requires_boolean_keep_together() -> None:
    profile = _layout_profile()
    profile["statblocks"] = {"keep_together": "yes"}

    with pytest.raises(ValueError):
        validate_document_config(profile)


def test_typography_roles_control_generated_styles_and_preserve_inline_emphasis(tmp_path: Path) -> None:
    profile = _layout_profile()
    profile["fonts"] = {"body": "Arial", "title": "Arial", "headings": "Arial", "footer": "Arial"}
    profile["typography"] = {
        role: {
            "font": "Courier New" if role == "toc" else "Arial",
            "size_pt": {"table_body": 10, "table_header": 10, "caption": 9, "statblock": 10.5, "toc": 12, "quote": 10.5}[role],
            "space_before_pt": 0,
            "space_after_pt": 2,
            "line_spacing": 1.1,
        }
        for role in ROLE_STYLE_NAMES
    }
    html_path = tmp_path / "roles.html"
    html_path.write_text(
        """<html><head><meta charset="utf-8"></head><body><main><nav id="toc" data-toc-font-size-pt="6"><h1>Содержание</h1><ol>
        <li><a href="#heading">Заголовок</a></li><li><a href="#text">Текст</a></li><li><a href="#picture">Рисунок</a></li><li><a href="#grid">Таблица</a></li>
        </ol></nav><h2 id="heading">Заголовок</h2><p id="text">Обычный <strong>жирный</strong> <em>курсивный</em>.</p>
        <figure id="picture"><img src="asset.png"/><figcaption>Подпись</figcaption></figure>
        <table id="grid"><tr><th>Шапка</th></tr><tr><td>Ячейка</td></tr></table></main></body></html>""",
        encoding="utf-8",
    )
    Image.new("RGB", (80, 40), "white").save(tmp_path / "asset.png")
    output = tmp_path / "roles.docx"

    build_docx(html_path, output, config=profile)

    document = Document(output)
    assert document.styles[ROLE_STYLE_NAMES["toc"]].font.name == "Courier New"
    assert document.styles[ROLE_STYLE_NAMES["toc"]].font.size.pt == 12
    assert document.styles[ROLE_STYLE_NAMES["table_body"]].font.size.pt == 10
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Заголовок")
    assert heading.style.name == "Heading 2"
    assert heading.style.font.bold is True
    assert heading.runs[0].bold is None
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', *ROLE_STYLE_NAMES.values()):
        rpr = document.styles[name].element.rPr
        assert not any('theme' in key.lower() for key in rpr.rFonts.attrib)
        assert rpr.find(f'{{{WORD_NS}}}szCs').get(f'{{{WORD_NS}}}val') == str(round(document.styles[name].font.size.pt * 2))
    body = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Обычный"))
    assert next(run for run in body.runs if run.text.startswith("Обычный")).bold is None
    assert next(run for run in body.runs if run.text == "жирный").bold is True
    assert next(run for run in body.runs if run.text == "курсивный").italic is True
    assert document.tables[1].cell(0, 0).paragraphs[0].style.name == ROLE_STYLE_NAMES["table_header"]
    assert document.tables[1].cell(1, 0).paragraphs[0].style.name == ROLE_STYLE_NAMES["table_body"]
    caption = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Подпись")
    assert caption.style.name == ROLE_STYLE_NAMES["caption"]
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    for target in ("heading", "text", "picture", "grid"):
        assert f'w:name="{bookmark_name(target)}"' in xml


def test_typography_rejects_partial_roles_and_duplicate_html_ids_before_output(tmp_path: Path) -> None:
    profile = _layout_profile()
    profile["typography"] = {"toc": {"font": "Arial", "size_pt": 10, "space_before_pt": 0, "space_after_pt": 0, "line_spacing": 1}}
    with pytest.raises(ValueError, match="DOCX typography"):
        validate_document_config(profile)

    html_path = tmp_path / "duplicate.html"
    html_path.write_text("<html><body><main><p id='same'>A</p><p id='same'>B</p></main></body></html>", encoding="utf-8")
    output = tmp_path / "duplicate.docx"
    with pytest.raises(ValueError, match="unique"):
        build_docx(html_path, output)
    assert not output.exists()
