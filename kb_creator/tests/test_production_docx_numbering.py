from pathlib import Path
import zipfile

import pytest
from lxml import etree

pytest.importorskip("docx")

from docx import Document

from kb_creator.scripts.production_docx import WORD_NS, build_docx


W = {"w": WORD_NS}


def _numbering(path: Path) -> tuple[dict[str, int], dict[str, str], str]:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/numbering.xml"))
    starts = {}
    abstracts = {}
    for num in root.xpath("./w:num", namespaces=W):
        num_id = num.get(f"{{{WORD_NS}}}numId")
        abstracts[num_id] = num.xpath("string(./w:abstractNumId/@w:val)", namespaces=W)
        override = num.xpath("string(./w:lvlOverride/w:startOverride/@w:val)", namespaces=W)
        if override:
            starts[num_id] = int(override)
    return starts, abstracts, etree.tostring(root, encoding="unicode")


def test_ordered_lists_have_independent_instances_and_explicit_values(tmp_path: Path) -> None:
    source = tmp_path / "lists.html"
    source.write_text(
        "<html><body><main>"
        "<ol><li>First one</li><li>First two</li></ol>"
        "<p>Intervening paragraph</p>"
        "<ol><li>Second one</li><li>Second two</li></ol>"
        "<ol start='7'><li>Seven</li><li>Eight</li><li value='11'>Eleven</li><li>Twelve</li></ol>"
        "<ol><li>Parent one<ol start='3'><li>Child three</li><li>Child four</li></ol></li>"
        "<li>Parent two</li></ol>"
        "</main></body></html>",
        encoding="utf-8",
    )
    output = tmp_path / "lists.docx"
    build_docx(source, output)

    paragraphs = {p.text: p for p in Document(output).paragraphs if p.text}
    starts, abstracts, numbering_xml = _numbering(output)

    def num_id(label: str) -> str:
        return paragraphs[label]._p.xpath("string(./w:pPr/w:numPr/w:numId/@w:val)")

    assert num_id("First one") == num_id("First two")
    assert num_id("Second one") == num_id("Second two")
    assert num_id("First one") != num_id("Second one")
    assert starts[num_id("First one")] == starts[num_id("Second one")] == 1
    assert num_id("Seven") == num_id("Eight")
    assert starts[num_id("Seven")] == 7
    assert num_id("Eleven") == num_id("Twelve")
    assert num_id("Eleven") != num_id("Eight")
    assert starts[num_id("Eleven")] == 11
    assert num_id("Parent one") == num_id("Parent two")
    assert num_id("Child three") == num_id("Child four")
    assert num_id("Parent one") != num_id("Child three")
    assert starts[num_id("Child three")] == 3
    assert len({abstracts[num_id(label)] for label in ("First one", "Second one", "Seven", "Eleven", "Parent one", "Child three")}) == 1
    assert '<w:suff w:val="space"' in numbering_xml
    assert paragraphs["Intervening paragraph"]._p.xpath("./w:pPr/w:numPr") == []


def test_long_ordered_list_keeps_one_numbering_instance(tmp_path: Path) -> None:
    source = tmp_path / "long-list.html"
    items = "".join(f"<li>Item {index}: " + ("long text " * 16) + "</li>" for index in range(1, 91))
    source.write_text(f"<html><body><main><ol>{items}</ol></main></body></html>", encoding="utf-8")
    output = tmp_path / "long-list.docx"
    build_docx(source, output)

    paragraphs = [p for p in Document(output).paragraphs if p.text.startswith("Item ")]
    ids = [p._p.xpath("string(./w:pPr/w:numPr/w:numId/@w:val)") for p in paragraphs]
    starts, _, _ = _numbering(output)
    assert len(paragraphs) == 90
    assert len(set(ids)) == 1
    assert starts[ids[0]] == 1


@pytest.mark.parametrize("markup", ["<ol start='next'><li>A</li></ol>", "<ol><li value='1.5'>A</li></ol>"])
def test_ordered_lists_reject_invalid_explicit_numbers(tmp_path: Path, markup: str) -> None:
    source = tmp_path / "invalid.html"
    source.write_text(f"<html><body><main>{markup}</main></body></html>", encoding="utf-8")
    output = tmp_path / "invalid.docx"
    with pytest.raises(ValueError, match="must be an integer"):
        build_docx(source, output)
    assert not output.exists()
