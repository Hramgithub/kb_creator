import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import html as lxml_html

from kb_creator.scripts.production_docx import BuildState, ROLE_STYLE_NAMES, add_block
from kb_creator.pipeline.five_step_docx import _default_config, _load_config


def _keep_next(paragraph) -> bool:
    p_pr = paragraph._p.find(qn("w:pPr"))
    keep = p_pr.find(qn("w:keepNext")) if p_pr is not None else None
    return keep is not None and keep.get(qn("w:val")) not in {"0", "false"}


def _keep_together_xml(paragraph_xml) -> bool:
    p_pr = paragraph_xml.find(qn("w:pPr"))
    return p_pr is not None and p_pr.find(qn("w:keepLines")) is not None


@pytest.mark.parametrize(
    ("source", "expected_keep_next", "expected_header_styles"),
    [
        (
            "<table data-table-header-policy='explicit' data-table-source-header-rows='1,2,4'>"
            "<tr data-table-source-row='1' data-table-header='true'><th>H1</th><th>H2</th></tr>"
            "<tr data-table-source-row='2' data-table-header='true'><th>H3</th><th>H4</th></tr>"
            "<tr data-table-source-row='3' data-table-header='false'><td>D1</td><td>D2</td></tr>"
            "<tr data-table-source-row='4' data-table-header='true'><th>Later</th><th>Header</th></tr></table>",
            [[True, True], [True, True], [False, False], [False, False]],
            [[True, True], [True, True], [False, False], [True, True]],
        ),
        (
            "<table data-table-header-policy='explicit' data-table-source-header-rows=''>"
            "<tr data-table-source-row='1' data-table-header='false'><td>D1</td></tr>"
            "<tr data-table-source-row='2' data-table-header='false'><td>D2</td></tr></table>",
            [[False], [False]],
            [[False], [False]],
        ),
        (
            "<table><tr><th>Only</th></tr></table>",
            [[False]],
            [[True]],
        ),
        (
            "<table><tr><td>Leading</td><td>Row</td></tr>"
            "<tr><th>Mixed</th><td>Data</td></tr><tr><td>Tail</td><td>Data</td></tr></table>",
            [[True, True], [False, False], [False, False]],
            [[True, True], [True, False], [False, False]],
        ),
    ],
)
def test_table_header_pagination_binds_only_leading_header_rows(
    source: str, expected_keep_next: list[list[bool]], expected_header_styles: list[list[bool]],
) -> None:
    document = Document()
    add_block(document, lxml_html.fromstring(source), None, BuildState())

    table = document.tables[0]
    assert [
        [_keep_next(cell.paragraphs[0]) for cell in row.cells]
        for row in table.rows
    ] == expected_keep_next
    assert [
        [cell.paragraphs[0].style.name == ROLE_STYLE_NAMES["table_header"] for cell in row.cells]
        for row in table.rows
    ] == expected_header_styles


def _body_paragraphs_in_order(document: Document):
    return [
        paragraph
        for child in document._element.body
        for paragraph in child.iter(qn("w:p"))
    ]


def _paragraph_text(paragraph_xml) -> str:
    return "".join(node.text or "" for node in paragraph_xml.iter(qn("w:t")))


def _keep_next_xml(paragraph_xml) -> bool:
    p_pr = paragraph_xml.find(qn("w:pPr"))
    keep = p_pr.find(qn("w:keepNext")) if p_pr is not None else None
    return keep is not None and keep.get(qn("w:val")) not in {"0", "false"}


@pytest.mark.parametrize(
    "card_markup, expected_texts",
    [
        (
            "<table data-table-header-policy='explicit' data-table-source-header-rows=''><tr data-table-source-row='1' data-table-header='false'><td>first-a</td></tr><tr data-table-source-row='2' data-table-header='false'><td>first-b</td></tr></table><p>first-after</p>",
            ["first-a", "first-b", "", "first-after"],
        ),
        (
            "<p>middle-before</p><table data-table-header-policy='explicit' data-table-source-header-rows=''><tr data-table-source-row='1' data-table-header='false'><td>middle-a</td></tr><tr data-table-source-row='2' data-table-header='false'><td>middle-b</td></tr></table><p>middle-after</p>",
            ["middle-before", "middle-a", "middle-b", "", "middle-after"],
        ),
        (
            "<p>last-before</p><table data-table-header-policy='explicit' data-table-source-header-rows=''><tr data-table-source-row='1' data-table-header='false'><td>last-a</td></tr><tr data-table-source-row='2' data-table-header='false'><td>last-b</td></tr></table>",
            ["last-before", "last-a", "last-b", ""],
        ),
    ],
)
def test_statblock_keep_together_chains_table_cells_without_leaking(
    card_markup: str, expected_texts: list[str],
) -> None:
    document = Document()
    source = lxml_html.fromstring(
        f"<div><section class='statblock-card'>{card_markup}</section><p>ordinary prose</p></div>"
    )
    add_block(document, source, None, BuildState(statblocks_keep_together=True))

    paragraphs = _body_paragraphs_in_order(document)
    assert [_paragraph_text(paragraph) for paragraph in paragraphs] == expected_texts + ["ordinary prose"]
    assert [_keep_together_xml(paragraph) for paragraph in paragraphs[:-1]] == [True] * len(expected_texts)
    assert [_keep_next_xml(paragraph) for paragraph in paragraphs] == [True] * (len(expected_texts) - 1) + [False, False]


def test_statblock_keep_together_false_leaves_table_and_card_flow_unbound() -> None:
    document = Document()
    source = lxml_html.fromstring(
        "<div><section class='statblock-card'><p>card-before</p>"
        "<table data-table-header-policy='explicit' data-table-source-header-rows=''><tr data-table-source-row='1' data-table-header='false'><td>cell</td></tr></table>"
        "<p>card-after</p></section><p>ordinary prose</p></div>"
    )
    add_block(document, source, None, BuildState(statblocks_keep_together=False))

    paragraphs = _body_paragraphs_in_order(document)
    assert [_paragraph_text(paragraph) for paragraph in paragraphs] == ["card-before", "cell", "", "card-after", "ordinary prose"]
    assert not any(_keep_together_xml(paragraph) for paragraph in paragraphs)
    assert not any(_keep_next_xml(paragraph) for paragraph in paragraphs)


def test_five_step_default_enables_statblock_keep_together() -> None:
    assert _default_config("Карточка")["statblocks"] == {"keep_together": True}


def test_five_step_caller_config_preserves_explicit_statblock_false(tmp_path: Path) -> None:
    config = _default_config("Caller title")
    config["statblocks"] = {"keep_together": False}
    source = tmp_path / "docx-config.json"
    source.write_text(json.dumps(config), encoding="utf-8")

    loaded, identity = _load_config(source, "Requested title", tmp_path)

    assert loaded["statblocks"] == {"keep_together": False}
    assert loaded["metadata"]["title"] == "Requested title"
    assert identity is not None
