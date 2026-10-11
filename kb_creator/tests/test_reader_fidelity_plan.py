"""Prospective source-to-DOCX plan closure; no writer is invoked here."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from pypdf import PdfWriter

from kb_creator.pipeline.identity import file_identity
from kb_creator.pipeline.coc_glossary import HEADER
from kb_creator.pipeline.manifests import sealed
from kb_creator.pipeline.reader_fidelity import (PLAN_SCHEMA, ReaderFidelityError,
    project_text_operations, validate_plan, verify_word_graphics, verify_word_text)
from kb_creator.tests.test_edition_projection_contract import _selector
from kb_creator.pipeline.source_graphics import INVENTORY_SCHEMA, build_inventory
from kb_creator.tests.test_source_graphics_contract import _graphic, _paint, _review


def _write(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def _plan(tmp_path: Path, *, include_graphic: bool = False) -> tuple[Path, dict]:
    renderer = shutil.which("pdftoppm")
    if renderer is None:
        pytest.skip("pdftoppm runtime is unavailable")
    source = tmp_path / "synthetic.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=288, height=360)
    with source.open("wb") as stream:
        writer.write(stream)
    inventory_path = build_inventory(source, Path(renderer), tmp_path / "inventory")
    inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    if include_graphic:
        inventory["paint_occurrences"] = [_paint("use-1")]
        _write(inventory_path, sealed(INVENTORY_SCHEMA,
                                      {k: v for k, v in inventory.items() if k not in {"schema", "inventory_sha256"}},
                                      digest_key="inventory_sha256"))
        inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    if include_graphic:
        graphic, region = _graphic(tmp_path, "use-1")
        review_path, review = _review(tmp_path, inventory_path, inventory, [graphic], [region])
        placement = {"placement_id": "place-scene-1", "graphic_id": "scene-1",
                     "story_id": "story-one", "anchor": "scene-heading", "position": "after",
                     "repeat_ordinal": 0, "selected_sha256": graphic["selected"]["sha256"],
                     "representation": "word_drawing", "master_relation": "absent_from_master",
                     "route": "reviewed_reader_insertion", "crop": {"left": 0, "top": 0,
                     "right": 0, "bottom": 0, "approved": True},
                     "fit": {"width_emu": 1000, "height_emu": 1000},
                     "decision_reference": "source scene review"}
    else:
        review_path, review = _review(tmp_path, inventory_path, inventory)
        placement = None
    files = {}
    for key in ("master_ru", "projection", "config", "glossary_registry", "glossary_overlay"):
        path = tmp_path / f"{key}.txt"
        if key == "glossary_registry":
            path.write_text(",".join(HEADER) + "\n", encoding="utf-8")
        elif key == "glossary_overlay":
            path.write_text('{"terms":[]}', encoding="utf-8")
        else:
            path.write_text(key, encoding="utf-8")
        files[key] = {"path": str(path.resolve()), **file_identity(path)}
    value = sealed(PLAN_SCHEMA, {
        "inventory": {"path": str(inventory_path.resolve()), **file_identity(inventory_path)},
        "review": {"path": str(review_path.resolve()), **file_identity(review_path)},
        "source_pdf": inventory["source_pdf"], **files,
        "policy": {"target_edition": "7e", "edition_applicability": "applicable",
                   "book_id": "synthetic",
                   "rules_version": "synthetic-v1", "reader_profile": "five_step_docx_v3",
                   "review_reference": "synthetic coverage"},
        "placements": [placement] if placement else [], "edition_domain": [], "edition_rows": [],
        "terminology_rows": [], "section_domain": ["story-one"],
        "section_reviews": [{"section_id": "story-one", "cards": 0, "prose": 0,
                             "pregens": 0, "checks": 0, "labels": 0, "reviewed": True,
                             "decision_reference": "manual full-section review", "residual_issues": []}],
        "state": "closed",
    }, digest_key="plan_sha256")
    path = tmp_path / "READER_FIDELITY_PLAN.json"
    _write(path, value)
    return path, value


def _reseal(path: Path, value: dict) -> None:
    _write(path, sealed(PLAN_SCHEMA, {key: row for key, row in value.items()
                                      if key not in {"schema", "plan_sha256"}},
                        digest_key="plan_sha256"))


def test_empty_graphic_domain_still_requires_complete_page_and_section_reviews(tmp_path: Path):
    path, value = _plan(tmp_path)
    assert validate_plan(path, source_pdf=tmp_path / "synthetic.pdf")["state"] == "closed"
    value["section_reviews"] = []
    _reseal(path, value)
    with pytest.raises(ReaderFidelityError, match="every source section"):
        validate_plan(path)


def test_reviewed_scene_missing_from_master_needs_exact_insertion_plan(tmp_path: Path):
    path, value = _plan(tmp_path, include_graphic=True)
    assert validate_plan(path, source_pdf=tmp_path / "synthetic.pdf")["placements"][0]["graphic_id"] == "scene-1"
    value["placements"] = []
    _reseal(path, value)
    with pytest.raises(ReaderFidelityError, match="every reviewed inclusion"):
        validate_plan(path)


@pytest.mark.parametrize("change,match", [
    ({"story_id": "wrong-story"}, "unresolved"),
    ({"route": "retain_exact"}, "explicit reader insertion"),
    ({"fit": {"width_emu": 0, "height_emu": 1000}}, "zero or unsupported extent"),
    ({"selected_sha256": "B" * 64}, "unresolved"),
])
def test_stale_or_wrong_graphic_placement_fails(tmp_path: Path, change: dict, match: str):
    path, value = _plan(tmp_path, include_graphic=True)
    value["placements"][0].update(change)
    _reseal(path, value)
    with pytest.raises(ReaderFidelityError, match=match):
        validate_plan(path)


def test_stale_review_or_master_identity_fails(tmp_path: Path):
    path, value = _plan(tmp_path)
    value["master_ru"]["sha256"] = "B" * 64
    _reseal(path, value)
    with pytest.raises(ReaderFidelityError, match="MASTER_RU identity mismatch"):
        validate_plan(path)


def _docx(tmp_path: Path, *, picture_before_story: bool = False,
          extra_picture: bool = False, wrong_extent: bool = False,
          unpaired_bookmark: bool = False, header_picture: bool = False,
          inline_before_anchor: bool = False) -> Path:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu
    from kb_creator.scripts.production_docx import bookmark_name

    def mark(paragraph, anchor, ident):
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(ident))
        start.set(qn("w:name"), bookmark_name(anchor))
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(ident))
        paragraph._p.append(start)
        paragraph._p.append(end)

    doc = Document()
    if picture_before_story:
        image_para = doc.add_paragraph()
    section = doc.add_paragraph("Story")
    mark(section, "story-one", 1)
    anchor = doc.add_paragraph("Scene")
    mark(anchor, "scene-heading", 2)
    if unpaired_bookmark:
        anchor._p.remove(anchor._p[-1])
    if inline_before_anchor:
        image_para = anchor
    elif not picture_before_story:
        image_para = doc.add_paragraph()
    inline = image_para.add_run().add_picture(str(tmp_path / "scene-1.png"),
                                               width=Emu(1000), height=Emu(1000))
    inline._inline.docPr.set("descr", "place-scene-1")
    if inline_before_anchor:
        run = image_para.runs[-1]._r
        image_para._p.remove(run)
        image_para._p.insert(0, run)
    if wrong_extent:
        inline._inline.extent.set("cx", "0")
    if extra_picture:
        doc.add_paragraph().add_run().add_picture(str(tmp_path / "scene-1.png"),
                                                  width=Emu(1000), height=Emu(1000))
    if header_picture:
        doc.sections[0].header.add_paragraph().add_run().add_picture(
            str(tmp_path / "scene-1.png"), width=Emu(1000), height=Emu(1000))
    path = tmp_path / "actual.docx"
    doc.save(path)
    return path


def test_word_observer_matches_referenced_media_owner_and_extent(tmp_path: Path):
    plan_path, _ = _plan(tmp_path, include_graphic=True)
    docx = _docx(tmp_path)
    rows = verify_word_graphics(plan_path, docx)
    assert len(rows) == 1 and rows[0]["placement_id"] == "place-scene-1"


@pytest.mark.parametrize("variant,match", [
    ({"picture_before_story": True}, "wrong story/anchor"),
    ({"extra_picture": True}, "orphaned"),
    ({"wrong_extent": True}, "extent differs"),
    ({"unpaired_bookmark": True}, "paired same-paragraph bookmark"),
    ({"header_picture": True}, "unreviewed non-body graphic story"),
    ({"inline_before_anchor": True}, "wrong side of inline anchor"),
])
def test_word_observer_rejects_wrong_story_orphan_and_zero_extent(tmp_path: Path,
                                                                   variant: dict, match: str):
    plan_path, _ = _plan(tmp_path, include_graphic=True)
    docx = _docx(tmp_path, **variant)
    with pytest.raises(ReaderFidelityError, match=match):
        verify_word_graphics(plan_path, docx)


def test_approved_reader_only_stat_value_is_projected_and_observed_in_actual_word(tmp_path: Path):
    from docx import Document
    from lxml import html

    plan_path, plan = _plan(tmp_path)
    plan["edition_rows"] = [{"occurrence_id": "stat-1", "book_id": "synthetic",
        "story_id": "story-one", "entity_id": "keeper", "source_page": 1,
        "source_pdf_sha256": plan["source_pdf"]["sha256"],
        "source_edition": "legacy_proven", "present_scale": "legacy",
        "entity_kind": "human_npc", "field": "STR", "before": 13, "after": 65,
        "action": "convert_deterministic", "rule_id": "coc7-characteristic-v1",
        "decision_reference": "synthetic source review",
        "selector": _selector("STR 13", "STR 65", before_value="13", after_value="65")}]
    plan["edition_domain"] = [{"occurrence_id": "stat-1", "book_id": "synthetic",
        "story_id": "story-one", "entity_id": "keeper", "source_page": 1,
        "field": "STR", "source_render_sha256": json.loads(
            (tmp_path / "inventory/SOURCE_GRAPHICS_INVENTORY.json").read_text(encoding="utf-8"))["pages"][0]["render"]["sha256"],
        "decision_reference": "synthetic source review"}]
    _reseal(plan_path, plan)
    tree = html.fromstring('<html><body><h1 id="story-one">Story</h1><p id="field-1">STR 13</p></body></html>')
    assert project_text_operations(plan_path, tree) == 1
    assert tree.xpath('//*[@id="field-1"]')[0].text_content() == "STR 65"
    hidden_tree = html.fromstring('<html><body><h1 id="story-one">Story</h1><p id="field-1">STR <span style="display:none">13</span></p></body></html>')
    with pytest.raises(ReaderFidelityError, match="hidden or unsupported inline"):
        project_text_operations(plan_path, hidden_tree)
    fragmented_tree = html.fromstring('<html><body><h1 id="story-one">Story</h1><p id="field-1">STR 1<b>3</b></p></body></html>')
    with pytest.raises(ReaderFidelityError, match="crosses styled runs"):
        project_text_operations(plan_path, fragmented_tree)
    doc = Document()
    section = doc.add_paragraph("Story")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from kb_creator.scripts.production_docx import bookmark_name
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), bookmark_name("story-one"))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    section._p.append(start)
    section._p.append(end)
    doc.add_paragraph("STR 65")
    plan["edition_rows"][0]["selector"]["paragraph"] = 1
    _reseal(plan_path, plan)
    output = tmp_path / "stat.docx"
    doc.save(output)
    assert verify_word_text(plan_path, output)[0]["after_sha256"] == plan["edition_rows"][0]["selector"]["after_sha256"]
    doc.paragraphs[1].text = "STR 13"
    doc.save(output)
    with pytest.raises(ReaderFidelityError, match="actual Word value"):
        verify_word_text(plan_path, output)
    paragraph = doc.paragraphs[1]
    paragraph.clear()
    paragraph.add_run("STR 6")
    paragraph.add_run("5")
    doc.save(output)
    with pytest.raises(ReaderFidelityError, match="crosses run boundaries"):
        verify_word_text(plan_path, output)
    paragraph.clear()
    paragraph.add_run("STR 65").font.hidden = True
    doc.save(output)
    with pytest.raises(ReaderFidelityError, match="became hidden"):
        verify_word_text(plan_path, output)
    paragraph.clear()
    paragraph.add_run("STR 65")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    paragraph.runs[0]._r.append(field)
    doc.save(output)
    with pytest.raises(ReaderFidelityError, match="unsupported inline"):
        verify_word_text(plan_path, output)


def test_accepted_book_overlay_label_projects_and_reopens_word(tmp_path: Path):
    from docx import Document
    from lxml import html

    plan_path, plan = _plan(tmp_path)
    overlay = Path(plan["glossary_overlay"]["path"])
    overlay.write_text(json.dumps({"terms": [{"source": "Know", "target": "Знание"}]},
                                  ensure_ascii=False), encoding="utf-8")
    plan["glossary_overlay"] = {"path": str(overlay.resolve()), **file_identity(overlay)}
    plan["terminology_rows"] = [{"occurrence_id": "skill-1", "entity_id": "keeper",
        "source_label": "Know", "source_context": "general", "target_label": "Знание",
        "percentage": 70, "registry_sha256": plan["glossary_registry"]["sha256"],
        "overlay_sha256": plan["glossary_overlay"]["sha256"],
        "disposition": "accepted_direct_mapping", "authority_reference": "overlay:Know",
        "collision_decision": None,
        "selector": _selector("Know", "Знание", before_value="Know", after_value="Знание")}]
    _reseal(plan_path, plan)
    assert validate_plan(plan_path)["terminology_rows"][0]["target_label"] == "Знание"
    tree = html.fromstring('<html><body><p id="field-1">Know</p></body></html>')
    assert project_text_operations(plan_path, tree) == 1
    assert tree.xpath('//*[@id="field-1"]')[0].text_content() == "Знание"
    word = Document()
    word.add_paragraph("Знание")
    output = tmp_path / "accepted-label.docx"
    word.save(output)
    assert verify_word_text(plan_path, output)[0]["kind"] == "terminology"
