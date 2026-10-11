"""Explicit label ownership and DOCX pagination without heading promotion."""
from pathlib import Path
import copy
import re

import pytest
from docx import Document
from lxml import html
from translator.translator import markdown_to_review_html

from kb_creator.pipeline import reader_projection
from kb_creator.scripts.production_docx import build_docx
from kb_creator.tests.test_reader_whole_artifact import _unit


def fixture(successor="Following body."):
    texts = ["<!-- source-page-018 -->", "Label **with emphasis**", successor, "Ordinary prose."]
    units = [_unit(f"label-{i}", text) for i, text in enumerate(texts)]
    bind = lambda i: {key: units[i][key] for key in ("unit_id", "ru_sha256")}
    role = {"role": "semantic_label", "unit": bind(1), "next_unit": bind(2),
            "source_evidence": "Inspected source label and retained successor.",
            "reason": "Keep the semantic label with its content."}
    options = {"source_toc": {"state": "absent", "source_pdf_sha256": "A" * 64, "reason": "fixture"},
               "body_links": [], "asset_choices": [], "hit_points": {"presentation": "numeric"},
               "display_roles": [role]}
    return {"units": units}, options


def project(projection, options, *, merge_lists=False):
    source = "\n\n".join(unit["ru_text"] for unit in projection["units"])
    body, toc, _ = markdown_to_review_html(source, source_path=Path("fixture.md"),
                                         html_path=Path("fixture.html"), blockquote_heading_offset=0)
    if merge_lists:
        body = re.sub(r"</ul>\s*<ul>", "", body)
    return reader_projection.project_reader_html(body, toc, projection, options)


@pytest.mark.parametrize("successor", ["Following body.", "> First verse.\n>\n> Second verse.",
                                       "- First item\n- Second item"])
def test_label_keeps_plain_style_text_and_inline_markup_with_supported_successor(tmp_path, successor):
    projection, options = fixture(successor)
    body, toc, report = project(projection, options)
    tree = html.fragment_fromstring(body, create_parent="main")
    label = tree.xpath('.//p[@data-semantic-kind="label"]')[0]
    assert label.text_content() == "Label with emphasis"
    assert label.xpath('./strong')[0].text == "with emphasis"
    assert label.get("id") is None
    assert toc == []
    assert report["display_roles"] == options["display_roles"]
    path = tmp_path / "label.html"
    path.write_text('<html><head><meta charset="utf-8"></head><body><main>' + body + '</main></body></html>', encoding="utf-8")
    build_docx(path, tmp_path / "label.docx")
    doc = Document(tmp_path / "label.docx")
    paragraph = next(p for p in doc.paragraphs if p.text == "Label with emphasis")
    assert paragraph.style.name == "Normal"
    assert paragraph.paragraph_format.keep_with_next is True
    assert paragraph.paragraph_format.keep_together is True
    ordinary = next(p for p in doc.paragraphs if p.text == "Ordinary prose.")
    assert ordinary.paragraph_format.keep_with_next is None
    assert next(r for r in paragraph.runs if r.text == "with emphasis").bold


@pytest.mark.parametrize("change, message", [
    (lambda p, o: o["display_roles"][0]["unit"].update(ru_sha256="0" * 64), "preimage"),
    (lambda p, o: o["display_roles"][0]["next_unit"].update(ru_sha256="0" * 64), "preimage"),
    (lambda p, o: o["display_roles"][0].update(source_evidence=""), "reviewed fields"),
    (lambda p, o: o["display_roles"][0].update(extra=True), "reviewed fields"),
    (lambda p, o: o["display_roles"][0].update(next_unit=o["display_roles"][0]["unit"]), "distinct whole"),
    (lambda p, o: p["units"][1].update(parent_unit_id="parent"), "unsplit"),
    (lambda p, o: p["units"].append(_unit("child", "Unused child.")) or p["units"][-1].update(parent_unit_id="label-2"), "unsplit"),
    (lambda p, o: o["display_roles"].append(copy.deepcopy(o["display_roles"][0])), "conflicts"),
    (lambda p, o: p["units"].insert(2, _unit("separator", "Intervening retained prose.")), "adjacent"),
    (lambda p, o: p["units"].insert(2, _unit("page", "<!-- source-page-019 -->")), "same page"),
])
def test_stale_ambiguous_nonadjacent_or_split_decisions_fail_closed(change, message):
    projection, options = fixture()
    change(projection, options)
    with pytest.raises(ValueError, match=message):
        project(projection, options)


@pytest.mark.parametrize("successor", ["## Heading", "![Figure](figure.png)",
                                       "| Key | Value |\n| --- | --- |\n| One | Two |"])
def test_heading_or_image_cannot_be_claimed_as_label_content(successor):
    projection, options = fixture(successor)
    with pytest.raises(ValueError, match="retained text successor"):
        project(projection, options)


def test_plain_body_is_not_automatically_bound(tmp_path):
    projection, options = fixture()
    options["display_roles"] = []
    body, _, _ = project(projection, options)
    assert 'data-keep-with-next' not in body
    path = tmp_path / "body.html"
    path.write_text('<html><head><meta charset="utf-8"></head><body><main>' + body + '</main></body></html>', encoding="utf-8")
    build_docx(path, tmp_path / "body.docx")
    assert all(p.paragraph_format.keep_with_next is None for p in Document(tmp_path / "body.docx").paragraphs)


def test_duplicate_labels_bind_complete_source_cohort_and_exact_successors():
    projection, options = fixture()
    label = _unit("second-label", projection["units"][1]["ru_text"])
    following = _unit("second-body", "Second content.")
    projection["units"][3:3] = [label, following]
    second = copy.deepcopy(options["display_roles"][0])
    second["unit"] = {key: label[key] for key in ("unit_id", "ru_sha256")}
    second["next_unit"] = {key: following[key] for key in ("unit_id", "ru_sha256")}
    options["display_roles"].append(second)
    body, _, _ = project(projection, options)
    labels = html.fragment_fromstring(body, create_parent="main").xpath('.//p[@data-semantic-kind="label"]')
    assert [node.getnext().text_content() for node in labels] == ["Following body.", "Second content."]
    options["display_roles"][0]["next_unit"] = second["next_unit"]
    options["display_roles"] = options["display_roles"][:1]
    with pytest.raises(ValueError, match="adjacent"):
        project(projection, options)


def test_retained_cohort_count_mismatch_fails_before_guessing():
    projection, options = fixture()
    source = "\n\n".join(unit["ru_text"] for unit in projection["units"])
    body, toc, _ = markdown_to_review_html(source, source_path=Path("fixture.md"), html_path=Path("fixture.html"))
    body = body.replace('<p>Ordinary prose.</p>', '<p>Label <strong>with emphasis</strong></p>')
    with pytest.raises(ValueError, match="source cohort"):
        reader_projection.project_reader_html(body, toc, projection, options)


def handout_list_fixture():
    texts = ["<!-- source-page-030 -->", "They stand at", "## Strange occurrences",
             "- First sighting.", "- Second sighting.", "- Third sighting.",
             "<!-- source-page-032 -->", "the station."]
    units = [_unit(f"handout-{i}", text) for i, text in enumerate(texts)]
    bind = lambda i: {key: units[i][key] for key in ("unit_id", "ru_sha256")}
    _, options = fixture()
    options["display_roles"] = [{"role": "interrupted_prose_reflow", "units": [bind(1), bind(7)],
        "intervening_units": [bind(i) for i in range(2, 7)], "join_mode": "space",
        "source_evidence": "Reviewed exact narrative and full handout list.", "reason": "Restore sentence.",
        "intervening_handout_groups": [{"units": [bind(i) for i in range(2, 6)], "title": bind(2),
             "source_evidence": "Reviewed title and all list items.", "reason": "Retain complete list."}]}]
    return {"units": units}, options


def test_retained_handout_full_list_preserves_order_text_and_list_structure():
    projection, options = handout_list_fixture()
    body, _, report = project(projection, options, merge_lists=True)
    tree = html.fragment_fromstring(body, create_parent="main")
    assert tree.xpath('.//p')[0].text_content() == "They stand at the station."
    assert [li.text_content() for li in tree.xpath('.//ul/li')] == ["First sighting.", "Second sighting.", "Third sighting."]
    assert len(tree.xpath('.//ul')) == 1
    assert report["display_roles"] == options["display_roles"]


def test_partial_handout_list_ownership_fails_closed():
    projection, options = handout_list_fixture()
    options["display_roles"][0]["intervening_handout_groups"][0]["units"].pop()
    with pytest.raises(ValueError, match="full flat list ownership"):
        project(projection, options, merge_lists=True)


def test_whole_list_unit_is_retained_as_one_owner():
    projection, options = handout_list_fixture()
    combined = _unit("combined-list", "\n".join(u["ru_text"] for u in projection["units"][3:6]))
    projection["units"][3:6] = [combined]
    binding = {key: combined[key] for key in ("unit_id", "ru_sha256")}
    role = options["display_roles"][0]
    role["intervening_units"][1:4] = [binding]
    role["intervening_handout_groups"][0]["units"][1:] = [binding]
    body, _, _ = project(projection, options)
    assert len(html.fragment_fromstring(body, create_parent="main").xpath('.//ul/li')) == 3
