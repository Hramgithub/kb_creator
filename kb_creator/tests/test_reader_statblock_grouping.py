import copy
import hashlib

import pytest
from docx import Document
from lxml import etree, html

from kb_creator.pipeline import reader_statblock_grouping as grouping
from kb_creator.pipeline.reader_projection import project_reader_html
from kb_creator.pipeline.five_step_docx import _default_config
from kb_creator.scripts.production_docx import (
    add_block, build_state_for_profile, configure_document, finalize_reader_statblock_groups,
)


def fixture(cross_page=True):
    bodies = ['<h2 id="npc">NPC <em>A</em></h2><p id="stats">HP: <strong>12</strong></p>',
              '<p id="last">Dodge 40%</p>']
    if not cross_page:
        bodies = [''.join(bodies)]
    texts = [('npc', '## NPC *A*'), ('stats', 'HP: **12**'), ('last', 'Dodge 40%')]
    units, pages = [], {}
    for i, (uid, text) in enumerate(texts):
        page = 2 if cross_page and i == 2 else 1
        if not units or pages[units[-1]['unit_id']] != page:
            marker = f'<!-- source-page-{page:03d} -->'
            units.append({'unit_id': f'm{page}', 'ru_text': marker,
                          'ru_sha256': hashlib.sha256(marker.encode()).hexdigest().upper()})
            pages[f'm{page}'] = page
        units.append({'unit_id': uid, 'ru_text': text,
                      'ru_sha256': hashlib.sha256(text.encode()).hexdigest().upper()})
        pages[uid] = page
    tree = html.fragment_fromstring(''.join(f'<section data-page="{p:03d}">{b}</section>'
                                          for p, b in enumerate(bodies, 1)), create_parent='main')
    nodes = tree.xpath('.//section/*')
    by_id = {u['unit_id']: u for u in units}
    role = {'role': grouping.ROLE, 'group_id': 'npc-a',
            'units': [{k: by_id[uid][k] for k in ('unit_id', 'ru_sha256')} for uid, _ in texts],
            'rendered_members': [{'unit_id': uid, 'source_page': pages[uid],
                                  'final_html_sha256': grouping.fingerprint(node)}
                                 for (uid, _), node in zip(texts, nodes)],
            'source_evidence': 'Reviewed finite source record', 'reason': 'Keep finite card'}
    return tree, {'units': units}, pages, role


def apply(tree, projection, pages, roles):
    by_id = {u['unit_id']: u for u in projection['units']}
    def bound(b):
        u = by_id[b['unit_id']]
        if b['ru_sha256'] != u['ru_sha256']:
            raise ValueError('stale unit')
        return u
    report = {'display_roles': []}
    grouping.apply(tree, projection, {'display_roles': roles}, pages, bound, report)
    return report


def build_memory(tree, keep=True):
    profile = _default_config('Fixture')
    profile['statblocks']['keep_together'] = keep
    document = Document()
    configure_document(document, profile)
    state = build_state_for_profile(profile)
    for n in tree:
        add_block(document, n, None, state)
    finalize_reader_statblock_groups(document, state)
    return document


@pytest.mark.parametrize('cross_page', [False, True])
@pytest.mark.parametrize('keep', [False, True])
def test_exact_group_preserves_containers_inline_and_word_payload(cross_page, keep):
    tree, projection, pages, role = fixture(cross_page)
    baseline = build_memory(tree, keep)
    before = etree.tostring(tree)
    report = apply(tree, projection, pages, [role])
    document = build_memory(tree, keep)
    grouping.verify_xml(report, document._element.xml.encode(), keep_together=keep)
    assert [p.text for p in document.paragraphs] == [p.text for p in baseline.paragraphs]
    assert [etree.tostring(r) for p in document.paragraphs for r in p._p.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')] == [etree.tostring(r) for p in baseline.paragraphs for r in p._p.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')]
    assert [p.style.name for p in document.paragraphs] == [p.style.name for p in baseline.paragraphs]
    assert document.paragraphs[0].runs[-1].italic
    assert document.paragraphs[1].runs[-1].bold
    if keep:
        assert [p.paragraph_format.keep_with_next for p in document.paragraphs] == [True, True, False]
        assert all(p.paragraph_format.keep_together for p in document.paragraphs)
    else:
        assert [p.paragraph_format.keep_with_next for p in document.paragraphs] == [None, None, None]
    for n in tree.xpath(f'.//*[@{grouping.ATTRIBUTE}]'):
        del n.attrib[grouping.ATTRIBUTE]
    assert etree.tostring(tree) == before


@pytest.mark.parametrize('mutation', ['stale-unit', 'stale-html', 'reorder', 'overlap', 'wrong-page', 'partial', 'gap', 'nested'])
def test_preflight_rejects_drift_before_any_mutation(mutation):
    tree, projection, pages, role = fixture()
    roles = [role]
    if mutation == 'stale-unit': role['units'][0]['ru_sha256'] = 'F' * 64
    if mutation == 'stale-html': role['rendered_members'][0]['final_html_sha256'] = 'F' * 64
    if mutation == 'reorder': role['units'].reverse(); role['rendered_members'].reverse()
    if mutation == 'overlap':
        other = copy.deepcopy(role); other['group_id'] = 'npc-b'; roles.append(other)
    if mutation == 'wrong-page': role['rendered_members'][0]['source_page'] = 2
    if mutation == 'partial': projection['units'][1]['parent_unit_id'] = 'parent'
    if mutation == 'gap': tree[0].append(html.fromstring('<figure><img src="map.png"></figure>'))
    if mutation == 'nested': tree[0].set('data-semantic-kind', 'callout')
    before = etree.tostring(tree)
    with pytest.raises(ValueError): apply(tree, projection, pages, roles)
    assert etree.tostring(tree) == before


def test_actual_word_verifier_rejects_missing_flags_and_intervening_member():
    tree, projection, pages, role = fixture()
    report = apply(tree, projection, pages, [role])
    doc = build_memory(tree)
    doc.paragraphs[1].paragraph_format.keep_with_next = False
    with pytest.raises(ValueError, match='keep flags'):
        grouping.verify_xml(report, doc._element.xml.encode(), keep_together=True)
    doc = build_memory(tree)
    gap = doc.add_paragraph('Unowned gap')._p
    doc.paragraphs[0]._p.addnext(gap)
    with pytest.raises(ValueError, match='consecutive'):
        grouping.verify_xml(report, doc._element.xml.encode(), keep_together=True)


def test_deferred_group_composes_with_prose_presentation_and_review_anchor():
    tree, projection, pages, role = fixture(False)
    projection['units'][2]['ru_text'] = '## HP: **12**'
    projection['units'][2]['ru_sha256'] = hashlib.sha256(projection['units'][2]['ru_text'].encode()).hexdigest().upper()
    role['units'][1]['ru_sha256'] = projection['units'][2]['ru_sha256']
    body = html.tostring(tree, encoding='unicode').replace('<p id="stats">', '<h2 id="stats">').replace('</strong></p>', '</strong></h2>').replace('<p id="last">', '<p>')
    options = {'source_toc': {'state': 'absent', 'source_pdf_sha256': 'A'*64, 'reason': 'fixture'},
               'body_links': [], 'asset_choices': [], 'hit_points': {'presentation': 'numeric', 'expected_counts': []},
               'display_roles': [{'role': 'prose_presentation', 'unit': role['units'][1], 'source_evidence': 'fixture', 'reason': 'field'},
                                 {'role': 'review_anchor', 'unit': role['units'][2], 'target_text': 'Dodge 40%', 'source_evidence': 'fixture', 'reason': 'review'}, role]}
    # Bind the exact post-presentation ids generated by the existing roles.
    baseline_options = copy.deepcopy(options); baseline_options['display_roles'].pop()
    baseline, _, _ = project_reader_html(body, [], projection, baseline_options)
    final_tree = html.fragment_fromstring(baseline, create_parent='main')
    for m, n in zip(role['rendered_members'], final_tree.xpath('.//section/*')):
        m['final_html_sha256'] = grouping.fingerprint(n)
    projected, _, report = project_reader_html(body, [], projection, options)
    assert projected.count(grouping.ATTRIBUTE) == 3
    assert len(report['statblock_groups']) == 1 and len(report['review_anchors']) == 1
