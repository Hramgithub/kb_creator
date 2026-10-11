"""Reviewed table presentation flows through reader HTML into the production DOCX."""
from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import html

from kb_creator.pipeline.reader_projection import project_reader_html
from kb_creator.scripts.production_docx import ROLE_STYLE_NAMES, build_docx


def _unit(text: str) -> dict:
    return {"unit_id": "table-1", "ru_text": text, "ru_sha256": hashlib.sha256(text.encode()).hexdigest().upper()}


def _options(unit: dict, **decision) -> dict:
    return {
        "source_toc": {"state": "absent", "source_pdf_sha256": "A" * 64, "reason": "Reviewed absence"},
        "body_links": [], "asset_choices": [], "hit_points": {"presentation": "numeric", "expected_counts": []},
        "display_roles": [{
            "role": "table_presentation", "unit": {key: unit[key] for key in ("unit_id", "ru_sha256")},
            "suppress_empty_transport_header": False, "source_header_rows": [],
            "source_evidence": "Reviewed source table", "reason": "Reviewed DOCX presentation", **decision,
        }],
    }


def _project(table: str, options: dict, unit: dict) -> tuple[html.HtmlElement, dict]:
    output, _, report = project_reader_html(table, [], {"units": [unit]}, options)
    return html.fragment_fromstring(output, create_parent="main"), report


def _unit_for_table(table: str) -> dict:
    source = html.fragment_fromstring(table)
    return _unit(" ".join(source.text_content().split()))


@pytest.mark.parametrize(
    ("table", "decision", "expected_rows", "headers"),
    [
        (
            "<table><thead><tr><th></th><th></th></tr></thead><tbody><tr><td>A</td><td><em>B</em></td></tr><tr><td>C</td><td>D</td></tr></tbody></table>",
            {"suppress_empty_transport_header": True, "source_header_rows": []},
            [["A", "B"], ["C", "D"]], set(),
        ),
        (
            "<table><thead><tr><th></th><th></th></tr></thead><tbody><tr><td>A</td><td><em>B</em></td></tr><tr><td>C</td><td>D</td></tr></tbody></table>",
            {"suppress_empty_transport_header": True, "source_header_rows": [2]},
            [["A", "B"], ["C", "D"]], {1},
        ),
        (
            "<table><tr><td>данные</td><td>0</td></tr><tr><td>Параметр</td><td>Значение</td></tr><tr><td>СИЛ</td><td>50</td></tr></table>",
            {"source_header_rows": [2]},
            [["данные", "0"], ["Параметр", "Значение"], ["СИЛ", "50"]], {1},
        ),
        (
            "<table><tr><td>Параметр</td><td>Значение</td></tr><tr><td>СИЛ</td><td>50</td></tr></table>",
            {"source_header_rows": [1]},
            [["Параметр", "Значение"], ["СИЛ", "50"]], {0},
        ),
    ],
    ids=["headerless-transport", "headerless-transport-internal-row-two", "internal-row-two", "ordinary-explicit-header"],
)
def test_reviewed_table_presentation_preserves_cells_and_styles_actual_docx(
    tmp_path: Path, table: str, decision: dict, expected_rows: list[list[str]], headers: set[int],
) -> None:
    unit = _unit_for_table(table)
    tree, report = _project(table, _options(unit, **decision), unit)
    projected = tree.xpath(".//table")[0]
    assert projected.get("data-table-header-policy") == "explicit"
    assert [[cell.text_content() for cell in row.xpath("./th|./td")] for row in projected.xpath("./thead/tr|./tbody/tr|./tr")] == expected_rows
    assert report["display_roles"][0]["source_header_rows"] == decision["source_header_rows"]

    html_path = tmp_path / "table.html"
    projected_html = "".join(html.tostring(child, encoding="unicode") for child in tree)
    html_path.write_text(f"<html><head><meta charset='utf-8'></head><body><main><nav id='toc'><h1>TOC</h1><ol><li><a href='#body'>Body</a></li></ol></nav><p id='body'>Body</p>{projected_html}</main></body></html>", encoding="utf-8")
    output = tmp_path / "table.docx"
    build_docx(html_path, output)
    document = Document(output)
    rendered = document.tables[-1]
    assert [[cell.text for cell in row.cells] for row in rendered.rows] == expected_rows
    for row_index, row in enumerate(rendered.rows):
        expected = ROLE_STYLE_NAMES["table_header" if row_index in headers else "table_body"]
        assert [cell.paragraphs[0].style.name for cell in row.cells] == [expected] * len(row.cells)
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    for cell in [value for row in expected_rows for value in row]:
        assert cell in document_xml


def test_legacy_table_first_row_remains_header(tmp_path: Path) -> None:
    html_path = tmp_path / "legacy.html"
    html_path.write_text("<html><body><main><nav id='toc'><h1>TOC</h1><ol><li><a href='#body'>Body</a></li></ol></nav><p id='body'>Body</p><table><tr><td>First</td></tr><tr><td>Second</td></tr></table></main></body></html>", encoding="utf-8")
    output = tmp_path / "legacy.docx"
    build_docx(html_path, output)
    document = Document(output)
    assert document.tables[-1].cell(0, 0).paragraphs[0].style.name == ROLE_STYLE_NAMES["table_header"]
    assert document.tables[-1].cell(1, 0).paragraphs[0].style.name == ROLE_STYLE_NAMES["table_body"]


@pytest.mark.parametrize(
    ("table", "change", "message"),
    [
        ("<table><tr><td>not empty</td></tr><tr><td>data</td></tr></table>", {"suppress_empty_transport_header": True}, "truly empty"),
        ("<table><tr><td><span style='display:none'>hidden</span></td></tr><tr><td>data</td></tr></table>", {"suppress_empty_transport_header": True}, "truly empty"),
        ("<table><tr><td><img src='x.png'></td></tr><tr><td>data</td></tr></table>", {"suppress_empty_transport_header": True}, "truly empty"),
        ("<table><tr><td id='bookmark'></td></tr><tr><td>data</td></tr></table>", {"suppress_empty_transport_header": True}, "truly empty"),
        ("<table><tr><td name='bookmark'></td></tr><tr><td>data</td></tr></table>", {"suppress_empty_transport_header": True}, "truly empty"),
        ("<table><tr><td>one</td></tr><tr><td>two</td></tr></table>", {"source_header_rows": [True]}, "header rows"),
        ("<table><tr><td>one</td></tr><tr><td>two</td></tr></table>", {"source_header_rows": [3]}, "header rows"),
    ],
)
def test_table_presentation_rejects_unsafe_header_operations(table: str, change: dict, message: str) -> None:
    unit = _unit_for_table(table)
    with pytest.raises(ValueError, match=message):
        _project(table, _options(unit, **change), unit)


def test_table_presentation_rejects_stale_ambiguous_duplicate_and_unknown_decisions() -> None:
    table = "<table><tr><td>A</td><td>B</td></tr><tr><td>x</td><td>y</td></tr><tr><td>m</td><td>n</td></tr></table>"
    unit = _unit_for_table(table)
    stale = _options(unit)
    stale["display_roles"][0]["unit"]["ru_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="preimage"):
        _project(table, stale, unit)
    with pytest.raises(ValueError, match="exact HTML matches"):
        _project(table + table, _options(unit), unit)
    duplicate = _options(unit)
    duplicate["display_roles"].append(duplicate["display_roles"][0].copy())
    with pytest.raises(ValueError, match="duplicate"):
        _project(table, duplicate, unit)
    unknown = _options(unit)
    unknown["display_roles"][0]["extra"] = True
    with pytest.raises(ValueError, match="unknown"):
        _project(table, unknown, unit)


def test_docx_rejects_invalid_explicit_table_markers(tmp_path: Path) -> None:
    html_path = tmp_path / "invalid.html"
    html_path.write_text("<html><body><main><nav id='toc'><h1>TOC</h1><ol><li><a href='#body'>Body</a></li></ol></nav><p id='body'>Body</p><table data-table-header-policy='explicit' data-table-source-header-rows='1'><tr data-table-source-row='1' data-table-header='false'><td>A</td></tr><tr data-table-source-row='2' data-table-header='false'><td>B</td></tr></table></main></body></html>", encoding="utf-8")
    with pytest.raises(ValueError, match="disagrees"):
        build_docx(html_path, tmp_path / "invalid.docx")
