from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image

from kb_creator.pipeline import reviewed_docx_revision as revision
from kb_creator.pipeline.reviewed_docx_revision import RevisionError, revise_docx
from kb_creator.pipeline.reviewed_reader_revision import _drawing_order, _relationships


def sha(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest().upper()


def picture(color: str, fmt: str) -> bytes:
    stream = BytesIO()
    Image.new("RGB", (24, 12), color).save(stream, format=fmt)
    return stream.getvalue()


def add_bookmark(paragraph, name: str, number: int) -> None:
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def fixture(root: Path):
    reviewed, source, authority = root / "reviewed.docx", root / "source.bin", root / "decisions.json"
    source.write_bytes(b"sealed source authority")
    authority.write_text('{"reviewed":true}', encoding="utf-8")
    old, blue, green = picture("red", "PNG"), picture("blue", "PNG"), picture("green", "JPEG")
    image1, image2 = root / "blue.jpg", root / "green.png"
    image1.write_bytes(blue)
    image2.write_bytes(green)
    document = Document()
    front = document.add_paragraph("Front matter")
    add_bookmark(front, "front", 1)
    h1 = document.add_heading("Story One", 1)
    add_bookmark(h1, "s1", 2)
    rich = document.add_paragraph()
    rich.add_run("Rich ").bold = True
    rich.add_run("old term")
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    h2 = document.add_heading("Story Two", 1)
    add_bookmark(h2, "s2", 3)
    document.add_paragraph("Second old")
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    document.save(reviewed)
    with zipfile.ZipFile(reviewed) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    pkg = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    ct = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    docroot = etree.fromstring(parts["word/document.xml"])
    for marker, number in (("Rich old term", "7"), ("Second old", "8"), ("Story One", "9")):
        paragraph = next(p for p in docroot.iter(w + "p") if marker == "".join(p.itertext()))
        paragraph.insert(0, etree.Element(w + "commentRangeStart", {w + "id": number}))
        paragraph.append(etree.Element(w + "commentRangeEnd", {w + "id": number}))
        run = etree.SubElement(paragraph, w + "r")
        etree.SubElement(run, w + "commentReference", {w + "id": number})
    parts["word/document.xml"] = etree.tostring(docroot, encoding="UTF-8")
    comments = etree.Element(w + "comments", nsmap={"w": w[1:-1]})
    for number in ("7", "8", "9"):
        item = etree.SubElement(comments, w + "comment", {w + "id": number, w + "author": "Reviewer"})
        etree.SubElement(etree.SubElement(item, w + "p"), w + "r")
    parts["word/comments.xml"] = etree.tostring(comments, encoding="UTF-8")
    parts["word/commentsExtended.xml"] = b"<native-reply-payload/>"
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(rels, pkg + "Relationship", Id="rIdReview", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", Target="comments.xml")
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels, encoding="UTF-8")
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, ct + "Override", PartName="/word/comments.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    parts["[Content_Types].xml"] = etree.tostring(types, encoding="UTF-8")
    with zipfile.ZipFile(reviewed, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    relationships = _relationships(etree.fromstring(parts["word/_rels/document.xml.rels"]))
    originals = _drawing_order(etree.fromstring(parts["word/document.xml"]), relationships, parts)
    assert len(originals) == 3
    rows = [
        {"graphic_id": "retained-1", "story_id": "one", "source_page": 1, "disposition": "must_include", "sha256": sha(old), "occurrence_ordinal": 0, "drawing_id": originals[0][1], "legacy_descr": "Picture 1"},
        {"graphic_id": "retained-2", "story_id": "one", "source_page": 1, "disposition": "must_include", "sha256": sha(old), "occurrence_ordinal": 1, "drawing_id": originals[1][1], "legacy_descr": "Picture 2"},
        {"graphic_id": "retained-3", "story_id": "two", "source_page": 2, "disposition": "must_include", "sha256": sha(old), "occurrence_ordinal": 0, "drawing_id": originals[2][1], "legacy_descr": "Picture 3"},
        {"graphic_id": "new-1", "story_id": "one", "source_page": 1, "disposition": "must_include", "sha256": sha(blue), "occurrence_ordinal": 0},
        {"graphic_id": "new-2", "story_id": "two", "source_page": 2, "disposition": "must_include", "sha256": sha(green), "occurrence_ordinal": 0},
    ]
    def selector():
        return {"kind": "exact_paragraph", "expected_matches": 1, "ordinal": 0}
    def text_op(gid, sid, bookmark, comment, expected, target, replacement):
        return {"kind": "text", "issue_id": gid, "story_id": sid, "word_bookmark": bookmark,
                "comment_id": comment, "disposition": "replace", "expected_text": expected,
                "expected_text_sha256": sha(expected.encode()), "selector": selector(),
                "target_text": target, "replacement": replacement}
    def art_op(gid, sid, bookmark, page, path, payload, position):
        return {"kind": "art", "graphic_id": gid, "story_id": sid, "word_bookmark": bookmark,
                "source_page": page, "source_image": str(path), "sha256": sha(payload),
                "dimensions": [24, 12], "word_width_inches": 1.0,
                "insertion": "after_exact_paragraph", "position": {"expected_text": position,
                "expected_text_sha256": sha(position.encode()), "selector": selector()}}
    manifest = {"schema": revision.SCHEMA, "provenance": "reviewed_artifact_successor",
                "inputs": {"docx_sha256": sha(reviewed.read_bytes()), "source_pdf_sha256": sha(source.read_bytes()),
                           "authorities": [{"path": str(authority), "sha256": sha(authority.read_bytes())}],
                           "original_drawing_occurrences": [{"sha256": x, "drawing_id": y} for x, y in originals]},
                "stories": [{"story_id": "front", "word_bookmark": "front"},
                            {"story_id": "one", "word_bookmark": "s1"},
                            {"story_id": "two", "word_bookmark": "s2"}],
                "source_graphic_ids": [r["graphic_id"] for r in rows], "source_graphics": rows,
                "operations": [text_op("issue-one", "one", "s1", "7", "Rich old term", "old", "new"),
                               art_op("new-1", "one", "s1", 1, image1, blue, "Rich new term"),
                               text_op("issue-two", "two", "s2", "8", "Second old", "old", "new"),
                               art_op("new-2", "two", "s2", 2, image2, green, "Second new"),
                               {"kind": "keep", "issue_id": "keep-one", "story_id": "one", "word_bookmark": "s1",
                                "comment_id": "9", "disposition": "keep", "reason": "reviewed wording retained",
                                "expected_text": "Story One", "expected_text_sha256": sha(b"Story One"), "selector": selector()}]}
    return manifest, reviewed, source, parts


def run_case(root: Path, manifest: dict, reviewed: Path, source: Path):
    path = root / "manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return revise_docx(path, reviewed, source, root / "candidate", root / "report.json")


class ReviewedDocxRevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="reviewed-docx-test-", dir=Path(__file__).parent)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_manifest_changed_after_initial_read_fails_before_publish(self):
        manifest, reviewed, source, _ = fixture(self.root)
        manifest_path = self.root / "manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        original = revision._source_bytes
        reads = 0

        def mutate_before_other_snapshot(path):
            nonlocal reads
            if path.resolve() == manifest_path.resolve():
                reads += 1
            if path.resolve() == reviewed.resolve():
                manifest_path.write_bytes(b"{}")
            return original(path)

        with patch.object(revision, "_source_bytes", side_effect=mutate_before_other_snapshot):
            with self.assertRaises(RevisionError):
                revise_docx(manifest_path, reviewed, source,
                            self.root / "candidate", self.root / "report.json")
        self.assertEqual(reads, 1)
        self.assertFalse((self.root / "candidate").exists())
        self.assertFalse((self.root / "report.json").exists())

    def test_docx_only_two_story_successor_preserves_review_and_rich_runs(self):
        manifest, reviewed, source, old_parts = fixture(self.root)
        report = run_case(self.root, manifest, reviewed, source)
        self.assertEqual(report["status"], "validated_candidate")
        self.assertFalse(report["accepted"])
        self.assertEqual(len(report["source_graphic_coverage"]), 5)
        self.assertEqual(len(report["operations"]), 5)
        self.assertEqual([p.name for p in (self.root / "candidate").iterdir()], ["document.docx"])
        self.assertFalse((self.root / "candidate" / "document.html").exists())
        self.assertFalse((self.root / "candidate" / "document.pdf").exists())
        out = self.root / "candidate" / "document.docx"
        with zipfile.ZipFile(out) as archive:
            names = archive.namelist()
            for name in old_parts:
                if name.startswith("word/media/") or "comment" in name.lower():
                    self.assertEqual(archive.read(name), old_parts[name])
            self.assertTrue(any(n.endswith(".png") and "revision-" in n for n in names))
            self.assertTrue(any(n.endswith(".jpg") and "revision-" in n for n in names))
            types = archive.read("[Content_Types].xml")
            self.assertIn(b'image/png', types)
            self.assertIn(b'image/jpeg', types)
            root = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        text = ["".join(p.xpath(".//w:t/text()", namespaces=ns)) for p in root.xpath(".//w:body/w:p", namespaces=ns)]
        self.assertIn("Rich new term", text)
        self.assertIn("Second new", text)
        self.assertTrue(root.xpath(".//w:rPr/w:b", namespaces=ns))
        self.assertEqual(len(Document(out).inline_shapes), 5)

    def test_fear_seed_closes_independent_26_openings_and_23_insertions(self):
        path = (Path(__file__).parents[1] / "docs/plans/20261004-user-docx-review/audit"
                / "fear-candidate/fear-docx-only-manifest-v01.json")
        manifest = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["schema"], revision.SCHEMA)
        self.assertEqual(len(manifest["source_graphic_ids"]), 26)
        self.assertEqual(len(set(manifest["source_graphic_ids"])), 26)
        self.assertEqual({row["graphic_id"] for row in manifest["source_graphics"]},
                         set(manifest["source_graphic_ids"]))
        self.assertEqual(sum(op["kind"] == "art" for op in manifest["operations"]), 23)
        self.assertEqual(sum(op["kind"] == "text" for op in manifest["operations"]), 6)
        self.assertEqual(sum(op["kind"] == "keep" for op in manifest["operations"]), 5)
        self.assertEqual(len(manifest["inputs"]["original_drawing_occurrences"]), 63)
        self.assertEqual({row["source_page"] for row in manifest["source_graphics"]
                          if row["graphic_id"] not in {op["graphic_id"] for op in manifest["operations"]
                                                       if op["kind"] == "art"}}, {46, 113, 129})

    def test_missing_art_position_selector_and_hash_fail(self):
        for number, mutation in enumerate((
            lambda m: m["operations"][1]["position"].pop("selector"),
            lambda m: m["operations"][1]["position"].update(expected_text_sha256="0" * 64),
            lambda m: m["operations"][1]["position"].update(expected_text="changed"),
        )):
            with self.subTest(case=number):
                case = self.root / str(number)
                case.mkdir()
                manifest, reviewed, source, _ = fixture(case)
                mutation(manifest)
                with self.assertRaises(RevisionError):
                    run_case(case, manifest, reviewed, source)
                self.assertFalse((case / "candidate").exists())

    def test_source_or_manifest_midbuild_mutation_fails_without_output(self):
        for number, target in enumerate(("source", "manifest")):
            with self.subTest(target=target):
                case = self.root / str(number)
                case.mkdir()
                manifest, reviewed, source, _ = fixture(case)
                manifest_path = case / "manifest.json"
                original = revision._add_word_image

                def mutate(*args):
                    result = original(*args)
                    changed = source if target == "source" else manifest_path
                    changed.write_bytes(changed.read_bytes() + b"changed")
                    return result

                with patch.object(revision, "_add_word_image", side_effect=mutate):
                    with self.assertRaisesRegex(RevisionError, "authority input changed"):
                        run_case(case, manifest, reviewed, source)
                self.assertFalse((case / "candidate").exists())

    def test_missing_or_reused_required_art_fails(self):
        for number, mutation in enumerate((
            lambda m: m["source_graphics"].pop(),
            lambda m: m["source_graphics"][0].update(disposition="reference_only"),
            lambda m: m["source_graphics"][1].update(occurrence_ordinal=0),
            lambda m: m["source_graphics"][0].update(story_id="two"),
            lambda m: m["source_graphic_ids"].extend([f"missing-{i}" for i in range(21)]),
        )):
            with self.subTest(case=number):
                case = self.root / str(number)
                case.mkdir()
                manifest, reviewed, source, _ = fixture(case)
                mutation(manifest)
                with self.assertRaises(RevisionError):
                    run_case(case, manifest, reviewed, source)
                self.assertFalse((case / "candidate").exists())

    def test_stale_hash_alias_comment_and_text_context_fail(self):
        for number, mutation in enumerate((
            lambda m: m["inputs"].update(source_pdf_sha256="0" * 64),
            lambda m: m["inputs"]["original_drawing_occurrences"].reverse(),
            lambda m: m["operations"][0].update(comment_id="99"),
            lambda m: m["operations"][0].update(target_text="Rich old"),
            lambda m: m["operations"][0]["selector"].update(expected_matches=2),
            lambda m: m["operations"][1].update(word_width_inches=0),
        )):
            with self.subTest(case=number):
                case = self.root / str(number)
                case.mkdir()
                manifest, reviewed, source, _ = fixture(case)
                mutation(manifest)
                with self.assertRaises(RevisionError):
                    run_case(case, manifest, reviewed, source)
        manifest, reviewed, source, _ = fixture(self.root)
        path = self.root / "manifest.json"
        path.write_text(json.dumps(manifest), encoding="utf-8")
        with self.assertRaises(RevisionError):
            revise_docx(path, reviewed, source, reviewed.parent, self.root / "report.json")


if __name__ == "__main__":
    unittest.main()
