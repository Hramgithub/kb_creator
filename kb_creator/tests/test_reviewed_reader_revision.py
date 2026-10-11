from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image

from kb_creator.pipeline.reviewed_reader_revision import RevisionError, revise_reader
from kb_creator.pipeline import reviewed_reader_revision as revision


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def png(color: str) -> bytes:
    stream = BytesIO()
    Image.new("RGB", (20, 10), color).save(stream, format="PNG")
    return stream.getvalue()


def bookmark(paragraph, name: str, number: int) -> None:
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def fixture(tmp_path: Path):
    docx = tmp_path / "reviewed.docx"
    html_path = tmp_path / "reviewed.html"
    pdf = tmp_path / "source.pdf"
    assets = tmp_path / "assets"
    assets.mkdir(exist_ok=True)
    old, first, second = png("red"), png("blue"), png("green")
    (assets / "old.png").write_bytes(old)
    pdf.write_bytes(b"%PDF-1.4\nsealed synthetic source\n")
    document = Document()
    heading1 = document.add_heading("Story One", level=1)
    bookmark(heading1, "story_one", 1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Rich ").bold = True
    paragraph.add_run("skill old")
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    heading2 = document.add_heading("Story Two", level=1)
    bookmark(heading2, "story_two", 2)
    document.add_paragraph("Second old")
    document.add_paragraph().add_run().add_picture(BytesIO(old), width=Inches(0.5))
    document.save(docx)
    # Retain native review payload and its anchors as immutable OOXML fixtures.
    with zipfile.ZipFile(docx) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    rel = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    ct = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    root = etree.fromstring(parts["word/document.xml"])
    target = next(p for p in root.iter(w + "p") if "skill old" in "".join(p.itertext()))
    target.insert(0, etree.Element(w + "commentRangeStart", {w + "id": "7"}))
    target.append(etree.Element(w + "commentRangeEnd", {w + "id": "7"}))
    ref_run = etree.SubElement(target, w + "r")
    etree.SubElement(ref_run, w + "commentReference", {w + "id": "7"})
    parts["word/document.xml"] = etree.tostring(root, encoding="UTF-8")
    comments = etree.Element(w + "comments", nsmap={"w": w[1:-1]})
    comment = etree.SubElement(comments, w + "comment", {w + "id": "7", w + "author": "Reviewer"})
    etree.SubElement(etree.SubElement(comment, w + "p"), w + "r").append(etree.Element(w + "t"))
    parts["word/comments.xml"] = etree.tostring(comments, encoding="UTF-8")
    parts["word/commentsExtended.xml"] = b"<review-thread-preserved/>"
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    etree.SubElement(rels, rel + "Relationship", Id="rIdReviewComment", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", Target="comments.xml")
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels, encoding="UTF-8")
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, ct + "Override", PartName="/word/comments.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    parts["[Content_Types].xml"] = etree.tostring(types, encoding="UTF-8")
    with zipfile.ZipFile(docx, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    html_path.write_text('<html><body><h1 id="one">Story One</h1><p><strong>Rich </strong>skill old</p>'
                         '<img src="assets/old.png"><img src="assets/old.png">'
                         '<h1 id="two">Story Two</h1><p>Second old</p><img src="assets/old.png"></body></html>', encoding="utf-8")
    image1, image2 = tmp_path / "first.png", tmp_path / "second.png"
    image1.write_bytes(first)
    image2.write_bytes(second)
    stories = [{"story_id": "s1", "word_bookmark": "story_one", "html_anchor": "one"},
               {"story_id": "s2", "word_bookmark": "story_two", "html_anchor": "two"}]
    rows = [
        {"graphic_id": "retained-1", "story_id": "s1", "source_page": 1, "disposition": "must_include", "sha256": digest(old), "occurrence_ordinal": 0},
        {"graphic_id": "retained-2", "story_id": "s1", "source_page": 1, "disposition": "must_include", "sha256": digest(old), "occurrence_ordinal": 1},
        {"graphic_id": "retained-3", "story_id": "s2", "source_page": 2, "disposition": "must_include", "sha256": digest(old), "occurrence_ordinal": 0},
        {"graphic_id": "new-1", "story_id": "s1", "source_page": 1, "disposition": "must_include", "sha256": digest(first), "occurrence_ordinal": 0},
        {"graphic_id": "new-2", "story_id": "s2", "source_page": 2, "disposition": "must_include", "sha256": digest(second), "occurrence_ordinal": 0},
    ]
    def text_op(issue, story, bookmark_name, anchor, expected, target, replacement, tag="p"):
        return {"kind": "text", "issue_id": issue, "story_id": story, "word_bookmark": bookmark_name,
                "html_anchor": anchor, "disposition": "replace", "expected_text": expected,
                "expected_text_sha256": digest(expected.encode()), "target_text": target, "replacement": replacement,
                "selector": {"word": {"kind": "exact_paragraph", "ordinal": 0, "expected_matches": 1},
                             "html": {"kind": "exact_element", "tag": tag, "ordinal": 0, "expected_matches": 1}}}
    def art_op(gid, story, bookmark_name, anchor, source, page, sha):
        return {"kind": "art", "graphic_id": gid, "story_id": story, "word_bookmark": bookmark_name,
                "html_anchor": anchor, "source_image": str(source), "source_page": page, "sha256": sha,
                "dimensions": [20, 10], "asset_path": f"assets/{gid}.png", "insertion": "after_anchor"}
    manifest = {"schema": "kb_creator.reviewed_reader_revision.v1", "provenance": "reviewed_artifact_successor",
                "inputs": {"docx_sha256": digest(docx.read_bytes()), "html_sha256": digest(html_path.read_bytes()),
                           "source_pdf_sha256": digest(pdf.read_bytes()), "existing_assets": {"assets/old.png": digest(old)}},
                "stories": stories, "source_graphic_ids": [r["graphic_id"] for r in rows], "source_graphics": rows,
                "operations": [text_op("issue-one", "s1", "story_one", "one", "Rich skill old", "old", "new"),
                               art_op("new-1", "s1", "story_one", "one", image1, 1, digest(first)),
                               text_op("issue-two", "s2", "story_two", "two", "Second old", "old", "new"),
                               art_op("new-2", "s2", "story_two", "two", image2, 2, digest(second)),
                               {"kind": "keep", "issue_id": "no-change", "disposition": "keep", "reason": "reviewed wording retained"}]}
    return manifest, (docx, html_path, assets, pdf), parts


def run_case(tmp_path, manifest, inputs):
    manifest_path = tmp_path / "decision.v1.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return revise_reader(manifest_path, *inputs, tmp_path / "successor")


MUTATORS = [
    lambda m, paths: m["inputs"].update(source_pdf_sha256="0" * 64),
    lambda m, paths: m["inputs"].update(html_sha256="0" * 64),
    lambda m, paths: m["inputs"]["existing_assets"].update({"assets/old.png": "0" * 64}),
    lambda m, paths: m["operations"][1].update(sha256="0" * 64),
    lambda m, paths: m["operations"][1].update(asset_path="assets/../escape.png"),
    lambda m, paths: m["operations"][1].update(word_width_inches=0),
    lambda m, paths: m["operations"][0].update(expected_text="changed context"),
    lambda m, paths: m["operations"][0].update(target_text="Rich skill"),
    lambda m, paths: m["operations"][0].update(target_text="l"),
    lambda m, paths: m["stories"][0].update(word_bookmark="missing_bookmark"),
    lambda m, paths: m["source_graphics"].pop(),
    lambda m, paths: m["source_graphics"][0].update(disposition="reference_only"),
    lambda m, paths: m["source_graphics"][0].update(story_id="s2"),
    lambda m, paths: m["source_graphics"][1].update(occurrence_ordinal=0),
]


class ReviewedReaderRevisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="reviewed-reader-test-", dir=Path(__file__).parent)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def _assert_retired(self, manifest, inputs, output):
        path = self.root / "decision.v1.json"
        path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(RevisionError, "DOCX-only"):
            revise_reader(path, *inputs, output)
        self.assertFalse(list(self.root.glob(".*.staging.*")))

    def test_valid_old_reader_writer_is_retired_before_output(self):
        manifest, inputs, old_parts = fixture(self.root)
        self._assert_retired(manifest, inputs, self.root / "successor")
        self.assertFalse((self.root / "successor").exists())
        self.assertEqual(revision._zip_parts(inputs[0])["word/document.xml"], old_parts["word/document.xml"])

    def test_all_old_manifest_variants_fail_before_output(self):
        for index, mutator in enumerate(MUTATORS):
            with self.subTest(case=index):
                case = self.root / str(index)
                case.mkdir()
                manifest, inputs, _ = fixture(case)
                mutator(manifest, inputs)
                path = case / "decision.v1.json"
                path.write_text(json.dumps(manifest), encoding="utf-8")
                with self.assertRaisesRegex(RevisionError, "DOCX-only"):
                    revise_reader(path, *inputs, case / "successor")
                self.assertFalse((case / "successor").exists())

    def test_drift_and_output_alias_do_not_revive_writer(self):
        manifest, inputs, _ = fixture(self.root)
        inputs[3].write_bytes(inputs[3].read_bytes() + b"changed")
        self._assert_retired(manifest, inputs, self.root / "successor")
        self._assert_retired(manifest, inputs, inputs[2])
        self.assertFalse((self.root / "successor").exists())

    def test_existing_output_remains_untouched(self):
        manifest, inputs, _ = fixture(self.root)
        output = self.root / "successor"
        output.mkdir()
        marker = output / "user-created.txt"
        marker.write_text("retain", encoding="utf-8")
        self._assert_retired(manifest, inputs, output)
        self.assertEqual(marker.read_text(encoding="utf-8"), "retain")

    def test_reviewed_exclusion_cannot_restart_dual_output(self):
        manifest, inputs, _ = fixture(self.root)
        manifest["source_graphics"].append({"graphic_id": "ornament", "story_id": "s1", "source_page": 1,
                                            "disposition": "exclude", "reviewed": True,
                                            "reason": "reviewed decorative exclusion"})
        manifest["source_graphic_ids"].append("ornament")
        self._assert_retired(manifest, inputs, self.root / "successor")
        self.assertFalse((self.root / "successor").exists())

    def test_pure_word_helper_still_rejects_unsupported_run_topology(self):
        _, inputs, _ = fixture(self.root)
        parts = revision._zip_parts(inputs[0])
        root = etree.fromstring(parts["word/document.xml"])
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        run = next(r for r in root.xpath(".//w:r", namespaces=ns)
                   if "skill old" in "".join(r.xpath("./w:t/text()", namespaces=ns)))
        etree.SubElement(run, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br")
        paragraph = run.getparent()
        while paragraph is not None and paragraph.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p":
            paragraph = paragraph.getparent()
        class Visible:
            def hidden(self, _paragraph, _run):
                return False
        with self.assertRaisesRegex(RevisionError, "unsupported run topology"):
            revision._plain_text_runs(paragraph, Visible())

    def test_pure_zip_helper_preserves_original_media_and_comments(self):
        _, inputs, original = fixture(self.root)
        reread = revision._zip_parts(inputs[0])
        for name, payload in original.items():
            if name.startswith("word/media/") or "comment" in name.lower():
                self.assertEqual(reread[name], payload)


if __name__ == "__main__":
    unittest.main()
