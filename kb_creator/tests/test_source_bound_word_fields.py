from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from kb_creator.pipeline.docx_review_comments import _StyleVisibility
from kb_creator.pipeline.reviewed_docx_revision import (
    FIELD_APPROVED_STATE, FIELD_LEDGER_SCHEMA, RevisionError,
    _approved_field_rule, _node_xml_sha, _predicted_node_xml_sha, _sha,
    _skill_consolidation_binding, _word_leaf,
    _source_bound_plan, revise_docx,
)
from kb_creator.pipeline.reviewed_reader_revision import NS, W, _drawing_order, _relationships, _xml


DOCX_HASH = "A" * 64
PDF_HASH = "B" * 64
BOOK = "synthetic-source-book"


def text_run(paragraph, value: str, *, bold: bool = False):
    run = etree.SubElement(paragraph, W + "r")
    if bold:
        etree.SubElement(etree.SubElement(run, W + "rPr"), W + "b")
    leaf = etree.SubElement(run, W + "t")
    leaf.text = value
    return leaf


def example():
    document = etree.Element(W + "document", nsmap={"w": NS["w"]})
    body = etree.SubElement(document, W + "body")
    anchor = etree.SubElement(body, W + "p")
    text_run(anchor, "Катерина, маленький вампир")
    table = etree.SubElement(body, W + "tbl")
    cell = etree.SubElement(etree.SubElement(etree.SubElement(table, W + "tr"), W + "tc"), W + "p")
    table_leaf = text_run(cell, "СИЛ: 13 Удача: 75")
    paragraph = etree.SubElement(body, W + "p")
    multi_leaf = text_run(paragraph, "СИЛ: 15 ЛВК: 12 РАС: 70 Навык: 70")
    effect = etree.SubElement(body, W + "p")
    effect_leaves = [text_run(effect, "1"), text_run(effect, "D", bold=True), text_run(effect, "6 СИЛ")]
    paths = ["word/document.xml/body[1]/tbl/tr[0]/tc[0]/p[0]/r[0]/t[0]",
             "word/document.xml/body[2]/p/r[0]/t[0]"]
    anchor_data = {"body_index": 0, "text": "Катерина, маленький вампир",
                   "text_sha256": _sha("Катерина, маленький вампир".encode("utf-8")),
                   "bookmarks": []}
    ops = []
    groups = []

    def add_group(path, leaf, edits):
        before_node = leaf.text
        after_node = before_node
        ids = []
        for ident, field, start, old, new in edits:
            ids.append(ident)
            ops.append({"id": ident, "book_id": BOOK, "story_id": "scenario",
                        "entity_id": "Катерина", "entity_kind": "human_npc", "entity_anchor": anchor_data,
                        "field": field, "kind": "characteristic", "disposition": "convert_deterministic",
                        "current_scale": "old_edition", "input_docx_sha256": DOCX_HASH,
                        "source_verified": True,
                        "source_pdf_sha256": PDF_HASH, "source_pdf_path": str(Path("D:/NRI/synthetic.pdf").resolve()),
                        "source_physical_page": 7, "source_edition_evidence": "fifth edition",
                        "selector": {"part": "word/document.xml", "path": path,
                                     "body_index": int(path.split("body[")[1].split("]")[0]),
                                     "table_row": 0 if "/tbl/" in path else None,
                                     "table_cell": 0 if "/tbl/" in path else None,
                                     "text_node_before": before_node,
                                     "text_node_sha256": _sha(before_node.encode("utf-8")),
                                     "span_start": start, "span_end": start + len(old)},
                        "before_text": old, "before_sha256": _sha(old.encode("utf-8")),
                        "after_text": new, "reverse_before": new, "reverse_after": old,
                        "rule_id": "official_old_characteristic_times_five",
                        "editorial_decision": None, "node_group_ids": ids})
        for _, _, start, old, new in sorted(edits, key=lambda item: item[2], reverse=True):
            after_node = after_node[:start] + new + after_node[start + len(old):]
        groups.append({"book_id": BOOK, "path": path, "operation_ids": ids,
                       "before_node": before_node, "before_node_sha256": _sha(before_node.encode("utf-8")),
                       "before_node_xml_sha256": _node_xml_sha(leaf),
                       "after_node": after_node, "after_node_sha256": _sha(after_node.encode("utf-8")),
                       "after_node_xml_sha256": _predicted_node_xml_sha(leaf, after_node),
                       "reverse_before_node": after_node, "reverse_after_node": before_node})

    add_group(paths[0], table_leaf, [("table-str", "СИЛ", 0, "СИЛ: 13", "СИЛ: 65")])
    add_group(paths[1], multi_leaf, [("inline-str", "СИЛ", 0, "СИЛ: 15", "СИЛ: 75"),
                                     ("inline-dex", "ЛВК", 8, "ЛВК: 12", "ЛВК: 60")])
    effect_paths = [f"word/document.xml/body[3]/p/r[{i}]/t[0]" for i in range(3)]
    effect_after = ["5 × (1", "D", "6) СИЛ"]
    expression_nodes = []
    for path, leaf, after in zip(effect_paths, effect_leaves, effect_after):
        expression_nodes.append({"path": path, "before_node": leaf.text,
                                 "before_node_sha256": _sha(leaf.text.encode("utf-8")),
                                 "before_node_xml_sha256": _node_xml_sha(leaf),
                                 "after_node": after, "after_node_sha256": _sha(after.encode("utf-8")),
                                 "after_node_xml_sha256": _predicted_node_xml_sha(leaf, after)})
    expression = {"id": "katerina-strength-drain", "book_id": BOOK, "field": "СИЛ",
                  "rule_id": "official_characteristic_effect_times_five",
                  "current_scale": "old_edition", "entity_id": "Катерина", "entity_anchor": anchor_data,
                  "source_verified": True,
                  "source_pdf_sha256": PDF_HASH, "source_pdf_path": str(Path("D:/NRI/synthetic.pdf").resolve()),
                  "source_physical_page": 7, "source_edition_evidence": "fifth edition",
                  "before_expression": "1D6 СИЛ", "after_expression": "5 × (1D6) СИЛ",
                  "reverse_before_expression": "5 × (1D6) СИЛ",
                  "reverse_after_expression": "1D6 СИЛ", "nodes": expression_nodes}
    ledger = {"schema": FIELD_LEDGER_SCHEMA, "state": FIELD_APPROVED_STATE,
              "source_qa_closed": True, "orchestrator_approved": True,
              "inputs": {BOOK: {"docx_sha256": DOCX_HASH, "source_pdf_sha256": PDF_HASH,
                                "source_pdf_path": str(Path("D:/NRI/synthetic.pdf").resolve())}},
              "operation_count": len(ops), "operations": ops, "node_groups": groups,
              "expression_groups": [expression]}
    choices = {"schema": "nri.user_conversion_choices.v1", "decisions": []}
    paths_by_type = {name: Path(f"D:/NRI/synthetic-{name}.json").resolve()
                     for name in ("ledger", "choices", "conversion_spec")}
    payloads = {paths_by_type["ledger"]: json.dumps(ledger, ensure_ascii=False).encode("utf-8"),
                paths_by_type["choices"]: json.dumps(choices, ensure_ascii=False).encode("utf-8"),
                paths_by_type["conversion_spec"]: b"sealed conversion authority"}
    config = {name: {"path": str(path), "sha256": _sha(payloads[path])}
              for name, path in paths_by_type.items()}
    config.update(book_id=BOOK, operation_ids=[op["id"] for op in ops],
                  expression_ids=[expression["id"]])
    inputs = {"docx_sha256": DOCX_HASH, "source_pdf_sha256": PDF_HASH}
    return document, ledger, choices, config, payloads, inputs


def rebind(ledger, choices, config, snapshots):
    for name, data in (("ledger", ledger), ("choices", choices)):
        path = Path(config[name]["path"]).resolve()
        snapshots[path] = json.dumps(data, ensure_ascii=False).encode("utf-8")
        config[name]["sha256"] = _sha(snapshots[path])


def plan(document, config, snapshots, inputs):
    return _source_bound_plan(config, snapshots, document, inputs,
                              Path("D:/NRI/synthetic.pdf").resolve())


class SourceBoundWordFieldsTests(unittest.TestCase):
    def test_split_expression_fragment_gaps_reject_and_exterior_prose_survives(self):
        for case in ("first_tail", "last_prefix", "middle_gap", "surrounding_prose"):
            with self.subTest(case=case):
                document, ledger, choices, config, snapshots, inputs = example()
                expression = ledger["expression_groups"][0]
                nodes = expression["nodes"]
                first = _word_leaf(document, nodes[0]["path"])[0]
                middle = _word_leaf(document, nodes[1]["path"])[0]
                last = _word_leaf(document, nodes[2]["path"])[0]
                if case == "first_tail":
                    first.text = "1 DAMAGE "
                    nodes[0].update(span_start=0, span_end=1, after_fragment="5 × (1",
                                    after_node="5 × (1 DAMAGE ")
                elif case == "last_prefix":
                    last.text = "DAMAGE 6 СИЛ"
                    nodes[2].update(span_start=7, span_end=12, after_fragment="6) СИЛ",
                                    after_node="DAMAGE 6) СИЛ")
                elif case == "middle_gap":
                    middle.text = "XDY"
                    nodes[1].update(span_start=1, span_end=2, after_fragment="D",
                                    after_node="XDY")
                else:
                    first.text = "The bite drains 1"
                    nodes[0].update(span_start=len(first.text) - 1, span_end=len(first.text),
                                    after_fragment="5 × (1", after_node="The bite drains 5 × (1")
                    last.text = "6 СИЛ each round"
                    nodes[2].update(span_start=0, span_end=5, after_fragment="6) СИЛ",
                                    after_node="6) СИЛ each round")
                for node in nodes:
                    leaf = _word_leaf(document, node["path"])[0]
                    node["before_node"] = leaf.text
                    node["before_node_sha256"] = _sha(leaf.text.encode("utf-8"))
                    node["before_node_xml_sha256"] = _node_xml_sha(leaf)
                    node["after_node_sha256"] = _sha(node["after_node"].encode("utf-8"))
                    node["after_node_xml_sha256"] = _predicted_node_xml_sha(leaf, node["after_node"])
                rebind(ledger, choices, config, snapshots)
                if case == "surrounding_prose":
                    self.assertEqual(len(plan(document, config, snapshots, inputs)), 5)
                else:
                    with self.assertRaises(RevisionError):
                        plan(document, config, snapshots, inputs)

    def test_partial_numeric_operand_rejected_both_directions(self):
        for before, after, start, end, final, accepted in (
            ("3", "15", 1, 2, "115", False),
            ("1", "5", 0, 1, "53", False),
            ("13", "65", 0, 2, "65", True),
        ):
            with self.subTest(before=before):
                document, ledger, choices, config, snapshots, inputs = example()
                group, op = ledger["node_groups"][0], ledger["operations"][0]
                leaf, _, table = _word_leaf(document, group["path"])
                leaf.text = "13"
                heading = etree.Element(W + "tr")
                header = etree.SubElement(etree.SubElement(etree.SubElement(heading, W + "tc"), W + "p"), W + "r")
                etree.SubElement(header, W + "t").text = "СИЛ"
                table.insert(0, heading)
                path = group["path"].replace("/tr[0]/", "/tr[1]/")
                op.update(before_text=before, before_sha256=_sha(before.encode("utf-8")),
                          after_text=after, reverse_before=after, reverse_after=before)
                op["selector"].update(path=path, table_row=1, text_node_before="13",
                                      text_node_sha256=_sha(b"13"), span_start=start, span_end=end)
                group.update(path=path, before_node="13", before_node_sha256=_sha(b"13"),
                             before_node_xml_sha256=_node_xml_sha(leaf), reverse_after_node="13",
                             after_node=final, after_node_sha256=_sha(final.encode("utf-8")),
                             reverse_before_node=final,
                             after_node_xml_sha256=_predicted_node_xml_sha(leaf, final))
                rebind(ledger, choices, config, snapshots)
                if accepted:
                    self.assertEqual(len(plan(document, config, snapshots, inputs)), 5)
                else:
                    with self.assertRaises(RevisionError):
                        plan(document, config, snapshots, inputs)

    def test_visible_numeric_operand_across_inline_word_boundaries(self):
        cases = (
            ("run_prefix", "3", "15", "1", False),
            ("run_suffix", "1", "5", "3", False),
            ("zero_width_bookmark", "3", "15", "1", False),
            ("hyperlink_prefix", "3", "15", "1", False),
            ("same_run_leaf", "3", "15", "1", False),
            ("delimiter", "13", "65", "; Удача: 75", True),
        )
        for case, before, after, neighbor, accepted in cases:
            with self.subTest(case=case):
                document, ledger, choices, config, snapshots, inputs = example()
                group, op = ledger["node_groups"][0], ledger["operations"][0]
                leaf, paragraph, table = _word_leaf(document, group["path"])
                leaf.text = before
                heading = etree.Element(W + "tr")
                header = etree.SubElement(etree.SubElement(etree.SubElement(heading, W + "tc"), W + "p"), W + "r")
                etree.SubElement(header, W + "t").text = "СИЛ"
                table.insert(0, heading)
                path = group["path"].replace("/tr[0]/", "/tr[1]/")
                if case == "same_run_leaf":
                    preceding = etree.Element(W + "t")
                    preceding.text = neighbor
                    leaf.addprevious(preceding)
                    path = path.replace("/t[0]", "/t[1]")
                elif case == "hyperlink_prefix":
                    link = etree.Element(W + "hyperlink")
                    text_run(link, neighbor)
                    leaf.getparent().addprevious(link)
                else:
                    run = etree.Element(W + "r")
                    etree.SubElement(run, W + "t").text = neighbor
                    if case == "run_suffix" or case == "delimiter":
                        leaf.getparent().addnext(run)
                    else:
                        leaf.getparent().addprevious(run)
                        path = path.replace("/r[0]/", "/r[1]/")
                        if case == "zero_width_bookmark":
                            marker = etree.Element(W + "bookmarkStart")
                            marker.set(W + "id", "17")
                            marker.set(W + "name", "between_digits")
                            run.addnext(marker)
                op.update(before_text=before, before_sha256=_sha(before.encode("utf-8")),
                          after_text=after, reverse_before=after, reverse_after=before)
                op["selector"].update(path=path, table_row=1, text_node_before=before,
                                      text_node_sha256=_sha(before.encode("utf-8")),
                                      span_start=0, span_end=len(before))
                group.update(path=path, before_node=before,
                             before_node_sha256=_sha(before.encode("utf-8")),
                             before_node_xml_sha256=_node_xml_sha(leaf),
                             after_node=after, after_node_sha256=_sha(after.encode("utf-8")),
                             after_node_xml_sha256=_predicted_node_xml_sha(leaf, after),
                             reverse_before_node=after, reverse_after_node=before)
                rebind(ledger, choices, config, snapshots)
                if accepted:
                    self.assertEqual(len(plan(document, config, snapshots, inputs)), 5)
                else:
                    with self.assertRaises(RevisionError):
                        plan(document, config, snapshots, inputs)

    def test_resealed_numeric_label_change_rejected_but_exact_crosswalk_allowed(self):
        document, ledger, choices, config, snapshots, inputs = example()
        operation, group = ledger["operations"][0], ledger["node_groups"][0]
        operation.update(after_text="Удача: 65", reverse_before="Удача: 65")
        after = "Удача: 65 Удача: 75"
        leaf, _, _ = _word_leaf(document, group["path"])
        group.update(after_node=after, reverse_before_node=after,
                     after_node_sha256=_sha(after.encode("utf-8")),
                     after_node_xml_sha256=_predicted_node_xml_sha(leaf, after))
        rebind(ledger, choices, config, snapshots)
        with self.assertRaises(RevisionError):
            plan(document, config, snapshots, inputs)

    def test_inherited_hidden_field_and_effect_rejected(self):
        for target in ("field", "effect"):
            with self.subTest(target=target):
                document, ledger, choices, config, snapshots, inputs = example()
                path = (ledger["node_groups"][0]["path"] if target == "field"
                        else ledger["expression_groups"][0]["nodes"][0]["path"])
                leaf, paragraph, _ = _word_leaf(document, path)
                properties = etree.Element(W + "pPr")
                etree.SubElement(etree.SubElement(properties, W + "rPr"), W + "vanish")
                paragraph.insert(0, properties)
                self.assertTrue(_StyleVisibility(None).hidden(paragraph, leaf.getparent()))
                with self.assertRaises(RevisionError):
                    plan(document, config, snapshots, inputs)

    def test_style_inherited_hidden_field_rejected(self):
        document, ledger, choices, config, snapshots, inputs = example()
        paragraph = _word_leaf(document, ledger["node_groups"][0]["path"])[1]
        paragraph.insert(0, etree.fromstring(
            f'<w:pPr xmlns:w="{NS["w"]}"><w:pStyle w:val="HiddenBase"/></w:pPr>'))
        styles = etree.fromstring(
            f'<w:styles xmlns:w="{NS["w"]}"><w:style w:styleId="HiddenBase">'
            '<w:rPr><w:vanish/></w:rPr></w:style></w:styles>')
        with self.assertRaises(RevisionError):
            _source_bound_plan(config, snapshots, document, inputs,
                               Path("D:/NRI/synthetic.pdf").resolve(), _StyleVisibility(styles))

    def test_intervening_hyperlink_and_break_reject_split_expression(self):
        for kind in ("hyperlink", "break"):
            with self.subTest(kind=kind):
                document, ledger, choices, config, snapshots, inputs = example()
                path = ledger["expression_groups"][0]["nodes"][0]["path"]
                leaf, _, _ = _word_leaf(document, path)
                if kind == "hyperlink":
                    link = etree.Element(W + "hyperlink")
                    text_run(link, " DAMAGE BETWEEN ")
                    leaf.getparent().addnext(link)
                else:
                    extra = etree.Element(W + "r")
                    etree.SubElement(extra, W + "br")
                    leaf.getparent().addnext(extra)
                with self.assertRaises(RevisionError):
                    plan(document, config, snapshots, inputs)

    def test_plain_line_break_kept_and_rich_field_hidden_runs_rejected(self):
        document, *_ = example()
        path = "word/document.xml/body[2]/p/r[0]/t[0]"
        leaf, paragraph, _ = _word_leaf(document, path)
        run = paragraph.findall(W + "r")[0]
        line_break = etree.Element(W + "br")
        run.insert(0, line_break)
        selected, _, _ = _word_leaf(document, path)
        self.assertIs(selected, leaf)
        selected.text = selected.text.replace("СИЛ: 15", "СИЛ: 75")
        self.assertEqual([child.tag for child in run], [W + "br", W + "t"])
        picture = etree.SubElement(run, W + "pict")
        with self.assertRaises(RevisionError):
            _word_leaf(document, path)
        run.remove(picture)
        line_break.set(W + "type", "page")
        with self.assertRaises(RevisionError):
            _word_leaf(document, path)
        line_break.attrib.clear()
        field = etree.SubElement(run, W + "fldChar")
        with self.assertRaises(RevisionError):
            _word_leaf(document, path)
        run.remove(field)
        properties = etree.Element(W + "rPr")
        etree.SubElement(properties, W + "vanish")
        run.insert(0, properties)
        with self.assertRaises(RevisionError):
            _word_leaf(document, path)

    def test_empty_word_text_round_trip_preserves_exact_xml_structure(self):
        document, *_ = example()
        leaf, _, _ = _word_leaf(document, "word/document.xml/body[2]/p/r[0]/t[0]")
        leaf.text = ""
        written = etree.fromstring(etree.tostring(document))
        checked, _, _ = _word_leaf(written, "word/document.xml/body[2]/p/r[0]/t[0]")
        predicted = etree.fromstring(etree.tostring(leaf, with_tail=False))
        self.assertEqual(checked.text or "", "")
        self.assertEqual(_node_xml_sha(checked), _node_xml_sha(predicted))

    def test_approved_consolidation_empty_cells_and_pair_topology(self):
        choices = {"decisions": [{"id": "len-pregen-persuade-consolidation",
                                  "disposition": "approved_local_skill_consolidation",
                                  "entities": ["Леон Стерндейл", "Коллин Палмер"],
                                  "source_values": {"Леон Стерндейл": {"Торг": 40, "Убеждение": 40},
                                                    "Коллин Палмер": {"Торг": 40, "Убеждение": 32}}}]}
        table = etree.Element(W + "tbl", nsmap={"w": NS["w"]})
        row = etree.SubElement(table, W + "tr")
        cells = []
        for value in ("Торг", "40%", "Убеждение", "32%"):
            cell = etree.SubElement(row, W + "tc")
            text_run(etree.SubElement(cell, W + "p"), value)
            cells.append(cell)
        paths = [f"word/document.xml/body[1]/tbl/tr[0]/tc[{i}]/p[0]/r[0]/t[0]" for i in range(4)]
        cases = [("clear_duplicate_label", "Торг", "", 0),
                 ("clear_duplicate_value", "40%", "", 1),
                 ("set_survivor_value", "32%", "40%", 3)]
        for action, before, after, index in cases:
            op = {"kind": "skill_consolidation", "field": "Убеждение",
                  "entity_id": "Коллин Палмер", "disposition": "convert_deterministic",
                  "before_text": before, "after_text": after,
                  "before_sha256": _sha(before.encode("utf-8")),
                  "reverse_before": after, "reverse_after": before,
                  "rule_id": "user_approved_local_skill_consolidation",
                  "editorial_decision": "len-pregen-persuade-consolidation", "action": action}
            _approved_field_rule(op, choices)
            _skill_consolidation_binding(table, op, paths[index])
            bad = dict(op, editorial_decision="missing")
            with self.assertRaises(RevisionError):
                _approved_field_rule(bad, choices)
        with self.assertRaises(RevisionError):
            _skill_consolidation_binding(table,
                {"kind": "skill_consolidation", "action": "clear_duplicate_value"}, paths[2])

    def test_approved_user_labels_and_exact_door_numeric(self):
        choices = {"decisions": [
            {"id": "kremlin-potapov-architecture", "entity": "Евгений Потапов",
             "disposition": "approved_contextual_reader_label"},
            {"id": "kremlin-language-parent-adaptation", "entities": ["Евгений Потапов"],
             "field_map": {"Русский язык": "Родной язык (Русский)"},
             "disposition": "approved_contextual_reader_label"}]}
        for field, before, after, decision in (
            ("architecture", "Архитектура", "Искусство/ремесло (Архитектура)",
             "kremlin-potapov-architecture"),
            ("language", "Русский язык", "Родной язык (Русский)",
             "kremlin-language-parent-adaptation"),
        ):
            op = {"kind": "user_label", "field": field, "entity_id": "Евгений Потапов",
                  "disposition": "convert_deterministic", "before_text": before,
                  "after_text": after, "before_sha256": _sha(before.encode("utf-8")),
                  "reverse_before": after, "reverse_after": before,
                  "rule_id": "user_approved_exact_label", "editorial_decision": decision}
            _approved_field_rule(op, choices)
            with self.assertRaises(RevisionError):
                _approved_field_rule(dict(op, after_text="wrong", reverse_before="wrong"), choices)
        door = {"kind": "characteristic_prose", "field": "СИЛ",
                "current_scale": "old_edition", "disposition": "convert_deterministic",
                "before_text": "СИЛ двери равна 10", "after_text": "СИЛ двери равна 50",
                "before_sha256": _sha("СИЛ двери равна 10".encode("utf-8")),
                "reverse_before": "СИЛ двери равна 50", "reverse_after": "СИЛ двери равна 10",
                "rule_id": "source_bound_door_str_times_five", "editorial_decision": None}
        _approved_field_rule(door, choices)
        with self.assertRaises(RevisionError):
            _approved_field_rule(dict(door, after_text="СИЛ двери равна 500",
                                      reverse_before="СИЛ двери равна 500"), choices)

    def test_end_to_end_synthetic_docx_only_field_route(self):
        with tempfile.TemporaryDirectory(prefix="source-field-test-") as temporary:
            root = Path(temporary)
            reviewed, source = root / "reviewed.docx", root / "source.bin"
            source.write_bytes(b"sealed original source")
            word = Document()
            heading = word.add_paragraph("Катерина, маленький вампир")
            bookmark = OxmlElement("w:bookmarkStart")
            bookmark.set(qn("w:id"), "1")
            bookmark.set(qn("w:name"), "story_one")
            heading._p.insert(0, bookmark)
            word.add_paragraph("СИЛ: 13 Удача: 75")
            stream = BytesIO()
            Image.new("RGB", (4, 4), "blue").save(stream, format="PNG")
            word.add_paragraph().add_run().add_picture(BytesIO(stream.getvalue()))
            word.save(reviewed)
            with zipfile.ZipFile(reviewed) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            document = _xml(parts["word/document.xml"], "synthetic Word")
            rels = _relationships(_xml(parts["word/_rels/document.xml.rels"], "synthetic rels"))
            leaf = document.xpath(".//w:body/w:p[2]/w:r/w:t", namespaces=NS)[0]
            before, after = leaf.text, "СИЛ: 65 Удача: 75"
            anchor = {"body_index": 0, "text": "Катерина, маленький вампир",
                      "text_sha256": _sha("Катерина, маленький вампир".encode("utf-8")),
                      "bookmarks": ["story_one"]}
            path = "word/document.xml/body[1]/p/r[0]/t[0]"
            operation = {"id": "str-13", "book_id": BOOK, "story_id": "story_one",
                         "entity_id": "Катерина", "entity_kind": "human_npc", "entity_anchor": anchor,
                         "field": "СИЛ", "kind": "characteristic", "disposition": "convert_deterministic",
                         "current_scale": "old_edition", "input_docx_sha256": _sha(reviewed.read_bytes()),
                         "source_verified": True,
                         "source_pdf_sha256": _sha(source.read_bytes()), "source_pdf_path": str(source),
                         "source_physical_page": 7, "source_edition_evidence": "fifth edition",
                         "selector": {"part": "word/document.xml", "path": path, "body_index": 1,
                                      "table_row": None, "table_cell": None,
                                      "text_node_before": before, "text_node_sha256": _sha(before.encode("utf-8")),
                                      "span_start": 0, "span_end": 7},
                         "before_text": "СИЛ: 13", "before_sha256": _sha("СИЛ: 13".encode("utf-8")),
                         "after_text": "СИЛ: 65", "reverse_before": "СИЛ: 65",
                         "reverse_after": "СИЛ: 13", "rule_id": "official_old_characteristic_times_five",
                         "editorial_decision": None, "node_group_ids": ["str-13"]}
            group = {"book_id": BOOK, "path": path, "operation_ids": ["str-13"],
                     "before_node": before, "before_node_sha256": _sha(before.encode("utf-8")),
                     "before_node_xml_sha256": _node_xml_sha(leaf),
                     "after_node": after, "after_node_sha256": _sha(after.encode("utf-8")),
                     "after_node_xml_sha256": _predicted_node_xml_sha(leaf, after),
                     "reverse_before_node": after, "reverse_after_node": before}
            ledger = {"schema": FIELD_LEDGER_SCHEMA, "state": FIELD_APPROVED_STATE,
                      "source_qa_closed": True, "orchestrator_approved": True,
                      "inputs": {BOOK: {"docx_sha256": _sha(reviewed.read_bytes()),
                                        "source_pdf_sha256": _sha(source.read_bytes()),
                                        "source_pdf_path": str(source)}},
                      "operation_count": 1, "operations": [operation], "node_groups": [group]}
            choices = {"schema": "nri.user_conversion_choices.v1", "decisions": []}
            authorities = {"ledger": json.dumps(ledger, ensure_ascii=False).encode("utf-8"),
                           "choices": json.dumps(choices).encode("utf-8"),
                           "conversion_spec": b"sealed official rules"}
            field_config = {name: {"path": str(root / f"{name}.json"), "sha256": _sha(payload)}
                            for name, payload in authorities.items()}
            field_config.update(book_id=BOOK, operation_ids=["str-13"], expression_ids=[])
            for name, payload in authorities.items():
                (root / f"{name}.json").write_bytes(payload)
            drawings = _drawing_order(document, rels, parts)
            self.assertEqual(len(drawings), 1)
            manifest = {"schema": "kb_creator.reviewed_docx_revision.v1",
                        "provenance": "reviewed_artifact_successor",
                        "inputs": {"docx_sha256": _sha(reviewed.read_bytes()),
                                   "source_pdf_sha256": _sha(source.read_bytes()),
                                   "authorities": [],
                                   "original_drawing_occurrences": [
                                       {"sha256": image_sha, "drawing_id": drawing_id}
                                       for image_sha, drawing_id in drawings]},
                        "stories": [{"story_id": "story_one", "word_bookmark": "story_one"}],
                        "source_graphic_ids": ["retained-picture"],
                        "source_graphics": [{"graphic_id": "retained-picture", "story_id": "story_one",
                                             "source_page": 1, "disposition": "must_include",
                                             "sha256": drawings[0][0], "occurrence_ordinal": 0,
                                             "drawing_id": drawings[0][1]}],
                        "operations": [], "source_bound_fields": field_config}
            manifest_path = root / "manifest.json"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            result = revise_docx(manifest_path, reviewed, source, root / "candidate", root / "report.json")
            self.assertFalse(result["accepted"])
            self.assertEqual(len(result["source_bound_field_nodes"]), 1)
            self.assertEqual([p.name for p in (root / "candidate").iterdir()], ["document.docx"])
            with zipfile.ZipFile(root / "candidate/document.docx") as archive:
                edited = _xml(archive.read("word/document.xml"), "edited Word")
                self.assertEqual(edited.xpath(".//w:body/w:p[2]/w:r/w:t/text()", namespaces=NS),
                                 ["СИЛ: 65 Удача: 75"])
            manifest["source_graphic_ids"] = []
            manifest["source_graphics"] = []
            manifest_path2 = root / "manifest-field-only.json"
            manifest_path2.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(RevisionError):
                revise_docx(manifest_path2, reviewed, source, root / "missing-scope", root / "missing-scope.json")
            manifest["art_scope"] = "preserve_original_drawings_only"
            manifest_path2.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            result2 = revise_docx(manifest_path2, reviewed, source, root / "field-only", root / "field-only.json")
            self.assertEqual(result2["source_graphic_coverage"], [])
            self.assertEqual(result2["preservation"]["old_drawing_occurrences"], 1)
            self.assertIn("source-art completeness was not assessed", result2["limitations"][-1])
            manifest["source_graphic_ids"] = ["retained-picture"]
            manifest["source_graphics"] = [{"graphic_id": "retained-picture", "story_id": "story_one",
                                             "source_page": 1, "disposition": "must_include",
                                             "sha256": drawings[0][0], "occurrence_ordinal": 0,
                                             "drawing_id": drawings[0][1]}]
            manifest_path2.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(RevisionError):
                revise_docx(manifest_path2, reviewed, source, root / "mixed-scope", root / "mixed-scope.json")

    def test_table_grouped_run_and_split_effect_preserve_protected_values(self):
        document, ledger, choices, config, snapshots, inputs = example()
        selected = plan(document, config, snapshots, inputs)
        self.assertEqual(len(selected), 5)
        for item in selected:
            item["leaf"].text = item["after"]
        text = document.xpath(".//w:t/text()", namespaces=NS)
        self.assertIn("СИЛ: 65 Удача: 75", text)
        self.assertIn("СИЛ: 75 ЛВК: 60 РАС: 70 Навык: 70", text)
        self.assertEqual("".join(text[-3:]), "5 × (1D6) СИЛ")
        self.assertTrue(document.xpath(".//w:rPr/w:b", namespaces=NS))

    def test_stale_alias_overlap_wrong_book_or_planning_ledger_fail(self):
        changes = (
            lambda d, l, c: l.update(state="planning_only_no_successor_authored"),
            lambda d, l, c: l.update(source_qa_closed=False),
            lambda d, l, c: l["operations"][0]["selector"].update(span_start=1),
            lambda d, l, c: l["operations"][2]["selector"].update(span_start=2),
            lambda d, l, c: l["node_groups"][1].update(path=l["node_groups"][0]["path"]),
            lambda d, l, c: l["node_groups"][0].update(before_node_xml_sha256="0" * 64),
            lambda d, l, c: l["operations"][0].update(node_group_ids=[]),
            lambda d, l, c: l["operations"][0].update(field="НАР"),
            lambda d, l, c: l["operations"][0].update(source_pdf_sha256="0" * 64),
            lambda d, l, c: l["operations"][0]["entity_anchor"].update(text="another entity"),
            lambda d, l, c: l["expression_groups"][0]["nodes"][1].update(after_node="changed"),
            lambda d, l, c: d.xpath(".//w:body/w:p[3]/w:r[1]", namespaces=NS)[0].insert(
                0, etree.Element(W + "fldChar")),
            lambda d, l, c: d.xpath(".//w:body/w:p[2]/w:r", namespaces=NS)[0].insert(
                0, etree.Element(W + "vanish")),
            lambda d, l, c: c["operation_ids"].pop(),
            lambda d, l, c: c["expression_ids"].pop(),
            lambda d, l, c: c.update(book_id="wrong-book"),
        )
        for number, change in enumerate(changes):
            with self.subTest(case=number):
                document, ledger, choices, config, snapshots, inputs = example()
                change(document, ledger, config)
                rebind(ledger, choices, config, snapshots)
                with self.assertRaises(RevisionError):
                    plan(document, config, snapshots, inputs)

    def test_edu_boundaries_and_protected_skill_scale(self):
        choices = {"decisions": []}
        template = {"disposition": "convert_deterministic", "kind": "characteristic",
                    "field": "ОБР", "current_scale": "old_edition", "entity_kind": "human_npc",
                    "editorial_decision": None}
        for old, new, rule in ((18, 90, "official_old_characteristic_times_five"),
                               (19, 91, "official_EDU_19_to_27_plus"),
                               (26, 98, "official_EDU_19_to_27_plus"),
                               (27, 99, "official_EDU_19_to_27_plus")):
            op = dict(template, before_text=str(old), after_text=str(new),
                      reverse_before=str(new), reverse_after=str(old),
                      before_sha256=_sha(str(old).encode("utf-8")), rule_id=rule)
            _approved_field_rule(op, choices)
        bad = dict(template, before_text="5", after_text="25", reverse_before="25",
                   reverse_after="5", before_sha256=_sha(b"5"),
                   current_scale="new_7e", rule_id="official_old_characteristic_times_five")
        with self.assertRaises(RevisionError):
            _approved_field_rule(bad, choices)
        bad_edu = dict(template, before_text="19", after_text="95", reverse_before="95",
                       reverse_after="19", before_sha256=_sha(b"19"),
                       rule_id="official_old_characteristic_times_five")
        with self.assertRaises(RevisionError):
            _approved_field_rule(bad_edu, choices)
        relabeled_edu = dict(template, before_text="ОБР: 19", after_text="Удача: 91",
                            reverse_before="Удача: 91", reverse_after="ОБР: 19",
                            before_sha256=_sha("ОБР: 19".encode("utf-8")),
                            rule_id="official_EDU_19_to_27_plus")
        with self.assertRaises(RevisionError):
            _approved_field_rule(relabeled_edu, choices)
        protected = dict(template, field="Удача", before_text="75", after_text="375",
                         reverse_before="375", reverse_after="75",
                         before_sha256=_sha(b"75"), rule_id="official_old_characteristic_times_five")
        with self.assertRaises(RevisionError):
            _approved_field_rule(protected, choices)
        skill = dict(template, kind="skill_label", field="skill:Пистолет", before_text="Пистолет 70%",
                     after_text="Стрельба (пистолет) 350%", reverse_before="Стрельба (пистолет) 350%",
                     reverse_after="Пистолет 70%", before_sha256=_sha("Пистолет 70%".encode("utf-8")),
                     rule_id="official_skill_mapping_392")
        with self.assertRaises(RevisionError):
            _approved_field_rule(skill, choices)

    def test_explicit_user_crosswalk_and_prose_authority(self):
        choices = {"decisions": [
            {"id": "kremlin-chr-app-crosswalk", "disposition": "approved_reader_adaptation",
             "user_answer": "ХАР к НАР ×5"},
            {"id": "len-stealth-context-values", "disposition": "approved_contextual_compatibility",
             "old": "Бесшумное передвижение 90%, Укрытие 75%",
             "new": "Скрытность: 90% при передвижении, 75% при укрытии",
             "user_answer": "retain both"},
            {"id": "len-rybin-bodily-conceal", "disposition": "approved_contextual_skill_mapping",
             "user_answer": "Скрытность"}]}
        crosswalk = {"kind": "characteristic", "field": "ХАР", "disposition": "convert_deterministic",
                     "current_scale": "old_edition", "before_text": "ХАР: 8", "after_text": "НАР: 40",
                     "before_sha256": _sha("ХАР: 8".encode("utf-8")),
                     "reverse_before": "НАР: 40", "reverse_after": "ХАР: 8",
                     "rule_id": "user_approved_CHR_to_APP_times_five",
                     "editorial_decision": "kremlin-chr-app-crosswalk"}
        _approved_field_rule(crosswalk, choices)
        with self.assertRaises(RevisionError):
            _approved_field_rule(dict(crosswalk, after_text="НАР: 40 Удача",
                                      reverse_before="НАР: 40 Удача"), choices)
        contextual = {"kind": "skill_label", "field": "contextual_stealth",
                      "disposition": "convert_deterministic",
                      "before_text": "Бесшумное передвижение 90%, Укрытие 75%",
                      "after_text": "Скрытность: 90% при передвижении, 75% при укрытии",
                      "before_sha256": _sha("Бесшумное передвижение 90%, Укрытие 75%".encode("utf-8")),
                      "reverse_before": "Скрытность: 90% при передвижении, 75% при укрытии",
                      "reverse_after": "Бесшумное передвижение 90%, Укрытие 75%",
                      "rule_id": "user_approved_contextual_stealth",
                      "editorial_decision": "len-stealth-context-values"}
        _approved_field_rule(contextual, choices)
        prose = {"kind": "prose_label", "field": "rybin_bodily_stealth",
                 "disposition": "convert_deterministic", "before_text": "Сокрытие +10%",
                 "after_text": "Скрытность +10%", "before_sha256": _sha("Сокрытие +10%".encode("utf-8")),
                 "reverse_before": "Скрытность +10%", "reverse_after": "Сокрытие +10%",
                 "rule_id": "user_approved_exact_prose",
                 "editorial_decision": "len-rybin-bodily-conceal"}
        _approved_field_rule(prose, choices)
        with self.assertRaises(RevisionError):
            _approved_field_rule(dict(prose, after_text="Скрытность +50%",
                                      reverse_before="Скрытность +50%"), choices)
        with self.assertRaises(RevisionError):
            _approved_field_rule(dict(crosswalk, editorial_decision="missing"), choices)


if __name__ == "__main__":
    unittest.main()
